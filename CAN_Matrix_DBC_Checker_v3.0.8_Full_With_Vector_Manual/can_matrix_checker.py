#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
CAN 通信矩阵与 DBC 一致性检查工具

特点：
1. 图形界面，支持 .xlsx/.xlsm/.csv 通信矩阵与 .dbc 文件。
2. 自动识别常见中英文列名。
3. 检查报文、信号及 DBC/矩阵自身结构问题。
4. 导出 Excel 差异报告；未安装 openpyxl 时可导出 CSV。
5. 使用 can_rules.json 执行可配置的 CAN/AUTOSAR 语义规则检查。
6. 支持“无CAN矩阵”模式，仅凭DBC结构、属性和保守经验规则完成预检查。
7. 支持导入客户E2E ID表（TX/Rx SignalGroup + DATA ID）并核对DBC E2EDataId。
8. E2E映射窗口直接显示客户值/DBC值/一致性结果，并支持按问题类型筛选主结果。
9. E2E ID表支持同一工作表多路TX/RX列、多个工作表汇总导入，以及客户表部分缺项。

Python: 3.9+
可选依赖: openpyxl（读取 xlsx、导出 xlsx 报告）
"""

from __future__ import annotations

import csv
import json
import hashlib
import math
import itertools
import queue
import os
import shutil
import tempfile
from datetime import datetime
import re
import sys
import threading
import traceback
import textwrap
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Callable, Dict, Iterable, List, Optional, Sequence, Tuple

try:
    import tkinter as tk
    from tkinter import filedialog, messagebox, simpledialog, ttk
except ImportError:  # pragma: no cover - 允许无GUI环境运行核心解析/自测
    tk = None
    filedialog = None
    messagebox = None
    simpledialog = None
    ttk = None

from dbc_transform import (
    apply_dbc_repair_plan,
    apply_node_completion,
    apply_rename_plan,
    build_dbc_repair_plan,
    build_node_completion_plan,
    build_rename_plan,
    default_rename_config,
    load_rename_config,
    object_key,
    save_rename_config,
    signal_key,
)

try:
    import openpyxl
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
except ImportError:
    openpyxl = None
    Workbook = None
    load_workbook = None


APP_NAME = "CAN矩阵-DBC一致性检查工具"
APP_VERSION = "3.0.9"
VECTOR_MANUAL_VERSION = "1.12"
VECTOR_MANUAL_TITLE = "Vector Rules for Legacy Communication Descriptions"


def text(value: Any) -> str:
    """将单元格或解析值转为干净字符串。"""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def norm_header(value: Any) -> str:
    """列名归一化：忽略大小写、空白和常见分隔符。"""
    return re.sub(r"[\s_\-./\\()（）\[\]【】:：]+", "", text(value)).lower()


def norm_name(value: Any) -> str:
    """名称匹配归一化，不改变报告中显示的原始名称。"""
    return re.sub(r"[^0-9a-zA-Z\u4e00-\u9fff]+", "", text(value)).lower()


def norm_enum(value: Any) -> str:
    return re.sub(r"[\s_\-./\\]+", "", text(value)).lower()


def parse_float(value: Any) -> Optional[float]:
    s = text(value)
    if not s:
        return None
    s = s.replace(",", "").strip()

    # 通信矩阵中的初始值、无效值和枚举值经常使用 0x14、-0x1 等十六进制形式。
    hex_match = re.fullmatch(r"([+-]?)0[xX]([0-9a-fA-F]+)", s)
    if hex_match:
        sign = -1 if hex_match.group(1) == "-" else 1
        return float(sign * int(hex_match.group(2), 16))

    try:
        return float(s)
    except ValueError:
        # 优先提取文本中完整的十六进制数，避免把 0x14 错读为 0。
        hex_search = re.search(r"([+-]?)0[xX]([0-9a-fA-F]+)", s)
        if hex_search:
            sign = -1 if hex_search.group(1) == "-" else 1
            return float(sign * int(hex_search.group(2), 16))
        match = re.search(r"[-+]?\d+(?:\.\d+)?(?:[eE][-+]?\d+)?", s)
        if match:
            try:
                return float(match.group(0))
            except ValueError:
                return None
    return None


def parse_int(value: Any) -> Optional[int]:
    number = parse_float(value)
    if number is None:
        return None
    return int(number)


def parse_can_id(value: Any) -> Optional[int]:
    """解析十六进制/十进制 CAN ID。"""
    s = text(value)
    if not s:
        return None
    s = s.replace("_", "").replace(" ", "")
    # Excel 可能把 0x520 当作文本，也可能直接存为 1312。
    try:
        if s.lower().startswith("0x"):
            return int(s, 16)
        if s.lower().endswith("h") and re.fullmatch(r"[0-9a-fA-F]+h", s):
            return int(s[:-1], 16)
        # 含 A-F 的纯字符串按十六进制处理。
        if re.fullmatch(r"[0-9a-fA-F]+", s) and re.search(r"[a-fA-F]", s):
            return int(s, 16)
        return int(float(s))
    except (ValueError, TypeError):
        match = re.search(r"0[xX]([0-9a-fA-F]+)", s)
        if match:
            return int(match.group(1), 16)
    return None


def format_can_id(can_id: Optional[int]) -> str:
    return "" if can_id is None else f"0x{can_id:X}"


def format_can_id_decimal(can_id: Optional[int]) -> str:
    """显示规范化后的十进制 CAN ID，便于在 DBC、代码和配置中搜索。"""
    return "" if can_id is None else str(can_id)


def parse_bool_signed(value: Any) -> Optional[bool]:
    s = norm_enum(value)
    if not s:
        return None
    unsigned_values = {
        "unsigned", "uint", "u", "无符号", "无符号数", "0", "+",
        "false", "否", "no",
    }
    signed_values = {
        "signed", "int", "s", "有符号", "有符号数", "1", "-",
        "true", "是", "yes",
    }
    if s in unsigned_values or "unsigned" in s or "无符号" in s:
        return False
    if s in signed_values or ("signed" in s and "unsigned" not in s) or "有符号" in s:
        return True
    return None


def parse_byte_order(value: Any) -> Optional[str]:
    s = norm_enum(value)
    if not s:
        return None
    intel = {"intel", "littleendian", "little", "小端", "小端序", "0", "lsb"}
    motorola = {"motorola", "bigendian", "big", "大端", "大端序", "1", "msb"}
    if s in intel or "intel" in s or "little" in s or "小端" in s:
        return "intel"
    if s in motorola or "motorola" in s or "big" in s or "大端" in s:
        return "motorola"
    return text(value).lower()


def parse_frame_format(value: Any) -> Optional[str]:
    s = norm_enum(value)
    if not s:
        return None
    if any(k in s for k in ("extended", "ext", "扩展", "29bit", "29位")):
        return "extended"
    if any(k in s for k in ("standard", "std", "标准", "11bit", "11位")):
        return "standard"
    return text(value).lower()


def parse_bus_format(value: Any) -> Optional[str]:
    """识别 Classical CAN / CAN FD。"""
    s = norm_enum(value)
    if not s:
        return None
    if "canfd" in s or "fdcan" in s or re.search(r"(^|[^a-z])fd([^a-z]|$)", text(value).lower()):
        return "can_fd"
    if "can" in s or any(k in s for k in ("standard", "extended", "标准", "扩展")):
        return "classic"
    return None


def parse_yes(value: Any) -> Optional[bool]:
    s = norm_enum(value)
    if not s:
        return None
    if s in {"yes", "true", "1", "on", "enable", "enabled", "是", "启用"}:
        return True
    if s in {"no", "false", "0", "off", "disable", "disabled", "否", "禁用"}:
        return False
    return None


def split_nodes(value: Any) -> Tuple[str, ...]:
    s = text(value)
    if not s:
        return tuple()
    parts = re.split(r"[,，;；/|\s]+", s)
    cleaned = sorted({p.strip() for p in parts if p.strip() and p.strip() != "Vector__XXX"}, key=str.lower)
    return tuple(cleaned)


def nearly_equal(a: Optional[float], b: Optional[float], rel_tol: float = 1e-9, abs_tol: float = 1e-9) -> bool:
    if a is None or b is None:
        return a is b
    return math.isclose(a, b, rel_tol=rel_tol, abs_tol=abs_tol)


def value_display(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "有符号" if value else "无符号"
    if isinstance(value, tuple):
        return ", ".join(value)
    if isinstance(value, float):
        return f"{value:.12g}"
    return str(value)


def _signal_region_candidate(signals: Sequence["Signal"]) -> Optional[Tuple[int, int, int, bool]]:
    """按一组信号在PDU中覆盖的字节区间估算序列化区长度。

    返回 (candidate_bits, start_byte, end_byte, fully_filled)。
    candidate_bits 使用“首个占用字节到最后占用字节”的完整字节跨度，
    这比简单相加各信号位长更接近 SignalGroupArray / transformationBuffer 的长度。
    """
    all_bits: set[int] = set()
    for sig in signals:
        bits = occupied_bits(sig)
        if bits:
            all_bits.update(bits)
    if not all_bits:
        return None
    start_byte = min(all_bits) // 8
    end_byte = max(all_bits) // 8
    candidate_bits = (end_byte - start_byte + 1) * 8
    byte_region = set(range(start_byte * 8, (end_byte + 1) * 8))
    fully_filled = all_bits == byte_region
    return candidate_bits, start_byte, end_byte, fully_filled


def e2e_data_length_suggestion(msg: "Message") -> str:
    """给 E2EDataLength 提供分层候选建议，不替代项目/Profile规范。

    优先级：
    1) 报文只有一个 SignalGroup：按该组成员覆盖的字节区间估算；
    2) 报文有多个 SignalGroup：列出各组候选值，不擅自选一个；
    3) 没有可用 SignalGroup：退回整帧 DLC*8 作为低置信度兜底。

    生成代码可进一步复核：若对应 RTE/Transformer 缓冲区为
    transformationBuffer[N]，通常应重点核对 N*8 bit 是否与 E2EDataLength 一致。
    """
    signal_by_name = {sig.name: sig for sig in msg.signals}
    group_candidates: List[Tuple[str, int, int, int, bool, List[str]]] = []
    for group_name, member_names in msg.signal_groups.items():
        members = [signal_by_name[name] for name in member_names if name in signal_by_name]
        candidate = _signal_region_candidate(members)
        if candidate is None:
            continue
        bits, start_byte, end_byte, fully_filled = candidate
        group_candidates.append((group_name, bits, start_byte, end_byte, fully_filled, [s.name for s in members]))

    if len(group_candidates) == 1:
        group_name, bits, start_byte, end_byte, fully_filled, members = group_candidates[0]
        byte_len = bits // 8
        fill_note = "组成员完整覆盖该字节区间" if fully_filled else "组成员在该字节区间内存在空洞/保留位，仍按序列化字节跨度估算"
        return (
            f"建议候选值：{bits} bit（优先按 SignalGroup={group_name} 推算："
            f"覆盖PDU字节{start_byte}~{end_byte}，共{byte_len} Byte；{fill_note}）。"
            f"建议再用生成代码复核：若对应 transformationBuffer[{byte_len}]，则 {byte_len}×8={bits} bit 与该候选值一致。"
        )

    if len(group_candidates) > 1:
        details = []
        for group_name, bits, start_byte, end_byte, _fully_filled, _members in group_candidates:
            details.append(f"{group_name}={bits} bit({bits // 8} Byte, PDU字节{start_byte}~{end_byte})")
        fallback = f"；整帧兜底={int(msg.dlc) * 8} bit" if msg.dlc is not None and msg.dlc > 0 else ""
        return (
            "检测到多个SignalGroup，无法仅凭DBC唯一确定哪一组被E2E保护。"
            f"候选：{'；'.join(details)}{fallback}。"
            "请结合E2E属性映射或生成代码中的 transformationBuffer[N] 复核，N×8 即可作为重点候选。"
        )

    # 没有显式SignalGroup时，如果仅少量信号带有E2E属性，可按这些信号的字节跨度给出辅助候选。
    e2e_signals = [sig for sig in msg.signals if any(key.startswith("e2e") for key in sig.attributes)]
    candidate = _signal_region_candidate(e2e_signals) if e2e_signals else None
    if candidate is not None:
        bits, start_byte, end_byte, _fully_filled = candidate
        byte_len = bits // 8
        names = ", ".join(sig.name for sig in e2e_signals[:5])
        suffix = "..." if len(e2e_signals) > 5 else ""
        return (
            f"建议候选值：{bits} bit（未找到SignalGroup，但检测到带E2E属性的信号 {names}{suffix}，"
            f"其PDU字节跨度为{start_byte}~{end_byte}，共{byte_len} Byte）。"
            f"请用生成代码 transformationBuffer[N] 复核；若N={byte_len}，则候选值为{bits} bit。"
        )

    if msg.dlc is None or msg.dlc <= 0:
        return (
            "建议值：无法自动计算（未找到可用SignalGroup且报文DLC缺失或为0）。"
            "请在生成代码中查看对应 transformationBuffer[N]，优先用 N×8 bit 复核E2EDataLength。"
        )

    bits = int(msg.dlc) * 8
    return (
        f"低置信度兜底候选：{bits} bit（未找到可唯一识别的E2E SignalGroup，按整帧 DLC={msg.dlc} Byte × 8 估算）。"
        "如果生成代码存在 transformationBuffer[N]，应优先用 N×8 bit 作为实际保护长度候选，而不是机械采用整帧值。"
    )



def e2e_data_length_auto_candidate(msg: "Message") -> Optional[Tuple[int, str]]:
    """返回可安全用于“一键修改”的 E2EDataLength 候选值。

    只接受能够从 DBC 内部唯一确定的高置信度场景：
    - 报文中只有一个可解析 SignalGroup，并能计算出其序列化字节跨度。

    多 SignalGroup、仅凭 DLC*8、或仅凭零散信号推测的情况不自动写回，
    仍只在检查报告中给建议，避免自动修改正确的项目配置。
    """
    signal_by_name = {sig.name: sig for sig in msg.signals}
    candidates: List[Tuple[str, int, int, int]] = []
    for group_name, member_names in msg.signal_groups.items():
        members = [signal_by_name[name] for name in member_names if name in signal_by_name]
        candidate = _signal_region_candidate(members)
        if candidate is None:
            continue
        bits, start_byte, end_byte, _fully_filled = candidate
        candidates.append((group_name, bits, start_byte, end_byte))

    if len(candidates) != 1:
        return None

    group_name, bits, start_byte, end_byte = candidates[0]
    if bits <= 0:
        return None
    reason = (
        f"SignalGroup={group_name} 唯一可确定，覆盖PDU字节{start_byte}~{end_byte}，"
        f"共{bits // 8} Byte，因此候选 E2EDataLength={bits} bit。"
    )
    return bits, reason


def collect_e2e_data_length_autofixes(db: "Database") -> List[Tuple["Message", int, str]]:
    """收集可以一键修复的报文，仅处理已启用E2E且DataLength缺失/<=0的情况。"""
    fixes: List[Tuple[Message, int, str]] = []
    for msg in db.messages:
        _profile_name, profile_value, _ = effective_attribute_any(msg, db, ("E2EProfile",))
        _length_name, current_value, _ = effective_attribute_any(msg, db, ("E2EDataLength",))
        if profile_value in (None, "", 0, "0", "No", "Off"):
            continue
        current_num = parse_float(current_value)
        if current_num is not None and current_num > 0:
            continue
        candidate = e2e_data_length_auto_candidate(msg)
        if candidate is None:
            continue
        bits, reason = candidate
        fixes.append((msg, int(bits), reason))
    return fixes


def _read_dbc_text_preserve_encoding(path: str) -> Tuple[str, str]:
    raw = Path(path).read_bytes()
    # 只有原文件本身带UTF-8 BOM时才用utf-8-sig写回，避免自动修改后凭空新增BOM。
    if raw.startswith(b"\xef\xbb\xbf"):
        return raw.decode("utf-8-sig"), "utf-8-sig"
    for encoding in ("utf-8", "gb18030", "latin-1"):
        try:
            return raw.decode(encoding), encoding
        except UnicodeDecodeError:
            continue
    raise DbcParseError("DBC 文本编码无法识别，无法自动修改。")


def apply_e2e_data_length_autofixes(
    dbc_path: str,
    db: "Database",
    fixes: Sequence[Tuple["Message", int, str]],
) -> Tuple[str, List[str]]:
    """把高置信度 E2EDataLength 候选值写回当前DBC，并自动创建时间戳备份。

    返回 (backup_path, change_descriptions)。
    """
    if not fixes:
        raise ValueError("没有可自动修改的E2EDataLength。")

    path = Path(dbc_path)
    content, encoding = _read_dbc_text_preserve_encoding(dbc_path)
    newline = "\r\n" if "\r\n" in content else "\n"
    lines = content.splitlines(keepends=True)

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S_%f")
    backup_path = path.with_name(f"{path.stem}.before_E2E_fix_{timestamp}{path.suffix}")
    shutil.copy2(path, backup_path)

    # 现有显式赋值的行号（parse_dbc为1-based）。
    usage_lines: Dict[int, List[int]] = {}
    for usage in db.attribute_usages:
        if usage.scope.upper() == "BO" and usage.name.lower() == "e2edatalength" and usage.can_id is not None and usage.line_number:
            usage_lines.setdefault(int(usage.can_id), []).append(int(usage.line_number))

    changes: List[str] = []
    append_lines: List[str] = []
    assignment_re = re.compile(
        r'^(\s*BA_\s+"E2EDataLength"\s+BO_\s+)(\d+)(\s+)(.+?)(\s*;\s*)(\r?\n)?$',
        re.IGNORECASE,
    )

    for msg, bits, reason in fixes:
        if msg.can_id is None:
            continue
        replaced = False
        for line_no in usage_lines.get(int(msg.can_id), []):
            idx = line_no - 1
            if idx < 0 or idx >= len(lines):
                continue
            original = lines[idx]
            m = assignment_re.match(original)
            if not m:
                continue
            ending = m.group(6) or ""
            lines[idx] = f"{m.group(1)}{m.group(2)}{m.group(3)}{bits}{m.group(5)}{ending}"
            replaced = True

        if not replaced:
            # 从 BO_ 定义行取得DBC原始ID（扩展帧可能带0x80000000标志）。
            raw_id: Optional[int] = None
            if msg.row_number and 0 < msg.row_number <= len(lines):
                bo_match = re.match(r"^\s*BO_\s+(\d+)\s+", lines[msg.row_number - 1])
                if bo_match:
                    raw_id = int(bo_match.group(1))
            if raw_id is None:
                raw_id = int(msg.can_id)
                if msg.id_frame_format == "extended" and raw_id <= 0x1FFFFFFF:
                    raw_id |= 0x80000000
            append_lines.append(f'BA_ "E2EDataLength" BO_ {raw_id} {bits};{newline}')

        changes.append(
            f"0x{msg.can_id:X} ({msg.can_id}) {msg.name}: E2EDataLength -> {bits} bit；{reason}"
        )

    if append_lines:
        if lines and not (lines[-1].endswith("\n") or lines[-1].endswith("\r")):
            lines[-1] = lines[-1] + newline
        if lines and lines[-1].strip():
            lines.append(newline)
        lines.extend(append_lines)

    try:
        path.write_text("".join(lines), encoding=encoding, newline="")
    except Exception:
        # 写回失败时尽量恢复原文件。
        shutil.copy2(backup_path, path)
        raise

    return str(backup_path), changes


def _dbc_attribute_storage_value(db: "Database", attribute_name: str, value: Any) -> str:
    """把检查器中的属性值编码回现有 BA_DEF_ 所定义的 DBC 存储形式。"""
    definition = definition_for(db, "BO", attribute_name)
    if definition is None:
        raise ValueError(f"属性“{attribute_name}”缺少 BO_ 级 BA_DEF_ 定义，不能猜测类型后写回。")
    if definition.value_type == "ENUM":
        normalized = norm_enum(value)
        for index, enum_value in enumerate(definition.enum_values):
            if norm_enum(enum_value) == normalized:
                return str(index)
        raise ValueError(f"属性“{attribute_name}”的值“{value_display(value)}”不在枚举定义中。")
    if definition.value_type in {"INT", "HEX"}:
        number = parse_int(value)
        if number is None:
            raise ValueError(f"属性“{attribute_name}”需要整数值。")
        return str(number)
    if definition.value_type == "FLOAT":
        number = parse_float(value)
        if number is None:
            raise ValueError(f"属性“{attribute_name}”需要数值。")
        return value_display(number)
    if definition.value_type == "STRING":
        return '"' + str(value).replace('"', '\\"') + '"'
    raise ValueError(f"属性“{attribute_name}”的类型“{definition.value_type}”暂不支持自动写回。")


def apply_message_attribute_autofix(
    dbc_path: str,
    db: "Database",
    message: "Message",
    attribute_name: str,
    value: Any,
) -> Tuple[str, str]:
    """对唯一报文补全或替换一个已定义的 BO_ 属性，并保留时间戳备份。"""
    if message.can_id is None or not message.row_number:
        raise ValueError("报文没有可定位的 CAN ID/BO_ 行，不能自动写回。")
    storage_value = _dbc_attribute_storage_value(db, attribute_name, value)
    path = Path(dbc_path)
    content, encoding = _read_dbc_text_preserve_encoding(dbc_path)
    lines = content.splitlines(keepends=True)
    row_index = int(message.row_number) - 1
    if row_index < 0 or row_index >= len(lines):
        raise ValueError("BO_ 行号已失效；请重新检查后再修复。")
    bo_match = re.match(r"^\s*BO_\s+(\d+)\s+", lines[row_index])
    if not bo_match:
        raise ValueError("BO_ 行内容已变化；请重新检查后再修复。")
    raw_id = int(bo_match.group(1))
    assignment_re = re.compile(
        rf'^(\s*BA_\s+"{re.escape(attribute_name)}"\s+BO_\s+)({raw_id})(\s+)(.+?)(\s*;\s*)(\r?\n)?$',
        re.IGNORECASE,
    )
    changed = False
    for index, line in enumerate(lines):
        match = assignment_re.match(line)
        if not match:
            continue
        lines[index] = f"{match.group(1)}{match.group(2)}{match.group(3)}{storage_value}{match.group(5)}{match.group(6) or ''}"
        changed = True
        break
    if not changed:
        newline = "\r\n" if "\r\n" in content else "\n"
        if lines and not (lines[-1].endswith("\n") or lines[-1].endswith("\r")):
            lines[-1] += newline
        lines.append(f'BA_ "{attribute_name}" BO_ {raw_id} {storage_value};{newline}')

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S_%f")
    backup_path = path.with_name(f"{path.stem}.before_{attribute_name}_{timestamp}{path.suffix}")
    shutil.copy2(path, backup_path)
    temp_handle, temp_name = tempfile.mkstemp(prefix=f".{path.name}.", suffix=".tmp", dir=str(path.parent))
    os.close(temp_handle)
    temp_path = Path(temp_name)
    try:
        temp_path.write_text("".join(lines), encoding=encoding, newline="")
        os.replace(temp_path, path)
    except Exception:
        shutil.copy2(backup_path, path)
        raise
    finally:
        temp_path.unlink(missing_ok=True)
    action = "更新" if changed else "补全"
    return str(backup_path), f"{action} {message.name} 的 {attribute_name}={storage_value}。"


@dataclass
class AttributeDefinition:
    name: str
    scope: str
    value_type: str
    minimum: Optional[float] = None
    maximum: Optional[float] = None
    enum_values: Tuple[str, ...] = field(default_factory=tuple)


@dataclass
class AttributeUsage:
    name: str
    scope: str
    object_name: str
    can_id: Optional[int]
    signal_name: str
    raw_value: str
    decoded_value: Any
    line_number: Optional[int] = None


@dataclass
class Signal:
    name: str
    start_bit: Optional[int] = None
    # dbc：DBC/Vector Motorola MSB起始位；lsb：矩阵以信号LSB位置表示起始位。
    start_bit_convention: str = "dbc"
    length: Optional[int] = None
    byte_order: Optional[str] = None
    signed: Optional[bool] = None
    factor: Optional[float] = None
    offset: Optional[float] = None
    minimum: Optional[float] = None
    maximum: Optional[float] = None
    unit: str = ""
    receivers: Tuple[str, ...] = field(default_factory=tuple)
    initial_value: Optional[float] = None
    invalid_value: Optional[float] = None
    comment: str = ""
    value_table: Dict[int, str] = field(default_factory=dict)
    multiplex: str = ""
    row_number: Optional[int] = None
    attributes: Dict[str, Any] = field(default_factory=dict)
    explicit_attributes: set[str] = field(default_factory=set)


@dataclass
class Message:
    can_id: Optional[int]
    name: str
    dlc: Optional[int] = None
    sender: str = ""
    cycle_time: Optional[float] = None
    send_type: str = ""
    frame_format: Optional[str] = None
    bus_format: Optional[str] = None
    id_frame_format: Optional[str] = None
    comment: str = ""
    signals: List[Signal] = field(default_factory=list)
    signal_groups: Dict[str, Tuple[str, ...]] = field(default_factory=dict)
    row_number: Optional[int] = None
    attributes: Dict[str, Any] = field(default_factory=dict)
    explicit_attributes: set[str] = field(default_factory=set)


@dataclass
class Database:
    source: str
    messages: List[Message]
    nodes: Tuple[str, ...] = field(default_factory=tuple)
    node_attributes: Dict[str, Dict[str, Any]] = field(default_factory=dict)
    global_attributes: Dict[str, Any] = field(default_factory=dict)
    attribute_definitions: Dict[Tuple[str, str], AttributeDefinition] = field(default_factory=dict)
    attribute_defaults: Dict[str, Any] = field(default_factory=dict)
    attribute_usages: List[AttributeUsage] = field(default_factory=list)
    syntax_issues: List[Tuple[int, str, str]] = field(default_factory=list)


@dataclass
class Difference:
    severity: str
    category: str
    can_id: Optional[int]
    message_name: str
    signal_name: str
    field_name: str
    matrix_value: str
    dbc_value: str
    description: str
    rule_id: str = ""


@dataclass
class E2EIdEntry:
    """客户E2E ID表中的一条 SignalGroup -> Data ID 映射。"""
    direction: str
    group_name: str
    data_id: Optional[int]
    row_number: int
    sheet_name: str = ""


@dataclass
class E2EComparisonRow:
    """用于“查看E2E映射”窗口的直观核对结果。"""
    direction: str
    customer_group: str
    dbc_group: str
    can_id: Optional[int]
    message_name: str
    customer_data_id: Optional[int]
    dbc_data_ids: Tuple[int, ...]
    status: str
    note: str = ""


def definition_for(db: Database, scope: str, name: str) -> Optional[AttributeDefinition]:
    key = name.lower()
    return db.attribute_definitions.get((scope.upper(), key)) or db.attribute_definitions.get(("GLOBAL", key))


def effective_attribute(obj: Any, db: Database, name: str) -> Tuple[Any, bool]:
    key = name.lower()
    attributes = getattr(obj, "attributes", {})
    explicit = getattr(obj, "explicit_attributes", set())
    if key in attributes:
        return attributes[key], key in explicit
    if key in db.attribute_defaults:
        return db.attribute_defaults[key], False
    return None, False


def effective_attribute_any(obj: Any, db: Database, names: Sequence[str]) -> Tuple[str, Any, bool]:
    for name in names:
        value, explicit = effective_attribute(obj, db, name)
        if value not in (None, ""):
            return name, value, explicit
    return "", None, False


def message_kind_flags(msg: Message, db: Database) -> Dict[str, bool]:
    haystack = f"{msg.name} {msg.comment}"
    _nm_name, nm_value, _ = effective_attribute_any(msg, db, ("NmMessage", "NMAsrMessage", "NmhMessage"))
    _msg_type_name, msg_type, _ = effective_attribute_any(msg, db, ("MsgType",))
    diag_values = [effective_attribute(msg, db, name)[0] for name in ("DiagRequest", "DiagResponse", "DiagState")]
    msg_type_text = norm_enum(msg_type)
    return {
        "nm": bool(re.search(r"(?i)(^|[_\-])(can)?nm([_\-]|$)|network\s*management", haystack))
              or parse_yes(nm_value) is True or msg_type_text in {"nm", "nmh"} or msg_type_text.startswith("nm"),
        "diag": bool(re.search(r"(?i)(diag|diagnostic|uds|obd|isotp|iso.?tp)", haystack))
                or any(parse_yes(v) is True for v in diag_values) or "isotp" in msg_type_text,
        "xcp": bool(re.search(r"(?i)(^|[_\-])xcp([_\-]|$)", haystack)) or "xcp" in msg_type_text,
    }


# 兼容不同公司常见通信矩阵列名。可继续在此扩展。
COLUMN_ALIASES: Dict[str, Sequence[str]] = {
    "message_id": (
        "报文ID", "帧ID", "消息ID", "CAN ID", "CANID", "Message ID", "Msg ID", "Frame ID",
        "Identifier", "CAN Identifier", "报文标识符", "ID",
    ),
    "message_name": (
        "报文名称", "帧名称", "消息名称", "Message Name", "Msg Name", "Frame Name", "报文名", "帧名",
    ),
    "dlc": (
        "DLC", "报文长度", "帧长度", "Message Length", "Frame Length", "Data Length", "Length(Byte)",
        "报文字节数", "字节长度",
    ),
    "sender": (
        "发送节点", "发送方", "发送ECU", "Tx Node", "Transmitter", "Sender", "发送者", "Source Node",
    ),
    "receiver": (
        "接收节点", "接收方", "接收ECU", "Rx Node", "Receiver", "Receivers", "接收者", "Destination Node",
    ),
    "cycle_time": (
        "周期时间", "报文周期", "周期(ms)", "周期时间(ms)", "Cycle Time", "Cycle Time(ms)", "Cycle Time (ms)", "CycleTime(ms)", "Cycle", "Period", "Period Time",
        "Msg Cycle", "Message Cycle", "Cycle/ms",
    ),
    "send_type": (
        "发送类型", "报文发送类型", "Tx Type", "Send Type", "Message Send Type", "Transmission Mode",
    ),
    "frame_format": (
        "帧格式", "报文格式", "Frame Format", "Message Type", "ID Type", "CAN ID Type", "标准扩展帧", "Standard/Extended",
    ),
    "message_comment": (
        "报文描述", "报文注释", "Message Comment", "Message Description", "Frame Comment", "报文说明",
    ),
    "signal_name": (
        "信号名称", "信号名", "Signal Name", "Sig Name", "Signal", "信号",
    ),
    "start_bit": (
        "起始位", "开始位", "Start Bit", "StartBit", "Start Bit(LSB Or MSB)", "Start Bit (LSB Or MSB)",
        "Start Bit(MSB Or LSB)", "Bit Start", "Start Position", "起始Bit", "起始位(bit)",
    ),
    "start_byte": (
        "起始字节", "Start Byte", "Byte Start", "StartByte", "起始Byte",
    ),
    "bit_in_byte": (
        "字节内起始位", "Byte Bit", "Bit In Byte", "Start Bit In Byte", "位偏移", "Bit Offset",
    ),
    "signal_length": (
        "信号长度", "位长度", "Signal Length", "Sig Length", "Signal Size(bit)", "Signal Size (bit)",
        "SignalSize(bit)", "Signal Size", "Bit Length", "Length(bit)", "长度(bit)", "Length",
    ),
    "byte_order": (
        "字节序", "端序", "Byte Order", "Endian", "Endianness", "Intel/Motorola", "排列格式",
    ),
    "signed": (
        "符号类型", "有无符号", "Signed", "Value Type", "Data Type", "数据类型", "Signed/Unsigned",
    ),
    "factor": (
        "精度", "因子", "比例因子", "Factor", "Scale", "Resolution", "Scaling", "分辨率",
    ),
    "offset": (
        "偏移量", "偏置", "Offset", "Bias",
    ),
    "minimum": (
        "最小值", "物理最小值", "Minimum", "Min", "Physical Min", "最小物理值",
    ),
    "maximum": (
        "最大值", "物理最大值", "Maximum", "Max", "Physical Max", "最大物理值",
    ),
    "unit": (
        "单位", "Unit", "Physical Unit",
    ),
    "initial_value": (
        "初始值", "默认值", "Initial Value", "Init Value", "Default Value", "Start Value",
    ),
    "invalid_value": (
        "无效值", "错误值", "Invalid Value", "Error Value", "Not Available Value", "NA Value",
    ),
    "signal_comment": (
        "信号描述", "信号注释", "Signal Comment", "Signal Description", "信号说明", "Description", "Comment",
    ),
    "value_table": (
        "枚举值", "值描述", "Value Table", "Value Description", "Value Definition", "取值说明", "编码定义", "Coding",
    ),
    "multiplex": (
        "复用", "复用值", "Multiplex", "Multiplexer", "Mux", "Mux Value",
    ),
}

ALIAS_NORMALIZED: Dict[str, str] = {}
for canonical, aliases in COLUMN_ALIASES.items():
    for alias in aliases:
        ALIAS_NORMALIZED[norm_header(alias)] = canonical


class MatrixReadError(RuntimeError):
    pass


class E2EIdTableError(RuntimeError):
    pass


E2E_GROUP_HEADER_ALIASES = {
    "tx": {"tx", "发送", "发送组", "txsignalgroup", "txsiggrp", "transmit", "transmitter"},
    "rx": {"rx", "接收", "接收组", "rxsignalgroup", "rxsiggrp", "receive", "receiver"},
}
E2E_DATA_ID_HEADER_ALIASES = {
    "dataid", "e2edataid", "e2eid", "dataidentifier", "数据id", "数据标识", "数据标识符"
}

E2E_ALL_SHEETS_LABEL = "全部可识别工作表"


def _e2e_header_role(value: Any) -> str:
    h = norm_header(value)
    if not h:
        return ""
    if h in E2E_DATA_ID_HEADER_ALIASES or ("data" in h and h.endswith("id")) or "e2edataid" in h:
        return "data_id"
    for role, aliases in E2E_GROUP_HEADER_ALIASES.items():
        if h in aliases:
            return role
    # 容忍“TX SignalGroup / RX SignalGroup / TX1 / RX2 / 发送SignalGroup”等表头。
    if h.startswith("tx") and "id" not in h:
        return "tx"
    if h.startswith("rx") and "id" not in h:
        return "rx"
    if "signalgroup" in h or "siggrp" in h or "group" in h:
        if h.startswith("tx") or "发送" in h:
            return "tx"
        if h.startswith("rx") or "接收" in h:
            return "rx"
    return ""


def detect_e2e_id_layout(rows: Sequence[Sequence[Any]], max_scan_rows: int = 30) -> Tuple[int, List[Tuple[str, int, int]], List[str]]:
    """识别客户E2E ID表的成对列，例如 TX|DATA ID、Rx|DATAID。"""
    best: Optional[Tuple[int, List[Tuple[str, int, int]], List[str], int]] = None
    for row_idx, row in enumerate(rows[:max_scan_rows]):
        roles = [_e2e_header_role(cell) for cell in row]
        headers = [text(cell) for cell in row]
        data_cols = [i for i, role in enumerate(roles) if role == "data_id"]
        pairs: List[Tuple[str, int, int]] = []
        for group_col, role in enumerate(roles):
            if role not in {"tx", "rx"}:
                continue
            # 优先匹配右侧最近的DATA ID列；没有时再取最近列。
            right = [c for c in data_cols if c > group_col and c - group_col <= 3]
            if right:
                id_col = min(right, key=lambda c: c - group_col)
            elif data_cols:
                id_col = min(data_cols, key=lambda c: abs(c - group_col))
            else:
                continue
            pairs.append((role.upper(), group_col, id_col))
        # 同一个ID列不应该被两个组列重复占用；按距离保留最合理的一对。
        unique: Dict[int, Tuple[str, int, int]] = {}
        for pair in pairs:
            _, gc, ic = pair
            old = unique.get(ic)
            if old is None or abs(ic - gc) < abs(old[2] - old[1]):
                unique[ic] = pair
        pairs = sorted(unique.values(), key=lambda x: x[1])
        score = len(pairs) * 10 + len(data_cols)
        if pairs and (best is None or score > best[3]):
            best = (row_idx, pairs, headers, score)
    if best is None:
        raise E2EIdTableError(
            "未识别到E2E ID表。支持的典型布局：TX | DATA ID | Rx | DATAID。\n"
            "SignalGroup列可写TX/Rx，Data ID列可写DATA ID、E2EDataID等。"
        )
    return best[0], best[1], best[2]


def _read_tabular_preview(path: str, sheet_name: Optional[str], max_rows: int = 80) -> Tuple[List[List[Any]], str]:
    suffix = Path(path).suffix.lower()
    if suffix == ".csv":
        last_error: Optional[Exception] = None
        for encoding in ("utf-8-sig", "gb18030", "utf-8"):
            try:
                with open(path, "r", encoding=encoding, newline="") as f:
                    rows = []
                    for i, row in enumerate(csv.reader(f)):
                        rows.append(row)
                        if i + 1 >= max_rows:
                            break
                return rows, "CSV"
            except UnicodeDecodeError as exc:
                last_error = exc
        raise E2EIdTableError(f"CSV编码无法识别：{last_error}")
    if suffix not in {".xlsx", ".xlsm"}:
        raise E2EIdTableError("E2E ID表目前支持 .xlsx、.xlsm、.csv。")
    if load_workbook is None:
        raise E2EIdTableError("读取Excel E2E ID表需要 openpyxl。请执行：pip install openpyxl")
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet_name and sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
        else:
            ws = wb[wb.sheetnames[0]]
        rows = [list(r) for r in itertools.islice(ws.iter_rows(values_only=True), max_rows)]
        return rows, ws.title
    finally:
        wb.close()


def inspect_e2e_id_sheets(path: str) -> Tuple[List[str], str, Dict[str, str]]:
    """返回E2E表工作表列表、最佳工作表、每页识别说明。"""
    suffix = Path(path).suffix.lower()
    if suffix == ".csv":
        rows, _ = _read_tabular_preview(path, "CSV")
        try:
            _h, pairs, _headers = detect_e2e_id_layout(rows)
            return ["CSV"], "CSV", {"CSV": f"识别到 {len(pairs)} 组 SignalGroup/Data ID 列"}
        except Exception as exc:
            return ["CSV"], "", {"CSV": str(exc)}
    if suffix not in {".xlsx", ".xlsm"}:
        raise E2EIdTableError("E2E ID表目前支持 .xlsx、.xlsm、.csv。")
    if load_workbook is None:
        raise E2EIdTableError("读取Excel E2E ID表需要 openpyxl。请执行：pip install openpyxl")
    wb = load_workbook(path, read_only=True, data_only=True)
    details: Dict[str, str] = {}
    best_sheet = ""
    best_score = -1
    try:
        names = list(wb.sheetnames)
        for name in names:
            ws = wb[name]
            rows = [list(r) for r in itertools.islice(ws.iter_rows(values_only=True), 60)]
            try:
                header_idx, pairs, _headers = detect_e2e_id_layout(rows)
                nonempty = 0
                for row in rows[header_idx + 1:]:
                    for _direction, gc, ic in pairs:
                        if gc < len(row) and text(row[gc]):
                            nonempty += 1
                score = len(pairs) * 100 + nonempty
                details[name] = f"识别到 {len(pairs)} 组列，预览有效映射约 {nonempty} 条"
                if score > best_score:
                    best_score = score
                    best_sheet = name
            except Exception as exc:
                details[name] = str(exc).splitlines()[0]
        return names, best_sheet, details
    finally:
        wb.close()


def _extract_e2e_entries_from_rows(
    rows: Sequence[Sequence[Any]], actual_sheet: str
) -> Tuple[List[E2EIdEntry], Dict[str, str], int]:
    """从单个工作表读取所有可识别的E2E列组。

    同一工作表可以同时存在 TX1/DATA ID、RX1/DATA ID、TX2/DATA ID 等多路列组。
    SignalGroup存在但Data ID为空或'-'时保留为 data_id=None，仅做提示，不判DBC错误。
    """
    if not rows:
        raise E2EIdTableError(f"工作表 {actual_sheet} 为空。")
    header_idx, pairs, headers = detect_e2e_id_layout(rows)
    entries: List[E2EIdEntry] = []
    mapping: Dict[str, str] = {}
    direction_count: Dict[str, int] = {}
    for direction, group_col, id_col in pairs:
        direction_count[direction] = direction_count.get(direction, 0) + 1
        route_no = direction_count[direction]
        route_label = f"{direction}{route_no}" if sum(1 for d, _g, _i in pairs if d == direction) > 1 else direction
        group_header = headers[group_col] if group_col < len(headers) else str(group_col + 1)
        id_header = headers[id_col] if id_col < len(headers) else str(id_col + 1)
        mapping[f"{actual_sheet}/{route_label} SignalGroup"] = group_header
        mapping[f"{actual_sheet}/{route_label} Data ID"] = id_header

    direction_seen: Dict[str, int] = {}
    route_labels: List[Tuple[str, int, int]] = []
    for direction, group_col, id_col in pairs:
        direction_seen[direction] = direction_seen.get(direction, 0) + 1
        route_no = direction_seen[direction]
        total_same = sum(1 for d, _g, _i in pairs if d == direction)
        route_label = f"{direction}{route_no}" if total_same > 1 else direction
        route_labels.append((route_label, group_col, id_col))

    for row_no, row in enumerate(rows[header_idx + 1:], start=header_idx + 2):
        for route_label, group_col, id_col in route_labels:
            group = text(row[group_col]) if group_col < len(row) else ""
            if not group:
                continue
            raw_id = text(row[id_col]) if id_col < len(row) else ""
            if raw_id in {"-", "--", "N/A", "NA", "n/a", "na"} or not raw_id:
                data_id = None
            else:
                data_id = parse_can_id(raw_id)
                if data_id is None:
                    raise E2EIdTableError(
                        f"工作表 {actual_sheet} 第{row_no}行：{route_label} SignalGroup={group} 的Data ID“{raw_id}”无法解析。"
                    )
            entries.append(E2EIdEntry(route_label, group, data_id, row_no, actual_sheet))
    return entries, mapping, header_idx + 1


def read_e2e_id_table(path: str, sheet_name: Optional[str] = None) -> Tuple[List[E2EIdEntry], Dict[str, str], int, str]:
    """读取客户E2E ID表。

    支持：
    1) 单工作表内多路 TX/RX + DATA ID 列组；
    2) Excel中选择“全部可识别工作表”后汇总所有能识别的工作表；
    3) 客户表只提供部分SignalGroup或部分Data ID，缺项不自动判为DBC错误。
    """
    suffix = Path(path).suffix.lower()
    if suffix == ".csv":
        last_error: Optional[Exception] = None
        rows: List[List[Any]] = []
        for encoding in ("utf-8-sig", "gb18030", "utf-8"):
            try:
                with open(path, "r", encoding=encoding, newline="") as f:
                    rows = [list(r) for r in csv.reader(f)]
                last_error = None
                break
            except UnicodeDecodeError as exc:
                last_error = exc
        if last_error is not None:
            raise E2EIdTableError(f"CSV编码无法识别：{last_error}")
        entries, mapping, header = _extract_e2e_entries_from_rows(rows, "CSV")
        if not entries:
            raise E2EIdTableError("E2E ID表没有读取到任何SignalGroup。")
        return entries, mapping, header, "CSV"

    if suffix not in {".xlsx", ".xlsm"}:
        raise E2EIdTableError("E2E ID表目前支持 .xlsx、.xlsm、.csv。")
    if load_workbook is None:
        raise E2EIdTableError("读取Excel E2E ID表需要 openpyxl。请执行：pip install openpyxl")

    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        selected_all = (sheet_name == E2E_ALL_SHEETS_LABEL)
        target_names: List[str]
        if selected_all:
            target_names = list(wb.sheetnames)
        elif sheet_name and sheet_name in wb.sheetnames:
            target_names = [sheet_name]
        else:
            _names, best, _details = inspect_e2e_id_sheets(path)
            target_names = [best or wb.sheetnames[0]]

        all_entries: List[E2EIdEntry] = []
        all_mapping: Dict[str, str] = {}
        headers: List[int] = []
        recognized_sheets: List[str] = []
        skipped: List[str] = []
        for name in target_names:
            ws = wb[name]
            rows = [list(r) for r in ws.iter_rows(values_only=True)]
            try:
                entries, mapping, header = _extract_e2e_entries_from_rows(rows, name)
            except E2EIdTableError as exc:
                if selected_all:
                    skipped.append(f"{name}：{str(exc).splitlines()[0]}")
                    continue
                raise
            if entries:
                all_entries.extend(entries)
                all_mapping.update(mapping)
                headers.append(header)
                recognized_sheets.append(name)

        if not all_entries:
            detail = "\n".join(skipped[:10])
            raise E2EIdTableError("E2E ID表没有读取到任何SignalGroup。" + (f"\n{detail}" if detail else ""))

        if selected_all:
            actual_sheet = "、".join(recognized_sheets)
            if len(recognized_sheets) > 1:
                actual_sheet += f"（共{len(recognized_sheets)}个工作表）"
            # 多工作表的表头行可能不同，0表示“各表自动识别”。
            header_display = headers[0] if len(set(headers)) == 1 else 0
        else:
            actual_sheet = recognized_sheets[0]
            header_display = headers[0]
        return all_entries, all_mapping, header_display, actual_sheet
    finally:
        wb.close()


def _e2e_data_id_candidates_for_group(db: Database, msg: Message, group_name: str) -> List[Tuple[str, int]]:
    """收集DBC中可归属于某SignalGroup的E2EDataId候选。"""
    candidates: List[Tuple[str, int]] = []
    # Message级E2EDataId仅在报文只有一个SignalGroup时可唯一归属；多个组时不贸然绑定。
    if len(msg.signal_groups) == 1:
        value, explicit = effective_attribute(msg, db, "E2EDataId")
        number = parse_int(value)
        if number is not None and (explicit or value not in (None, "")):
            candidates.append(("Message.E2EDataId", number))
    member_names = set(msg.signal_groups.get(group_name, ()))
    for sig in msg.signals:
        if sig.name not in member_names:
            continue
        value, explicit = effective_attribute(sig, db, "E2EDataId")
        number = parse_int(value)
        if number is not None and explicit:
            candidates.append((f"Signal.{sig.name}.E2EDataId", number))
    return candidates


def compare_e2e_id_table(entries: Sequence[E2EIdEntry], db: Database) -> List[Difference]:
    """将客户E2E SignalGroup/Data ID表与DBC中的E2EDataId进行核对。"""
    diffs: List[Difference] = []
    group_index_exact: Dict[str, List[Tuple[Message, str]]] = {}
    group_index_norm: Dict[str, List[Tuple[Message, str]]] = {}
    for msg in db.messages:
        for group_name in msg.signal_groups:
            group_index_exact.setdefault(group_name, []).append((msg, group_name))
            group_index_norm.setdefault(norm_name(group_name), []).append((msg, group_name))

    # 客户表自身重复/冲突检查。
    table_values: Dict[str, List[E2EIdEntry]] = {}
    for entry in entries:
        table_values.setdefault(norm_name(entry.group_name), []).append(entry)
    for key, same_group in table_values.items():
        ids = {e.data_id for e in same_group if e.data_id is not None}
        if len(ids) > 1:
            first = same_group[0]
            diffs.append(Difference(
                "错误", "E2E ID表核对", None, "", first.group_name, "客户E2E ID表重复冲突",
                " / ".join(format_can_id(v) for v in sorted(ids)), "",
                f"客户E2E ID表中SignalGroup“{first.group_name}”出现多个不同Data ID，请先确认客户表。"
                f"来源：{'；'.join(f'{e.sheet_name}第{e.row_number}行={format_can_id(e.data_id)}' for e in same_group if e.data_id is not None)}",
                "E2E_TABLE_DUP_001",
            ))

    processed: set[Tuple[str, Optional[int]]] = set()
    for entry in entries:
        marker = (norm_name(entry.group_name), entry.data_id)
        if marker in processed:
            continue
        processed.add(marker)
        if entry.data_id is None:
            # '-'表示客户没有给ID，不作为DBC错误，只做提示。
            diffs.append(Difference(
                "提示", "E2E ID表核对", None, "", entry.group_name, "E2E Data ID",
                "-", "",
                f"客户E2E表 {entry.sheet_name} 第{entry.row_number}行未给SignalGroup“{entry.group_name}”配置Data ID，本项跳过数值核对，不判DBC错误。",
                "E2E_TABLE_NO_ID_001",
            ))
            continue
        matches = group_index_exact.get(entry.group_name, [])
        source_note = f"来源：{entry.sheet_name} 第{entry.row_number}行。" if entry.sheet_name else ""
        name_note = ""
        if not matches:
            norm_matches = group_index_norm.get(norm_name(entry.group_name), [])
            if len(norm_matches) == 1:
                matches = norm_matches
                name_note = f"客户表名称为“{entry.group_name}”，DBC中为“{norm_matches[0][1]}”，名称大小写/分隔符不完全一致。"
            elif len(norm_matches) > 1:
                matches = norm_matches
        if not matches:
            diffs.append(Difference(
                "错误", "E2E ID表核对", None, "", entry.group_name, "SignalGroup",
                f"{entry.direction} / {format_can_id(entry.data_id)} ({entry.data_id})", "未找到",
                f"客户E2E ID表中的SignalGroup“{entry.group_name}”在DBC中未找到。",
                "E2E_TABLE_GROUP_001",
            ))
            continue
        if len(matches) > 1:
            names = "; ".join(f"{format_can_id(m.can_id)} {m.name}/{g}" for m, g in matches[:8])
            diffs.append(Difference(
                "错误", "E2E ID表核对", None, "", entry.group_name, "SignalGroup定位",
                f"{format_can_id(entry.data_id)} ({entry.data_id})", names,
                f"SignalGroup“{entry.group_name}”在DBC中匹配到多个位置，无法唯一核对Data ID。",
                "E2E_TABLE_GROUP_DUP_001",
            ))
            continue
        msg, actual_group = matches[0]
        if name_note:
            diffs.append(Difference(
                "警告", "E2E ID表核对", msg.can_id, msg.name, actual_group, "SignalGroup名称",
                entry.group_name, actual_group, name_note,
                "E2E_TABLE_NAME_001",
            ))
        candidates = _e2e_data_id_candidates_for_group(db, msg, actual_group)
        unique_ids = sorted({v for _src, v in candidates})
        customer_display = f"{format_can_id(entry.data_id)} ({entry.data_id})"
        if not candidates:
            extra = ""
            if len(msg.signal_groups) > 1:
                msg_level, explicit = effective_attribute(msg, db, "E2EDataId")
                if explicit and parse_int(msg_level) is not None:
                    extra = " 当前报文包含多个SignalGroup且只发现Message级E2EDataId，无法安全判断该值属于哪一个Group。"
            diffs.append(Difference(
                "错误", "E2E ID表核对", msg.can_id, msg.name, actual_group, "E2EDataId",
                customer_display, "未配置/无法唯一归属",
                f"客户要求 {actual_group} 的E2E Data ID={customer_display}，但DBC中未找到可唯一归属于该SignalGroup的E2EDataId。"
                f"建议按客户表补充/确认E2EDataId。{extra}",
                "E2E_TABLE_DATAID_MISSING_001",
            ))
            continue
        if len(unique_ids) > 1:
            dbc_display = "；".join(f"{src}={format_can_id(v)} ({v})" for src, v in candidates)
            diffs.append(Difference(
                "错误", "E2E ID表核对", msg.can_id, msg.name, actual_group, "E2EDataId冲突",
                customer_display, dbc_display,
                f"同一SignalGroup关联到多个不同的DBC E2EDataId，需先消除DBC内部冲突。",
                "E2E_TABLE_DATAID_MULTI_001",
            ))
            continue
        dbc_id = unique_ids[0]
        dbc_display = f"{format_can_id(dbc_id)} ({dbc_id})"
        if dbc_id != entry.data_id:
            sources = ", ".join(src for src, _v in candidates)
            diffs.append(Difference(
                "错误", "E2E ID表核对", msg.can_id, msg.name, actual_group, "E2EDataId",
                customer_display, dbc_display,
                f"客户E2E ID表与DBC不一致。建议值：{customer_display}；DBC来源：{sources}。",
                "E2E_TABLE_DATAID_MISMATCH_001",
            ))
    return diffs


def build_e2e_mapping_rows(entries: Sequence[E2EIdEntry], db: Database) -> List[E2EComparisonRow]:
    """把客户E2E表与DBC做成逐SignalGroup可视化核对结果。

    该函数只负责展示，不改变 compare_e2e_id_table 的错误判定。
    """
    group_index_exact: Dict[str, List[Tuple[Message, str]]] = {}
    group_index_norm: Dict[str, List[Tuple[Message, str]]] = {}
    for msg in db.messages:
        for group_name in msg.signal_groups:
            group_index_exact.setdefault(group_name, []).append((msg, group_name))
            group_index_norm.setdefault(norm_name(group_name), []).append((msg, group_name))

    # 标记客户表自身同名SignalGroup出现不同Data ID的情况。
    conflicting_customer_groups: set[str] = set()
    grouped_entries: Dict[str, List[E2EIdEntry]] = {}
    for entry in entries:
        grouped_entries.setdefault(norm_name(entry.group_name), []).append(entry)
    for key, same_group in grouped_entries.items():
        ids = {e.data_id for e in same_group if e.data_id is not None}
        if len(ids) > 1:
            conflicting_customer_groups.add(key)

    rows: List[E2EComparisonRow] = []
    seen: set[Tuple[str, str, Optional[int]]] = set()
    for entry in entries:
        marker = (entry.direction, norm_name(entry.group_name), entry.data_id)
        if marker in seen:
            continue
        seen.add(marker)

        matches = group_index_exact.get(entry.group_name, [])
        source_note = f"来源：{entry.sheet_name} 第{entry.row_number}行。" if entry.sheet_name else ""
        name_note = ""
        if not matches:
            norm_matches = group_index_norm.get(norm_name(entry.group_name), [])
            if len(norm_matches) == 1:
                matches = norm_matches
                name_note = f"客户表名称“{entry.group_name}”与DBC名称“{norm_matches[0][1]}”大小写/分隔符不同。"
            elif len(norm_matches) > 1:
                matches = norm_matches

        if norm_name(entry.group_name) in conflicting_customer_groups:
            rows.append(E2EComparisonRow(
                entry.direction, entry.group_name, "", None, "", entry.data_id, tuple(),
                "客户表冲突", source_note + " 同一SignalGroup在客户表中配置了多个不同Data ID。"
            ))
            continue

        if not matches:
            rows.append(E2EComparisonRow(
                entry.direction, entry.group_name, "", None, "", entry.data_id, tuple(),
                "DBC未找到", source_note + " 客户表中的SignalGroup在DBC中未找到。"
            ))
            continue

        if len(matches) > 1:
            locs = "; ".join(f"{format_can_id(m.can_id)} {m.name}/{g}" for m, g in matches[:6])
            rows.append(E2EComparisonRow(
                entry.direction, entry.group_name, "多个位置", None, locs, entry.data_id, tuple(),
                "DBC多处匹配", source_note + " 无法唯一定位该SignalGroup。"
            ))
            continue

        msg, actual_group = matches[0]
        candidates = _e2e_data_id_candidates_for_group(db, msg, actual_group)
        unique_ids = tuple(sorted({v for _src, v in candidates}))

        if entry.data_id is None:
            status = "客户未给ID"
            note = source_note + " 客户表该项为'-'或空白，仅展示DBC值，不做数值一致性判断。"
        elif not candidates:
            status = "DBC未配置"
            note = source_note + " DBC中未找到可唯一归属于该SignalGroup的E2EDataId。"
        elif len(unique_ids) > 1:
            status = "DBC多值冲突"
            note = source_note + " 同一SignalGroup关联到多个不同的DBC E2EDataId。"
        elif unique_ids[0] == entry.data_id:
            status = "一致（名称差异）" if name_note else "一致"
            note = source_note + (" " + name_note if name_note else "")
        else:
            status = "不一致"
            sources = ", ".join(src for src, _v in candidates)
            note = source_note + f" 客户Data ID与DBC不一致；DBC来源：{sources}。"
            if name_note:
                note = source_note + " " + name_note + f" 客户Data ID与DBC不一致；DBC来源：{sources}。"

        rows.append(E2EComparisonRow(
            entry.direction, entry.group_name, actual_group, msg.can_id, msg.name,
            entry.data_id, unique_ids, status, note,
        ))
    return rows


def detect_header_and_mapping(rows: Sequence[Sequence[Any]], max_scan_rows: int = 50) -> Tuple[int, Dict[str, int], List[str]]:
    """返回 0 基头行、字段到列号映射、原始表头。"""
    best_row = -1
    best_mapping: Dict[str, int] = {}
    best_headers: List[str] = []
    best_score = -1

    for row_idx, row in enumerate(rows[:max_scan_rows]):
        mapping: Dict[str, int] = {}
        headers = [text(cell) for cell in row]
        for col_idx, cell in enumerate(row):
            canonical = ALIAS_NORMALIZED.get(norm_header(cell))
            if canonical and canonical not in mapping:
                mapping[canonical] = col_idx

        # 关键列给予更高权重。
        score = len(mapping)
        score += 3 if "signal_name" in mapping else 0
        score += 2 if "message_id" in mapping else 0
        score += 2 if "message_name" in mapping else 0
        score += 1 if "start_bit" in mapping else 0
        score += 1 if "signal_length" in mapping else 0

        if score > best_score:
            best_score = score
            best_row = row_idx
            best_mapping = mapping
            best_headers = headers

    required_msg = "message_id" in best_mapping or "message_name" in best_mapping
    if best_row < 0 or not required_msg or "signal_name" not in best_mapping:
        raise MatrixReadError(
            "未能识别通信矩阵表头。至少需要：报文ID或报文名称、信号名称。\n"
            "可将列名改为常见名称，例如：报文ID、报文名称、信号名称、起始位、信号长度。"
        )
    return best_row, best_mapping, best_headers


def list_matrix_sheets(path: str) -> List[str]:
    suffix = Path(path).suffix.lower()
    if suffix == ".csv":
        return ["CSV"]
    if suffix not in {".xlsx", ".xlsm"}:
        raise MatrixReadError("目前支持 .xlsx、.xlsm 和 .csv，不支持旧版 .xls。")
    if load_workbook is None:
        raise MatrixReadError("读取 Excel 需要 openpyxl。请执行：pip install openpyxl")
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        return list(wb.sheetnames)
    finally:
        wb.close()


def read_csv_rows(path: str) -> List[List[str]]:
    last_error: Optional[Exception] = None
    for encoding in ("utf-8-sig", "gb18030", "gbk"):
        try:
            with open(path, "r", encoding=encoding, newline="") as f:
                return [list(row) for row in csv.reader(f)]
        except UnicodeDecodeError as exc:
            last_error = exc
    raise MatrixReadError(f"CSV 编码无法识别：{last_error}")


def parse_value_table(value: Any) -> Dict[int, str]:
    """解析十进制/十六进制枚举及范围，例如 0=Off、0x1:On、0x2-0x7:Reserve。"""
    s = text(value)
    if not s:
        return {}

    def parse_token(token: str) -> Optional[int]:
        token = token.strip()
        sign = -1 if token.startswith("-") else 1
        unsigned = token[1:] if token[:1] in "+-" else token
        try:
            if unsigned.lower().startswith("0x"):
                return sign * int(unsigned[2:], 16)
            return int(token, 10)
        except ValueError:
            return None

    result: Dict[int, str] = {}
    item_re = re.compile(
        r"^\s*([+-]?(?:0[xX][0-9a-fA-F]+|\d+))"
        r"(?:\s*-\s*([+-]?(?:0[xX][0-9a-fA-F]+|\d+)))?"
        r"\s*[:=：]\s*[\"']?(.+?)[\"']?\s*$"
    )
    for item in re.split(r"[;；\n]+", s):
        item = item.strip()
        if not item:
            continue
        match = item_re.match(item)
        if not match:
            continue
        start = parse_token(match.group(1))
        end = parse_token(match.group(2)) if match.group(2) else start
        if start is None or end is None:
            continue
        description = match.group(3).strip()
        step = 1 if end >= start else -1
        # 防止异常范围导致大量展开。
        if abs(end - start) > 4096:
            continue
        for raw_value in range(start, end + step, step):
            result[raw_value] = description
    return result


def matrix_rows_to_database(
    rows: Iterable[Sequence[Any]],
    source_name: str,
    progress: Optional[Callable[[str], None]] = None,
) -> Tuple[Database, Dict[str, str], int]:
    """流式读取矩阵行，避免把大型 Excel 工作表一次性装入内存。"""
    row_iter = iter(rows)
    preview: List[List[Any]] = []
    for _ in range(50):
        try:
            preview.append(list(next(row_iter)))
        except StopIteration:
            break

    if not preview:
        raise MatrixReadError("通信矩阵工作表为空。")

    header_idx, mapping, headers = detect_header_and_mapping(preview)
    mapping_names = {
        key: headers[col] if col < len(headers) else f"第{col + 1}列"
        for key, col in mapping.items()
    }

    # 该 Vector 通信矩阵模板的“Start Bit(LSB Or MSB)”实际填写的是信号 LSB 位号。
    # DBC 对 Motorola 信号保存的是 MSB 位号，不能直接比较数字，需转换为占用位集合。
    start_header_norm = norm_header(mapping_names.get("start_bit", ""))
    matrix_start_bit_convention = (
        "lsb" if start_header_norm in {
            "startbitlsbormsb", "startbitlsb", "lsbstartbit", "起始位lsb",
        } else "dbc"
    )

    if progress:
        progress(f"已识别表头：第 {header_idx + 1} 行，共识别 {len(mapping)} 个字段")

    messages: List[Message] = []
    message_index: Dict[Tuple[Optional[int], str], Message] = {}

    # 合并单元格或按报文分组的矩阵通常只在首行填写报文字段，需要向下继承。
    last_msg_values: Dict[str, Any] = {
        "message_id": None,
        "message_name": "",
        "dlc": None,
        "sender": "",
        "cycle_time": None,
        "send_type": "",
        "frame_format": None,
        "bus_format": None,
        "message_comment": "",
    }

    def get_cell(row: Sequence[Any], field_name: str) -> Any:
        idx = mapping.get(field_name)
        if idx is None or idx >= len(row):
            return None
        return row[idx]

    data_rows = itertools.chain(preview[header_idx + 1 :], row_iter)
    processed = 0
    for data_idx, row in enumerate(data_rows, start=header_idx + 2):
        processed += 1
        if progress and processed % 500 == 0:
            progress(f"正在读取通信矩阵：已处理 {processed} 行……")

        if not any(text(cell) for cell in row):
            continue

        raw_msg_id = get_cell(row, "message_id")
        raw_msg_name = get_cell(row, "message_name")

        if text(raw_msg_id):
            last_msg_values["message_id"] = parse_can_id(raw_msg_id)
        if text(raw_msg_name):
            last_msg_values["message_name"] = text(raw_msg_name)

        for field_name, parser in (
            ("dlc", parse_int),
            ("cycle_time", parse_float),
        ):
            raw = get_cell(row, field_name)
            if text(raw):
                last_msg_values[field_name] = parser(raw)

        raw_frame_format = get_cell(row, "frame_format")
        if text(raw_frame_format):
            last_msg_values["frame_format"] = parse_frame_format(raw_frame_format)
            last_msg_values["bus_format"] = parse_bus_format(raw_frame_format)

        for field_name in ("sender", "send_type", "message_comment"):
            raw = get_cell(row, field_name)
            if text(raw):
                last_msg_values[field_name] = text(raw)

        can_id = last_msg_values["message_id"]
        msg_name = last_msg_values["message_name"]
        if can_id is None and not msg_name:
            continue

        key = (can_id, norm_name(msg_name))
        message = message_index.get(key)
        if message is None:
            message = Message(
                can_id=can_id,
                name=msg_name,
                dlc=last_msg_values["dlc"],
                sender=last_msg_values["sender"],
                cycle_time=last_msg_values["cycle_time"],
                send_type=last_msg_values["send_type"],
                frame_format=last_msg_values["frame_format"] or (
                    "standard" if can_id is not None and can_id <= 0x7FF else "extended"
                ),
                bus_format=last_msg_values["bus_format"],
                id_frame_format=("standard" if can_id is not None and can_id <= 0x7FF else "extended"),
                comment=last_msg_values["message_comment"],
                row_number=data_idx,
            )
            message_index[key] = message
            messages.append(message)
        else:
            # 后续行补充首行缺失的报文字段。
            for attr, value in (
                ("dlc", last_msg_values["dlc"]),
                ("sender", last_msg_values["sender"]),
                ("cycle_time", last_msg_values["cycle_time"]),
                ("send_type", last_msg_values["send_type"]),
                ("frame_format", last_msg_values["frame_format"]),
                ("bus_format", last_msg_values["bus_format"]),
                ("comment", last_msg_values["message_comment"]),
            ):
                current = getattr(message, attr)
                if (current is None or current == "") and value not in (None, ""):
                    setattr(message, attr, value)

        signal_name = text(get_cell(row, "signal_name"))
        if not signal_name:
            continue

        start_bit = parse_int(get_cell(row, "start_bit"))
        if start_bit is None:
            start_byte = parse_int(get_cell(row, "start_byte"))
            bit_in_byte = parse_int(get_cell(row, "bit_in_byte"))
            if start_byte is not None and bit_in_byte is not None:
                # 默认起始字节从 0 计数。若矩阵从 1 计数，可直接提供“起始位”列以避免歧义。
                start_bit = start_byte * 8 + bit_in_byte

        signal = Signal(
            name=signal_name,
            start_bit=start_bit,
            start_bit_convention=matrix_start_bit_convention,
            length=parse_int(get_cell(row, "signal_length")),
            byte_order=parse_byte_order(get_cell(row, "byte_order")),
            signed=parse_bool_signed(get_cell(row, "signed")),
            factor=parse_float(get_cell(row, "factor")),
            offset=parse_float(get_cell(row, "offset")),
            minimum=parse_float(get_cell(row, "minimum")),
            maximum=parse_float(get_cell(row, "maximum")),
            unit=text(get_cell(row, "unit")),
            receivers=split_nodes(get_cell(row, "receiver")),
            initial_value=parse_float(get_cell(row, "initial_value")),
            invalid_value=parse_float(get_cell(row, "invalid_value")),
            comment=text(get_cell(row, "signal_comment")),
            value_table=parse_value_table(get_cell(row, "value_table")),
            multiplex=text(get_cell(row, "multiplex")),
            row_number=data_idx,
        )
        message.signals.append(signal)

    if not messages:
        raise MatrixReadError("通信矩阵中没有读取到有效报文。")
    if progress:
        progress(f"通信矩阵读取完成：{processed} 行，{len(messages)} 条报文。")
    return Database(source=source_name, messages=messages), mapping_names, header_idx + 1


def _mapping_is_matrix_definition(mapping: Dict[str, int]) -> bool:
    """区别原始通信矩阵和本工具导出的差异报告。"""
    has_position = "start_bit" in mapping or ("start_byte" in mapping and "bit_in_byte" in mapping)
    return has_position and "signal_length" in mapping


def _preview_excel_sheet(ws: Any, max_rows: int = 50) -> Tuple[bool, int, str]:
    """仅扫描前若干行，判断工作表是否像通信矩阵。"""
    rows = [list(row) for row in itertools.islice(ws.iter_rows(values_only=True), max_rows)]
    try:
        _header_idx, mapping, _headers = detect_header_and_mapping(rows, max_scan_rows=max_rows)
        if not _mapping_is_matrix_definition(mapping):
            return False, 0, "缺少信号起始位或信号长度列，不像原始通信矩阵"
        score = len(mapping)
        return True, score, ""
    except MatrixReadError as exc:
        return False, 0, str(exc)


def inspect_matrix_sheets(path: str) -> Tuple[List[str], Optional[str], Dict[str, str]]:
    """返回工作表列表、最可能的矩阵工作表和各表检测结果。"""
    suffix = Path(path).suffix.lower()
    if suffix == ".csv":
        rows = read_csv_rows(path)[:50]
        try:
            _idx, mapping, _headers = detect_header_and_mapping(rows)
            if not _mapping_is_matrix_definition(mapping):
                return ["CSV"], None, {"CSV": "缺少信号起始位或信号长度列，不像原始通信矩阵"}
            return ["CSV"], "CSV", {"CSV": "可识别为通信矩阵"}
        except MatrixReadError as exc:
            return ["CSV"], None, {"CSV": str(exc)}

    if suffix not in {".xlsx", ".xlsm"}:
        raise MatrixReadError("目前支持 .xlsx、.xlsm 和 .csv，不支持旧版 .xls。")
    if load_workbook is None:
        raise MatrixReadError("读取 Excel 需要 openpyxl。请执行：pip install openpyxl")

    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        names = list(wb.sheetnames)
        best_sheet: Optional[str] = None
        best_score = -1
        details: Dict[str, str] = {}
        for name in names:
            valid, score, reason = _preview_excel_sheet(wb[name])
            if valid:
                details[name] = f"可识别为通信矩阵（识别字段 {score} 个）"
                if score > best_score:
                    best_score = score
                    best_sheet = name
            else:
                details[name] = reason
        return names, best_sheet, details
    finally:
        wb.close()


def validate_matrix_selection(path: str, sheet_name: Optional[str]) -> None:
    """开始检查前快速验证输入，避免把导出的检查报告误当成通信矩阵。"""
    suffix = Path(path).suffix.lower()
    if suffix == ".csv":
        rows = read_csv_rows(path)[:50]
        _idx, mapping, _headers = detect_header_and_mapping(rows)
        if not _mapping_is_matrix_definition(mapping):
            raise MatrixReadError(
                "所选 CSV 缺少信号起始位或信号长度列，不像原始 CAN 通信矩阵。"
            )
        return

    if load_workbook is None:
        raise MatrixReadError("读取 Excel 需要 openpyxl。请执行：pip install openpyxl")
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        if not sheet_name or sheet_name not in wb.sheetnames:
            raise MatrixReadError("请选择有效的通信矩阵工作表。")
        rows = [list(row) for row in itertools.islice(wb[sheet_name].iter_rows(values_only=True), 50)]
        try:
            _idx, mapping, _headers = detect_header_and_mapping(rows)
            if not _mapping_is_matrix_definition(mapping):
                raise MatrixReadError("缺少信号起始位或信号长度列")
        except MatrixReadError as exc:
            raise MatrixReadError(
                f"工作表“{sheet_name}”不是可识别的通信矩阵。\n"
                "当前文件很可能是检查结果报告，而不是原始 CAN 通信矩阵。\n\n"
                "请重新选择包含报文ID/报文名称、信号名称、起始位、长度等字段的原始矩阵工作表。"
            ) from exc
    finally:
        wb.close()


def read_matrix(
    path: str,
    sheet_name: Optional[str] = None,
    progress: Optional[Callable[[str], None]] = None,
) -> Tuple[Database, Dict[str, str], int]:
    suffix = Path(path).suffix.lower()
    if suffix == ".csv":
        with open(path, "r", encoding="utf-8-sig", newline="") as f:
            try:
                return matrix_rows_to_database(csv.reader(f), f"{Path(path).name} / CSV", progress)
            except UnicodeDecodeError:
                pass
        # 非 UTF-8 CSV 回退到原来的多编码读取逻辑。
        rows = read_csv_rows(path)
        return matrix_rows_to_database(rows, f"{Path(path).name} / CSV", progress)

    if suffix not in {".xlsx", ".xlsm"}:
        raise MatrixReadError("目前支持 .xlsx、.xlsm 和 .csv，不支持旧版 .xls。")
    if load_workbook is None:
        raise MatrixReadError("读取 Excel 需要 openpyxl。请执行：pip install openpyxl")

    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        if not sheet_name:
            sheet_name = wb.sheetnames[0]
        if sheet_name not in wb.sheetnames:
            raise MatrixReadError(f"Excel 中不存在工作表：{sheet_name}")
        ws = wb[sheet_name]
        rows = (list(row) for row in ws.iter_rows(values_only=True))
        return matrix_rows_to_database(rows, f"{Path(path).name} / {sheet_name}", progress)
    finally:
        wb.close()


class DbcParseError(RuntimeError):
    pass


def parse_dbc(path: str, progress: Optional[Callable[[str], None]] = None) -> Database:
    try:
        raw = Path(path).read_bytes()
    except OSError as exc:
        raise DbcParseError(f"无法读取 DBC：{exc}") from exc

    content = None
    for encoding in ("utf-8-sig", "gb18030", "latin-1"):
        try:
            content = raw.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    if content is None:
        raise DbcParseError("DBC 文本编码无法识别。")

    messages: List[Message] = []
    nodes: List[str] = []
    current_message: Optional[Message] = None
    message_by_raw_id: Dict[int, List[Message]] = {}
    signal_by_key: Dict[Tuple[int, str], Signal] = {}
    syntax_issues: List[Tuple[int, str, str]] = []

    bo_re = re.compile(r"^\s*BO_\s+(\d+)\s+([^:]+)\s*:\s*(\d+)\s+(\S+)")
    sg_re = re.compile(
        r"^\s*SG_\s+(\S+)"
        r"(?:\s+([mM]\d+|[mM]))?\s*:\s*"
        r"(\d+)\|(\d+)@([01])([+-])\s*"
        r"\(\s*([^,]+)\s*,\s*([^\)]+)\s*\)\s*"
        r"\[\s*([^|]+)\|([^\]]+)\]\s*"
        r'"([^"]*)"\s*(.*)$'
    )
    bu_re = re.compile(r"^\s*BU_\s*:\s*(.*)$")

    lines = content.splitlines()
    for line_no, line in enumerate(lines, start=1):
        bu_match = bu_re.match(line)
        if bu_match:
            nodes.extend([item for item in re.split(r"\s+", bu_match.group(1).strip()) if item])
            continue

        bo_match = bo_re.match(line)
        if bo_match:
            raw_id = int(bo_match.group(1))
            normalized_id = raw_id & 0x1FFFFFFF if raw_id & 0x80000000 else raw_id
            id_frame_format = "extended" if (raw_id & 0x80000000 or normalized_id > 0x7FF) else "standard"
            current_message = Message(
                can_id=normalized_id,
                name=bo_match.group(2).strip(),
                dlc=int(bo_match.group(3)),
                sender=bo_match.group(4).strip(),
                frame_format=id_frame_format,
                id_frame_format=id_frame_format,
                row_number=line_no,
            )
            messages.append(current_message)
            message_by_raw_id.setdefault(raw_id, []).append(current_message)
            message_by_raw_id.setdefault(normalized_id, []).append(current_message)
            continue
        if re.match(r"^\s*BO_\s", line):
            syntax_issues.append((line_no, "BO_", line.strip()))
            current_message = None
            continue

        sg_match = sg_re.match(line)
        if sg_match and current_message is not None:
            receivers = split_nodes(sg_match.group(12))
            signal = Signal(
                name=sg_match.group(1),
                multiplex=sg_match.group(2) or "",
                start_bit=int(sg_match.group(3)),
                length=int(sg_match.group(4)),
                byte_order="intel" if sg_match.group(5) == "1" else "motorola",
                signed=sg_match.group(6) == "-",
                factor=parse_float(sg_match.group(7)),
                offset=parse_float(sg_match.group(8)),
                minimum=parse_float(sg_match.group(9)),
                maximum=parse_float(sg_match.group(10)),
                unit=sg_match.group(11),
                receivers=receivers,
                row_number=line_no,
            )
            current_message.signals.append(signal)
            if current_message.can_id is not None:
                signal_by_key[(current_message.can_id, signal.name)] = signal
            continue
        if re.match(r"^\s*SG_\s", line) and not re.match(r"^\s*SG_(MUL_VAL|TYPE)_", line):
            syntax_issues.append((line_no, "SG_", line.strip()))

    if not messages:
        raise DbcParseError("DBC 中没有解析到 BO_ 报文定义。")

    attribute_definitions: Dict[Tuple[str, str], AttributeDefinition] = {}
    attribute_defaults: Dict[str, Any] = {}
    attribute_usages: List[AttributeUsage] = []
    node_attributes: Dict[str, Dict[str, Any]] = {}
    global_attributes: Dict[str, Any] = {}

    ba_def_re = re.compile(
        r'^\s*BA_DEF_\s+(?:(BU_|BO_|SG_|EV_)\s+)?\s*"([^"]+)"\s+'
        r'(INT|HEX|FLOAT|STRING|ENUM)\b(.*?)\s*;\s*$'
    )
    enum_item_re = re.compile(r'"([^"]*)"')
    for line_no, definition_line in enumerate(lines, start=1):
        definition_match = ba_def_re.match(definition_line)
        if not definition_match:
            continue
        scope = (definition_match.group(1) or "GLOBAL").rstrip("_").upper()
        name = definition_match.group(2)
        value_type = definition_match.group(3).upper()
        tail = definition_match.group(4).strip()
        minimum = maximum = None
        enum_values: Tuple[str, ...] = tuple()
        if value_type == "ENUM":
            enum_values = tuple(enum_item_re.findall(tail))
        elif value_type in {"INT", "HEX", "FLOAT"}:
            parts = tail.split()
            if len(parts) >= 2:
                minimum = parse_float(parts[0])
                maximum = parse_float(parts[1])
        attribute_definitions[(scope, name.lower())] = AttributeDefinition(
            name=name,
            scope=scope,
            value_type=value_type,
            minimum=minimum,
            maximum=maximum,
            enum_values=enum_values,
        )

    def any_definition(name: str) -> Optional[AttributeDefinition]:
        key = name.lower()
        for (_scope, attr_key), definition in attribute_definitions.items():
            if attr_key == key:
                return definition
        return None

    def decode_attribute(scope: str, name: str, raw_value: str) -> Any:
        definition = (
            attribute_definitions.get((scope.upper(), name.lower()))
            or attribute_definitions.get(("GLOBAL", name.lower()))
            or any_definition(name)
        )
        raw_text = raw_value.strip()
        unquoted = raw_text[1:-1] if len(raw_text) >= 2 and raw_text[0] == raw_text[-1] == '"' else raw_text
        if definition is None:
            number = parse_float(unquoted)
            return number if number is not None and re.fullmatch(r"[+-]?(?:0[xX][0-9a-fA-F]+|\d+(?:\.\d+)?)", unquoted) else unquoted
        if definition.value_type == "STRING":
            return unquoted
        if definition.value_type == "ENUM":
            if raw_text.startswith('"'):
                return unquoted
            index = parse_int(unquoted)
            if index is not None and 0 <= index < len(definition.enum_values):
                return definition.enum_values[index]
            return unquoted
        if definition.value_type in {"INT", "HEX"}:
            return parse_int(unquoted)
        if definition.value_type == "FLOAT":
            return parse_float(unquoted)
        return unquoted

    ba_default_re = re.compile(r'^\s*BA_DEF_DEF_\s+"([^"]+)"\s+(.+?)\s*;\s*$')
    for line in lines:
        match = ba_default_re.match(line)
        if match:
            attribute_defaults[match.group(1).lower()] = decode_attribute("GLOBAL", match.group(1), match.group(2))

    cm_bo_re = re.compile(r'^\s*CM_\s+BO_\s+(\d+)\s+"(.*)"\s*;\s*$')
    cm_sg_re = re.compile(r'^\s*CM_\s+SG_\s+(\d+)\s+(\S+)\s+"(.*)"\s*;\s*$')
    ba_bo_re = re.compile(r'^\s*BA_\s+"([^"]+)"\s+BO_\s+(\d+)\s+(.+?)\s*;\s*$')
    ba_sg_re = re.compile(r'^\s*BA_\s+"([^"]+)"\s+SG_\s+(\d+)\s+(\S+)\s+(.+?)\s*;\s*$')
    ba_bu_re = re.compile(r'^\s*BA_\s+"([^"]+)"\s+BU_\s+(\S+)\s+(.+?)\s*;\s*$')
    ba_global_re = re.compile(r'^\s*BA_\s+"([^"]+)"\s+(.+?)\s*;\s*$')
    value_re = re.compile(r'^\s*VAL_\s+(\d+)\s+(\S+)\s+(.+?)\s*;\s*$')
    value_pair_re = re.compile(r'(-?\d+)\s+"([^"]*)"')
    sig_group_re = re.compile(r'^\s*SIG_GROUP_\s+(\d+)\s+(\S+)\s+\d+\s*:\s*(.*?)\s*;\s*$')

    # BA_/BA_DEF_ 是 CANdb++ 导入时常见的停止点。以前自定义解析器会静默
    # 跳过格式错误的属性语句，导致界面无法指出实际行号。
    for line_no, line in enumerate(lines, start=1):
        if re.match(r"^\s*BA_DEF_\s+\S", line) and not ba_def_re.match(line):
            syntax_issues.append((line_no, "BA_DEF_", line.strip()))
        elif re.match(r"^\s*BA_DEF_DEF_\s+\S", line) and not ba_default_re.match(line):
            syntax_issues.append((line_no, "BA_DEF_DEF_", line.strip()))
        elif re.match(r'^\s*BA_\s+"', line):
            global_match = ba_global_re.match(line)
            recognized = bool(
                ba_bo_re.match(line)
                or ba_sg_re.match(line)
                or ba_bu_re.match(line)
                or (global_match and not re.match(r"^(?:BO_|SG_|BU_|EV_)\b", global_match.group(2).strip()))
            )
            if not recognized:
                syntax_issues.append((line_no, "BA_", line.strip()))

    def messages_for_id(raw_id: int) -> List[Message]:
        candidates = message_by_raw_id.get(raw_id, [])
        seen: set[int] = set()
        result: List[Message] = []
        for msg in candidates:
            marker = id(msg)
            if marker not in seen:
                seen.add(marker)
                result.append(msg)
        return result

    def signal_for(msg_id: int, sig_name: str) -> Optional[Signal]:
        for msg in messages_for_id(msg_id):
            for sig in msg.signals:
                if sig.name == sig_name:
                    return sig
        normalized = msg_id & 0x1FFFFFFF if msg_id & 0x80000000 else msg_id
        return signal_by_key.get((normalized, sig_name))

    for line_no, line in enumerate(lines, start=1):
        match = sig_group_re.match(line)
        if match:
            raw_id, group_name, members = int(match.group(1)), match.group(2), match.group(3)
            member_names = tuple(item for item in re.split(r"\s+", members.strip()) if item)
            for msg in messages_for_id(raw_id):
                msg.signal_groups[group_name] = member_names
            continue

        match = cm_bo_re.match(line)
        if match:
            for msg in messages_for_id(int(match.group(1))):
                msg.comment = match.group(2)
            continue

        match = cm_sg_re.match(line)
        if match:
            sig = signal_for(int(match.group(1)), match.group(2))
            if sig:
                sig.comment = match.group(3)
            continue

        match = ba_bo_re.match(line)
        if match:
            attr, raw_id, raw_value = match.group(1), int(match.group(2)), match.group(3).strip()
            decoded = decode_attribute("BO", attr, raw_value)
            targets = messages_for_id(raw_id)
            for msg in targets:
                key = attr.lower()
                msg.attributes[key] = decoded
                msg.explicit_attributes.add(key)
            attribute_usages.append(AttributeUsage(
                attr, "BO", targets[0].name if targets else f"BO_ {raw_id}",
                raw_id & 0x1FFFFFFF, "", raw_value, decoded, line_no,
            ))
            continue

        match = ba_sg_re.match(line)
        if match:
            attr, raw_id, sig_name, raw_value = match.group(1), int(match.group(2)), match.group(3), match.group(4).strip()
            decoded = decode_attribute("SG", attr, raw_value)
            sig = signal_for(raw_id, sig_name)
            if sig:
                key = attr.lower()
                sig.attributes[key] = decoded
                sig.explicit_attributes.add(key)
            attribute_usages.append(AttributeUsage(
                attr, "SG", sig.name if sig else f"SG_ {raw_id} {sig_name}",
                raw_id & 0x1FFFFFFF, sig_name, raw_value, decoded, line_no,
            ))
            continue

        match = ba_bu_re.match(line)
        if match:
            attr, node_name, raw_value = match.group(1), match.group(2), match.group(3).strip()
            decoded = decode_attribute("BU", attr, raw_value)
            node_attributes.setdefault(node_name, {})[attr.lower()] = decoded
            attribute_usages.append(AttributeUsage(attr, "BU", node_name, None, "", raw_value, decoded, line_no))
            continue

        match = ba_global_re.match(line)
        if match:
            attr, raw_value = match.group(1), match.group(2).strip()
            # BO_/SG_/BU_ 行已在前面处理；这里仅剩全局属性。
            if not re.match(r"^(?:BO_|SG_|BU_|EV_)\b", raw_value):
                decoded = decode_attribute("GLOBAL", attr, raw_value)
                global_attributes[attr.lower()] = decoded
                attribute_usages.append(AttributeUsage(attr, "GLOBAL", "Database", None, "", raw_value, decoded, line_no))
            continue

        match = value_re.match(line)
        if match:
            raw_id, sig_name, pairs = int(match.group(1)), match.group(2), match.group(3)
            sig = signal_for(raw_id, sig_name)
            if sig:
                sig.value_table = {int(v): desc for v, desc in value_pair_re.findall(pairs)}

    db = Database(
        source=Path(path).name,
        messages=messages,
        nodes=tuple(nodes),
        node_attributes=node_attributes,
        global_attributes=global_attributes,
        attribute_definitions=attribute_definitions,
        attribute_defaults=attribute_defaults,
        attribute_usages=attribute_usages,
        syntax_issues=syntax_issues,
    )

    for msg in messages:
        _name, cycle, _explicit = effective_attribute_any(msg, db, ("GenMsgCycleTime", "CycleTime", "GenMsgCycledTime"))
        msg.cycle_time = parse_float(cycle)
        _name, send_type, _explicit = effective_attribute_any(msg, db, ("GenMsgSendType", "VFrameSendType", "SendType"))
        msg.send_type = text(send_type)
        _name, vframe, _explicit = effective_attribute_any(msg, db, ("VFrameFormat",))
        if vframe not in (None, ""):
            msg.frame_format = parse_frame_format(vframe) or msg.frame_format
            msg.bus_format = parse_bus_format(vframe)
        if msg.bus_format is None and msg.dlc is not None:
            msg.bus_format = "can_fd" if msg.dlc > 8 else "classic"
        for sig in msg.signals:
            _name, initial, _explicit = effective_attribute_any(sig, db, ("GenSigStartValue", "StartValue", "InitialValue"))
            if initial is not None:
                sig.initial_value = parse_float(initial)
            _name, invalid, _explicit = effective_attribute_any(sig, db, ("GenSigInvalidValue", "InvalidValue", "Fehlerwert"))
            if invalid is not None:
                sig.invalid_value = parse_float(invalid)

    if progress:
        progress(
            f"DBC解析完成：{len(messages)} 条报文，{sum(len(m.signals) for m in messages)} 个信号，"
            f"{len(attribute_definitions)} 个属性定义"
        )
    return db


def occupied_bits(signal: Signal) -> Optional[List[int]]:
    """将信号转换为统一的物理bit集合，兼容矩阵LSB位号与DBC Motorola MSB位号。"""
    if signal.start_bit is None or signal.length is None or signal.length <= 0:
        return None

    if signal.byte_order == "motorola":
        bits: List[int] = []
        current = signal.start_bit
        if signal.start_bit_convention == "lsb":
            # 从信号LSB反向走到MSB：bit7的前一位是下一字节的bit0，即索引减15。
            for _ in range(signal.length):
                bits.append(current)
                current = current - 15 if current % 8 == 7 else current + 1
        else:
            # DBC/Vector Motorola：从MSB走到LSB，bit0之后跳到下一字节bit7。
            for _ in range(signal.length):
                bits.append(current)
                current = current + 15 if current % 8 == 0 else current - 1
        return bits

    return list(range(signal.start_bit, signal.start_bit + signal.length))


def start_bits_equivalent(matrix_signal: Signal, dbc_signal: Signal) -> bool:
    """比较两个信号是否占用相同bit，避免Motorola起始位定义差异造成误报。"""
    if matrix_signal.start_bit is None or dbc_signal.start_bit is None:
        return matrix_signal.start_bit is dbc_signal.start_bit
    matrix_bits = occupied_bits(matrix_signal)
    dbc_bits = occupied_bits(dbc_signal)
    if matrix_bits is not None and dbc_bits is not None:
        return set(matrix_bits) == set(dbc_bits)
    return matrix_signal.start_bit == dbc_signal.start_bit


def multiplex_branch(signal: Signal) -> Optional[int]:
    """返回简单DBC复用分支号（m0、m1等）；M为复用选择器，不属于分支。"""
    match = re.fullmatch(r"m(\d+)", text(signal.multiplex), flags=re.IGNORECASE)
    return int(match.group(1)) if match else None


def signals_are_mutually_exclusive(a: Signal, b: Signal) -> bool:
    """不同简单复用分支不会同时有效，允许占用相同bit。"""
    branch_a = multiplex_branch(a)
    branch_b = multiplex_branch(b)
    return branch_a is not None and branch_b is not None and branch_a != branch_b


def self_check_database(db: Database, label: str) -> List[Difference]:
    diffs: List[Difference] = []
    id_map: Dict[int, List[Message]] = {}
    name_map: Dict[str, List[Message]] = {}
    is_dbc = label.lower() == "dbc"
    valid_name_re = re.compile(r"^[A-Za-z_][A-Za-z0-9_]*$")
    valid_fd_lengths = {0, 1, 2, 3, 4, 5, 6, 7, 8, 12, 16, 20, 24, 32, 48, 64}

    def add(
        severity: str,
        msg: Optional[Message],
        sig: Optional[Signal],
        field_name: str,
        description: str,
        rule_id: str = "",
        actual: Any = "",
    ) -> None:
        matrix_value = value_display(actual) if not is_dbc else ""
        dbc_value = value_display(actual) if is_dbc else ""
        diffs.append(Difference(
            severity, f"{label}自身检查", msg.can_id if msg else None,
            msg.name if msg else "", sig.name if sig else "", field_name,
            matrix_value, dbc_value, description, rule_id,
        ))

    for msg in db.messages:
        if msg.can_id is not None:
            id_map.setdefault(msg.can_id, []).append(msg)
        name_map.setdefault(norm_name(msg.name), []).append(msg)

    for can_id, group in id_map.items():
        unique_names = {m.name for m in group}
        if len(group) > 1 and len(unique_names) > 1:
            add("错误", group[0], None, "CAN ID", f"{label}中同一 CAN ID 被多个报文使用：{', '.join(sorted(unique_names))}。", "DBC_ID_DUP_001", format_can_id(can_id))

    for normalized_name, group in name_map.items():
        ids = {m.can_id for m in group}
        if normalized_name and len(group) > 1 and len(ids) > 1:
            add("警告", group[0], None, "报文名称", f"{label}中归一化后同名报文对应多个 CAN ID：{', '.join(format_can_id(i) for i in sorted(x for x in ids if x is not None))}", "DBC_MSG_NAME_DUP_001")

    if is_dbc:
        for line_no, token, raw_line in db.syntax_issues:
            add("错误", None, None, "DBC语法", f"第 {line_no} 行 {token} 定义无法解析：{raw_line}", "DBC_SYNTAX_001")

        node_counts: Dict[str, int] = {}
        for node in db.nodes:
            node_counts[node] = node_counts.get(node, 0) + 1
            if not valid_name_re.fullmatch(node):
                add("错误", None, None, "节点命名", f"BU_ 节点名“{node}”不符合英文字母/数字/下划线规则，且不能以数字开头。", "DBC_NAME_001", node)
        for node, count in node_counts.items():
            if count > 1:
                add("错误", None, None, "节点重复", f"BU_ 中节点“{node}”重复定义 {count} 次。", "DBC_NODE_DUP_001", node)

        def value_validation_error(definition: AttributeDefinition, value: Any) -> str:
            if definition.value_type == "ENUM":
                if norm_enum(value) not in {norm_enum(v) for v in definition.enum_values}:
                    return f"枚举值“{value_display(value)}”不在允许列表 {list(definition.enum_values)} 中"
            elif definition.value_type in {"INT", "HEX"}:
                number = parse_float(value)
                if number is None or not nearly_equal(number, int(number)):
                    return f"值“{value_display(value)}”不是整数"
                if definition.minimum is not None and number < definition.minimum or definition.maximum is not None and number > definition.maximum:
                    return f"值 {value_display(value)} 超出定义范围 {value_display(definition.minimum)}~{value_display(definition.maximum)}"
            elif definition.value_type == "FLOAT":
                number = parse_float(value)
                if number is None:
                    return f"值“{value_display(value)}”不是数值"
                if definition.minimum is not None and number < definition.minimum or definition.maximum is not None and number > definition.maximum:
                    return f"值 {value_display(value)} 超出定义范围 {value_display(definition.minimum)}~{value_display(definition.maximum)}"
            return ""

        for usage in db.attribute_usages:
            definition = definition_for(db, usage.scope, usage.name)
            msg = next((m for m in db.messages if m.can_id == usage.can_id), None) if usage.can_id is not None else None
            sig = None
            if msg and usage.signal_name:
                sig = next((item for item in msg.signals if item.name == usage.signal_name), None)
            if definition is None:
                add("错误", msg, sig, "属性定义", f"第 {usage.line_number} 行使用属性“{usage.name}”，但未找到对应 BA_DEF_ 定义。", "DBC_ATTR_UNDEFINED_001", usage.raw_value)
                continue
            error = value_validation_error(definition, usage.decoded_value)
            if error:
                # 许多项目DBC模板会把数值属性的 BA_DEF_ 范围保留为 0~0，
                # 但仍通过 BA_ 赋予有效的非零工程值。此类范围不一致默认作为警告，
                # 避免把可正常使用的项目DBC批量判成严重错误。类型和枚举错误仍为错误。
                is_range_mismatch = "超出定义范围" in error
                severity = "警告" if is_range_mismatch else "错误"
                suffix = "（项目模板可能使用占位范围，请结合项目规范确认）" if is_range_mismatch else ""
                add(severity, msg, sig, "属性值", f"属性“{usage.name}”{error}。{suffix}", "DBC_ATTR_VALUE_001", usage.decoded_value)

        for default_name, default_value in db.attribute_defaults.items():
            # 部分Vector模板会为枚举/字符串属性保留空默认值，空值本身不参与范围校验。
            if default_value in (None, ""):
                continue
            definitions = [d for (_scope, key), d in db.attribute_definitions.items() if key == default_name]
            if not definitions:
                add("警告", None, None, "属性默认值", f"BA_DEF_DEF_ 为属性“{default_name}”设置默认值，但没有找到 BA_DEF_ 定义。", "DBC_ATTR_DEFAULT_001", default_value)
            else:
                validation_errors = [value_validation_error(definition, default_value) for definition in definitions]
                if all(validation_errors):
                    range_only = all("超出定义范围" in error for error in validation_errors)
                    if range_only and vector_spec_for(definitions[0].scope, definitions[0].name) is not None:
                        continue
                    severity = "警告" if range_only else "错误"
                    suffix = "（项目模板可能使用占位范围，请结合项目规范确认）" if range_only else ""
                    add(severity, None, None, "属性默认值", f"属性“{definitions[0].name}”的默认值 {value_display(default_value)} 不符合定义范围或枚举。{suffix}", "DBC_ATTR_DEFAULT_002", default_value)

    for msg in db.messages:
        if is_dbc and not valid_name_re.fullmatch(msg.name):
            add("错误", msg, None, "报文命名", "BO_ 报文名只能使用英文字母、数字和下划线，且不能以数字开头。", "DBC_NAME_001", msg.name)

        if msg.can_id is not None:
            if msg.can_id < 0 or msg.can_id > 0x1FFFFFFF:
                add("错误", msg, None, "CAN ID", "CAN ID 超出 CAN 29位标识符允许范围 0x0~0x1FFFFFFF。", "DBC_ID_RANGE_001", format_can_id(msg.can_id))
            elif msg.frame_format == "standard" and msg.can_id > 0x7FF:
                add("错误", msg, None, "帧格式", "报文标记为标准帧，但 CAN ID 超过 0x7FF。", "DBC_FRAME_ID_001", msg.frame_format)

        if msg.dlc is not None and msg.dlc < 0:
            add("错误", msg, None, "DLC", f"DLC={msg.dlc}，不能小于0。", "DBC_DLC_001", msg.dlc)
        if msg.bus_format == "classic" and msg.dlc is not None and msg.dlc > 8:
            add("错误", msg, None, "DLC", "VFrameFormat为Classical CAN，但DLC超过8。", "DBC_CLASSIC_DLC_001", msg.dlc)
        if msg.bus_format == "can_fd" and msg.dlc is not None and msg.dlc not in valid_fd_lengths:
            add("错误", msg, None, "CAN FD长度", f"CAN FD 数据长度 {msg.dlc} 不在合法集合 {sorted(valid_fd_lengths)} 中。", "DBC_CANFD_DLC_001", msg.dlc)
        if msg.id_frame_format and msg.frame_format and msg.id_frame_format != msg.frame_format:
            add("错误", msg, None, "VFrameFormat", f"BO_标识符判定为{msg.id_frame_format}，但VFrameFormat配置为{msg.frame_format}。", "DBC_VFRAME_ID_001", msg.frame_format)

        if is_dbc:
            if msg.sender and msg.sender != "Vector__XXX" and msg.sender not in db.nodes:
                add("错误", msg, None, "发送节点", f"发送节点“{msg.sender}”未在 BU_ 节点列表中定义。", "DBC_NODE_REF_001", msg.sender)
            if msg.sender and not valid_name_re.fullmatch(msg.sender) and msg.sender != "Vector__XXX":
                add("错误", msg, None, "发送节点命名", f"发送节点“{msg.sender}”命名不合法。", "DBC_NAME_001", msg.sender)

            _send_name, send_value, send_explicit = effective_attribute_any(msg, db, ("GenMsgSendType", "VFrameSendType", "SendType"))
            _cycle_name, cycle_value, cycle_explicit = effective_attribute_any(msg, db, ("GenMsgCycleTime", "CycleTime", "GenMsgCycledTime"))
            send_norm = norm_enum(send_value)
            cycle_number = parse_float(cycle_value)
            is_cyclic = "cyclic" in send_norm or "periodic" in send_norm
            is_event = any(token in send_norm for token in ("event", "notused", "nomsgsendtype"))
            if cycle_number is not None and cycle_number > 0 and not send_norm:
                add("错误", msg, None, "发送类型", "周期大于0，但没有有效的GenMsgSendType。", "DBC_TX_ATTR_001", cycle_number)
            if is_cyclic:
                if not send_explicit:
                    add("错误", msg, None, "GenMsgSendType", "周期型报文必须通过BA_显式配置GenMsgSendType，不能只依赖默认值。", "DBC_TX_EXPLICIT_001", send_value)
                if not cycle_explicit:
                    add("错误", msg, None, "GenMsgCycleTime", "周期型报文必须通过BA_显式配置GenMsgCycleTime，不能只依赖默认值。", "DBC_TX_EXPLICIT_002", cycle_value)
                if cycle_number is None or cycle_number <= 0:
                    add("错误", msg, None, "GenMsgCycleTime", "周期型报文的周期必须大于0。", "DBC_TX_CYCLE_001", cycle_value)
            elif is_event and cycle_number is not None and cycle_number > 0:
                add("错误", msg, None, "发送属性冲突", "事件型/NotUsed报文配置了大于0的GenMsgCycleTime，周期与发送类型冲突。", "DBC_TX_CONFLICT_001", f"{send_value}/{cycle_number}")

            _vframe_name, vframe_value, vframe_explicit = effective_attribute_any(msg, db, ("VFrameFormat",))
            _brs_name, brs_value, brs_explicit = effective_attribute_any(msg, db, ("CANFD_BRS",))
            if brs_explicit and parse_yes(brs_value) is True and msg.bus_format != "can_fd":
                add("警告", msg, None, "CANFD_BRS", "Classical CAN报文启用了CANFD_BRS，请确认属性配置。", "DBC_CANFD_BRS_001", brs_value)
            if msg.dlc is not None and msg.dlc > 8 and not vframe_explicit:
                add("警告", msg, None, "VFrameFormat", "DLC大于8，疑似CAN FD报文，但VFrameFormat没有显式赋值。部分转换工具需要该属性才能正确识别CAN FD。", "DBC_EXPERIENCE_FD_001", vframe_value)
            if not msg.sender or msg.sender == "Vector__XXX":
                add("提示", msg, None, "发送节点", "报文发送节点为空或为Vector__XXX。作为占位DBC可能正常；用于ECU集成前建议明确Tx节点。", "DBC_EXPERIENCE_TX_001", msg.sender)
            if not msg.signals and not any(message_kind_flags(msg, db).values()):
                add("提示", msg, None, "信号定义", "普通应用报文中没有SG_信号定义。若该报文按原始字节处理可忽略，否则建议确认是否漏配信号。", "DBC_EXPERIENCE_EMPTY_MSG_001")
            if msg.can_id is not None and 0x700 <= msg.can_id <= 0x7FF and not message_kind_flags(msg, db)["diag"]:
                add("提示", msg, None, "报文用途", "该标准帧ID位于0x700~0x7FF，工程中常用于诊断/ISO-TP，但仅凭ID不能定论；请确认报文用途和发送属性。", "DBC_EXPERIENCE_DIAG_ID_001", format_can_id(msg.can_id))

        sig_names: Dict[str, List[Signal]] = {}
        bit_owner: Dict[int, Signal] = {}
        crc_signals: List[Signal] = []
        counter_signals: List[Signal] = []
        for sig in msg.signals:
            sig_names.setdefault(norm_name(sig.name), []).append(sig)
            if is_dbc and not valid_name_re.fullmatch(sig.name):
                add("错误", msg, sig, "信号命名", "SG_ 信号名只能使用英文字母、数字和下划线，且不能以数字开头。", "DBC_NAME_001", sig.name)
            if is_dbc:
                for receiver in sig.receivers:
                    if receiver not in db.nodes:
                        add("错误", msg, sig, "接收节点", f"接收节点“{receiver}”未在 BU_ 节点列表中定义。", "DBC_NODE_REF_002", receiver)
                    if not valid_name_re.fullmatch(receiver):
                        add("错误", msg, sig, "接收节点命名", f"接收节点“{receiver}”命名不合法。", "DBC_NAME_001", receiver)

            if sig.length is not None and sig.length <= 0:
                add("错误", msg, sig, "信号长度", f"信号长度={sig.length}，必须大于0。", "DBC_SIG_LENGTH_001", sig.length)
            if sig.factor is not None and nearly_equal(sig.factor, 0.0):
                add("错误", msg, sig, "精度/Factor", "Factor=0 会导致所有原始值映射到同一个物理值，请确认配置。", "DBC_FACTOR_001", sig.factor)
            if sig.minimum is not None and sig.maximum is not None and sig.minimum > sig.maximum:
                add("错误", msg, sig, "物理范围", f"最小值 {value_display(sig.minimum)} 大于最大值 {value_display(sig.maximum)}。", "DBC_PHY_RANGE_001")

            bits = occupied_bits(sig)
            if bits is not None and msg.dlc is not None:
                limit = msg.dlc * 8
                invalid_bits = [bit for bit in bits if bit < 0 or bit >= limit]
                if invalid_bits:
                    add("错误", msg, sig, "信号范围", f"信号占用位超出 DLC 范围 0~{limit - 1}，异常位：{invalid_bits[:12]}" + ("..." if len(invalid_bits) > 12 else ""), "DBC_SIG_OUT_OF_RANGE_001")

            if bits is not None:
                conflicting_bits = []
                conflicting_signals: Dict[str, Signal] = {}
                for bit in bits:
                    owner = bit_owner.get(bit)
                    if owner is not None and not signals_are_mutually_exclusive(owner, sig):
                        conflicting_bits.append(bit)
                        conflicting_signals[owner.name] = owner
                if conflicting_bits:
                    other_names = sorted(conflicting_signals)
                    add("错误", msg, sig, "信号重叠", f"与信号 {', '.join(other_names)} 重叠，占用位：{sorted(set(conflicting_bits))[:12]}" + ("..." if len(set(conflicting_bits)) > 12 else ""), "DBC_SIG_OVERLAP_001")
                for bit in bits:
                    owner = bit_owner.get(bit)
                    if owner is None or not signals_are_mutually_exclusive(owner, sig):
                        bit_owner.setdefault(bit, sig)

            if is_dbc and sig.length is not None and sig.length > 0 and sig.signed is not None:
                raw_min = -(1 << (sig.length - 1)) if sig.signed else 0
                raw_max = (1 << (sig.length - 1)) - 1 if sig.signed else (1 << sig.length) - 1
                for attr_label, raw_value, rule_id in (
                    ("初始值", sig.initial_value, "DBC_INIT_RAW_RANGE_001"),
                    ("无效值", sig.invalid_value, "DBC_INVALID_RAW_RANGE_001"),
                ):
                    if raw_value is None:
                        continue
                    if not nearly_equal(raw_value, int(raw_value)):
                        add("警告", msg, sig, attr_label, f"{attr_label}通常应是原始整数值，当前为 {value_display(raw_value)}。请确认该属性是否采用物理值语义。", rule_id, raw_value)
                    elif raw_value < raw_min or raw_value > raw_max:
                        add("错误", msg, sig, attr_label, f"{attr_label}={value_display(raw_value)} 超出 {sig.length} bit {'有符号' if sig.signed else '无符号'}原始范围 {raw_min}~{raw_max}。", rule_id, raw_value)

            if is_dbc and not sig.receivers:
                add("提示", msg, sig, "接收节点", "信号没有有效接收节点。对于内部、保留或诊断原始载荷信号可能是正常情况；普通应用信号建议确认。", "DBC_EXPERIENCE_RX_001")

            if sig.length is not None and sig.length > 0 and sig.value_table:
                if sig.signed:
                    min_raw = -(1 << (sig.length - 1)); max_raw = (1 << (sig.length - 1)) - 1
                else:
                    min_raw = 0; max_raw = (1 << sig.length) - 1
                bad_values = [v for v in sig.value_table if v < min_raw or v > max_raw]
                if bad_values:
                    add("警告", msg, sig, "枚举范围", f"枚举值 {bad_values[:10]} 超出 {sig.length} bit 信号可表达范围 {min_raw}~{max_raw}。", "DBC_ENUM_RANGE_001")

            if (sig.length is not None and sig.length > 0 and sig.factor is not None and sig.offset is not None
                    and sig.minimum is not None and sig.maximum is not None and sig.signed is not None):
                if sig.signed:
                    raw_min = -(1 << (sig.length - 1)); raw_max = (1 << (sig.length - 1)) - 1
                else:
                    raw_min = 0; raw_max = (1 << sig.length) - 1
                physical_a = raw_min * sig.factor + sig.offset
                physical_b = raw_max * sig.factor + sig.offset
                calc_min, calc_max = min(physical_a, physical_b), max(physical_a, physical_b)
                if sig.minimum < calc_min - 1e-9 or sig.maximum > calc_max + 1e-9:
                    add("警告", msg, sig, "物理范围", f"配置范围 [{value_display(sig.minimum)}, {value_display(sig.maximum)}] 超出按位宽/精度计算的理论范围 [{value_display(calc_min)}, {value_display(calc_max)}]。", "DBC_PHY_CALC_001")

            _func_name, func_value, _ = effective_attribute_any(sig, db, ("GenSigFuncType", "SigFuncType")) if is_dbc else ("", None, False)
            sig_text = f"{sig.name} {sig.comment} {value_display(func_value)}"
            if re.search(r"(?i)(crc|checksum|crc-8|crc-32)", sig_text):
                crc_signals.append(sig)
            if re.search(r"(?i)(alive.*counter|rolling.*counter|message.*counter|msgcounter|counter16|(^|[_\-])(cnt|counter)([_\-]|$))", sig_text):
                counter_signals.append(sig)

        for group in sig_names.values():
            if len(group) > 1:
                add("错误", msg, group[0], "信号名称", f"同一报文中存在 {len(group)} 个归一化后同名信号。", "DBC_SIG_NAME_DUP_001")

        if (crc_signals or counter_signals) and not (crc_signals and counter_signals):
            add("警告", msg, None, "E2E组成", f"检测到CRC类信号 {[s.name for s in crc_signals]}、Counter类信号 {[s.name for s in counter_signals]}，未成对出现。", "E2E_PAIR_001")
        if is_dbc:
            _profile_name, profile_value, _ = effective_attribute_any(msg, db, ("E2EProfile",))
            _length_name, e2e_length, _ = effective_attribute_any(msg, db, ("E2EDataLength",))
            if profile_value not in (None, "", 0, "0", "No", "Off") and (e2e_length is None or parse_float(e2e_length) is None or parse_float(e2e_length) <= 0):
                suggestion = e2e_data_length_suggestion(msg)
                add(
                    "警告", msg, None, "E2EDataLength",
                    f"已配置E2EProfile，但E2EDataLength缺失或不大于0。{suggestion}",
                    "E2E_ATTR_001", e2e_length,
                )

    return diffs


# ---------------------------------------------------------------------------
# Vector Technical Reference 1.12 baseline rules
# ---------------------------------------------------------------------------
# This catalog is intentionally independent from the local BA_DEF_ ranges.
# Many project templates use 0..0 placeholders. Values are therefore checked
# against the published Vector ranges; a local definition mismatch is emitted
# once per attribute definition instead of once per BA_ assignment.

VECTOR_EXACT_SPECS: Dict[str, Dict[str, Any]] = {
    # General attributes
    "baudrate": {"scopes": ("GLOBAL",), "types": ("INT",), "minimum": 0, "maximum": 1_000_000},
    "samplepointmin": {"scopes": ("GLOBAL",), "types": ("INT",), "minimum": 50, "maximum": 100},
    "samplepointmax": {"scopes": ("GLOBAL",), "types": ("INT",), "minimum": 50, "maximum": 100},
    "syncjumpwidthmin": {"scopes": ("GLOBAL",), "types": ("INT",), "minimum": 1, "maximum": 4},
    "syncjumpwidthmax": {"scopes": ("GLOBAL",), "types": ("INT",), "minimum": 1, "maximum": 4},
    "nbtmin": {"scopes": ("GLOBAL",), "types": ("INT",), "minimum": 6, "maximum": 25},
    "nbtmax": {"scopes": ("GLOBAL",), "types": ("INT",), "minimum": 8, "maximum": 25},
    "manufacturer": {"scopes": ("GLOBAL",), "types": ("STRING",), "allowed": ("Vector",)},
    "dbname": {"scopes": ("GLOBAL",), "types": ("STRING",)},
    "bustype": {"scopes": ("GLOBAL",), "types": ("STRING",), "allowed": ("CAN", "CAN FD")},
    "vframeformat": {"scopes": ("BO",), "types": ("ENUM",),
        "allowed": ("CAN Standard", "CAN Extended", "CAN FD Standard", "CAN FD Extended", "StandardCAN", "ExtendedCAN", "StandardCAN_FD", "ExtendedCAN_FD"),
        "valid_indices": (0, 1, 14, 15)},

    # COM
    "genmsgilsupport": {"scopes": ("BO",), "types": ("ENUM",), "allowed": ("No", "Yes"), "valid_indices": (0, 1)},
    "genmsgsendtype": {"scopes": ("BO",), "types": ("ENUM",), "allowed": ("Cyclic", "NotUsed", "NoMsgSendType", "CyclicAndSpontanX", "Event"), "valid_indices": tuple(range(9))},
    "gensigsendtype": {"scopes": ("SG",), "types": ("ENUM",), "allowed": ("Cyclic", "OnWrite", "OnWriteWithRepetition", "OnChange", "OnChangeWithRepetition", "NotUsed", "NoSigSendType", "IfActive", "IfActiveWithRepetition"), "valid_indices": tuple(range(8))},
    "genmsgcycletime": {"scopes": ("BO",), "types": ("INT",), "minimum": 0, "maximum": 65535},
    "genmsgcycletimefast": {"scopes": ("BO",), "types": ("INT",), "minimum": 0, "maximum": 65535},
    "gensigstartvalue": {"scopes": ("SG",), "types": ("INT", "FLOAT", "STRING"), "minimum": 0, "maximum": 2147483647},
    "gensiginactivevalue": {"scopes": ("SG",), "types": ("INT", "STRING"), "minimum": 0, "maximum": 2147483647},
    "gensiginvalidvalue": {"scopes": ("SG",), "types": ("INT", "STRING"), "minimum": 0, "maximum": 2147483647},
    "genmsgdelaytime": {"scopes": ("BO",), "types": ("INT",), "minimum": 0, "maximum": 65535},
    "genmsgstartdelaytime": {"scopes": ("BO",), "types": ("INT",), "minimum": 0, "maximum": 65535},
    "genmsgnrofrepetition": {"scopes": ("BO",), "types": ("INT",), "minimum": 0, "maximum": 255},

    # E2E message-level
    "e2eprofile": {"scopes": ("BO",), "types": ("STRING",), "allowed": ("P01", "P02", "P04", "P05", "P06", "P07", "P11")},
    "e2edataid": {"scopes": ("BO", "SG"), "types": ("INT",), "minimum": 0, "maximum": 65535},
    "e2edatalength": {"scopes": ("BO", "SG"), "types": ("INT",), "minimum": 0, "maximum": 65535},

    # SecOC
    "sc_message": {"scopes": ("BO",), "types": ("ENUM",), "allowed": ("No", "Yes", "Yes, split message")},
    "sc_linkpos": {"scopes": ("BO",), "types": ("INT",), "minimum": 0, "maximum": 65535},
    "scl_linklen": {"scopes": ("BO",), "types": ("INT",), "minimum": 0, "maximum": 65535},
    "scl_cryptographicmessageid": {"scopes": ("BO",), "types": ("INT",), "minimum": 0, "maximum": 65535},
    "scp_authinfotxlength": {"scopes": ("BO",), "types": ("INT",), "minimum": 0, "maximum": 512},
    "scp_dataid": {"scopes": ("BO",), "types": ("INT",), "minimum": 0, "maximum": 65535},
    "scp_freshnessvalueid": {"scopes": ("BO",), "types": ("INT",), "minimum": 0, "maximum": 65535},
    "scp_freshnessvaluelength": {"scopes": ("BO",), "types": ("INT",), "minimum": 0, "maximum": 512},
    "scp_freshnessvaluetxlength": {"scopes": ("BO",), "types": ("INT",), "minimum": 0, "maximum": 512},

    # AUTOSAR NM
    "nmtype": {"scopes": ("GLOBAL",), "types": ("STRING",), "allowed": ("NmAsr", "Vector")},
    "nmasrnode": {"scopes": ("BU",), "types": ("ENUM",), "allowed": ("No", "Yes")},
    "nmasrtimeouttime": {"scopes": ("GLOBAL",), "types": ("INT",), "minimum": 1, "maximum": 65535},
    "nmasrwaitbussleeptime": {"scopes": ("GLOBAL",), "types": ("INT",), "minimum": 1, "maximum": 65535},
    "nmasrrepeatmessagetime": {"scopes": ("GLOBAL",), "types": ("INT",), "minimum": 1, "maximum": 65535},
    "nmasrmessage": {"scopes": ("BO",), "types": ("ENUM",), "allowed": ("No", "Yes")},
    "nmasrmessagecount": {"scopes": ("GLOBAL",), "types": ("INT",), "minimum": 1, "maximum": 256},
    "nmasrbaseaddress": {"scopes": ("GLOBAL",), "types": ("HEX", "INT"), "minimum": 0, "maximum": 0x1FFFFFFF},
    "nmasrcanmsgcycletime": {"scopes": ("GLOBAL",), "types": ("INT",), "minimum": 1, "maximum": 65535},
    "nmasrcanmsgreducedtime": {"scopes": ("BU",), "types": ("INT",), "minimum": 1, "maximum": 65535},
    "nmasrcanmsgcycleoffset": {"scopes": ("BU",), "types": ("INT",), "minimum": 0, "maximum": 65535},
    "nmasrnodeidentifier": {"scopes": ("BU",), "types": ("HEX", "INT"), "minimum": 0, "maximum": 255},

    # OSEK NM
    "nmnode": {"scopes": ("BU",), "types": ("ENUM",), "allowed": ("No", "Yes")},
    "nmmessage": {"scopes": ("BO",), "types": ("ENUM",), "allowed": ("No", "Yes")},
    "nmmessagecount": {"scopes": ("GLOBAL",), "types": ("INT",), "minimum": 16, "maximum": 256},
    "nmbaseaddress": {"scopes": ("GLOBAL",), "types": ("HEX", "INT"), "minimum": 0, "maximum": 0x1FFFFFFF},
    "nmstationaddress": {"scopes": ("BU",), "types": ("HEX", "INT"), "minimum": 0, "maximum": 255},

    # CanTp / DCM
    "tptxindex": {"scopes": ("BO",), "types": ("INT",), "minimum": 0, "maximum": 255},
    "diagstate": {"scopes": ("BO",), "types": ("ENUM",), "allowed": ("No", "Yes")},
    "diagrequest": {"scopes": ("BO",), "types": ("ENUM",), "allowed": ("No", "Yes")},
    "diagresponse": {"scopes": ("BO",), "types": ("ENUM",), "allowed": ("No", "Yes")},
    "diagconnection": {"scopes": ("BO",), "types": ("INT", "HEX"), "minimum": 0, "maximum": 0xFFFF},
    "diagfdonly": {"scopes": ("BO",), "types": ("ENUM",), "allowed": ("No", "Yes")},

    # J1939
    "baudratecanfd": {"scopes": ("GLOBAL",), "types": ("INT",), "minimum": 0, "maximum": 16_000_000},
    "protocoltype": {"scopes": ("GLOBAL",), "types": ("STRING",), "allowed": ("J1939",)},
    "j1939transportlayer": {"scopes": ("GLOBAL",), "types": ("ENUM",), "allowed": ("Normed", "ClassicTpOnly", "ClassicTpAndMultiPg")},
    "reroutj1939tocom": {"scopes": ("BU",), "types": ("ENUM",), "allowed": ("No", "Yes")},
    "nmj1939aac": {"scopes": ("BU",), "types": ("INT",), "minimum": 0, "maximum": 1},
    "nmj1939industrygroup": {"scopes": ("BU",), "types": ("INT",), "minimum": 0, "maximum": 7},
    "nmj1939systeminstance": {"scopes": ("BU",), "types": ("INT",), "minimum": 0, "maximum": 15},
    "nmj1939system": {"scopes": ("BU",), "types": ("INT",), "minimum": 0, "maximum": 127},
    "nmj1939function": {"scopes": ("BU",), "types": ("INT",), "minimum": 0, "maximum": 255},
    "nmj1939functioninstance": {"scopes": ("BU",), "types": ("INT",), "minimum": 0, "maximum": 7},
    "nmj1939ecuinstance": {"scopes": ("BU",), "types": ("INT",), "minimum": 0, "maximum": 3},
    "nmj1939manufacturercode": {"scopes": ("BU",), "types": ("INT",), "minimum": 0, "maximum": 2047},
    "nmj1939identitynumber": {"scopes": ("BU",), "types": ("INT",), "minimum": 0, "maximum": 2097151},
    "tpj1939vardlc": {"scopes": ("BO",), "types": ("ENUM",), "allowed": ("No", "Yes")},
    "genmsgrequestable": {"scopes": ("BO",), "types": ("INT",), "minimum": 0, "maximum": 1},
}

VECTOR_PATTERN_SPECS: Tuple[Tuple[re.Pattern[str], Dict[str, Any]], ...] = (
    (re.compile(r"^gensigtimeouttime_.+$", re.I), {"scopes": ("SG",), "types": ("INT",), "minimum": 0, "maximum": 65535}),
    (re.compile(r"^e2ep(?:01|02|04|05|06|07|11)(?:counteroffset|crcoffset|dataidnibbleoffset|maxdeltacounter|maxerrorstateinit|maxerrorstateinvalid|maxerrorstatevalid|maxnoneworrepeateddata|minokstateinit|minokstateinvalid|minokstatevalid|syncounterinit|upperheaderbitstoshift|windowsize)$", re.I), {"scopes": ("GLOBAL",), "types": ("INT",), "minimum": 0, "maximum": 65535}),
    (re.compile(r"^e2ep(?:01|02|04|05|06|07|11)profilebehavior$", re.I), {"scopes": ("GLOBAL",), "types": ("ENUM",), "allowed": ("PRE_R4_2", "R4_2")}),
    (re.compile(r"^e2ep(?:01|02|04|05|06|07|11)dataidmode$", re.I), {"scopes": ("GLOBAL",), "types": ("ENUM",), "allowed": ("all16Bit", "alternating8Bit", "lower8Bit", "lower12Bit")}),
    (re.compile(r"^e2ep(?:01|02|04|05|06|07|11)profilename$", re.I), {"scopes": ("GLOBAL",), "types": ("STRING",), "allowed": ("PROFILE_01", "PROFILE_02", "PROFILE_04", "PROFILE_05", "PROFILE_06", "PROFILE_07", "PROFILE_11")}),
    (re.compile(r"^e2edataid(?:0[1-9]|1[0-6])$", re.I), {"scopes": ("BO", "SG"), "types": ("INT",), "minimum": 0, "maximum": 65535}),
)

VECTOR_ENUM_ORDER: Dict[str, Tuple[str, ...]] = {
    "genmsgilsupport": ("No", "Yes"),
    "genmsgsendtype": ("Cyclic", "NotUsed", "NotUsed", "NotUsed", "NotUsed", "NotUsed", "NotUsed", "NotUsed", "NoMsgSendType"),
    "gensigsendtype": ("Cyclic", "OnWrite", "OnWriteWithRepetition", "OnChange", "OnChangeWithRepetition", "NotUsed", "NotUsed", "NoSigSendType"),
    "vframeformat": ("CAN Standard", "CAN Extended", "CAN FD Standard", "CAN FD Extended"),
}


def vector_spec_for(scope: str, name: str) -> Optional[Dict[str, Any]]:
    key = name.lower()
    spec = VECTOR_EXACT_SPECS.get(key)
    if spec is not None:
        return spec
    for pattern, candidate in VECTOR_PATTERN_SPECS:
        if pattern.fullmatch(key):
            return candidate
    return None


def _db_global(db: Database, name: str) -> Any:
    key = name.lower()
    if key in db.global_attributes:
        return db.global_attributes[key]
    return db.attribute_defaults.get(key)


def _node_attr(db: Database, node: str, name: str) -> Any:
    attrs = db.node_attributes.get(node, {})
    key = name.lower()
    if key in attrs:
        return attrs[key]
    return db.attribute_defaults.get(key)


def _is_power_of_two(value: Optional[int]) -> bool:
    return value is not None and value > 0 and (value & (value - 1)) == 0


def _msg_for_usage(db: Database, usage: AttributeUsage) -> Optional[Message]:
    if usage.can_id is None:
        return None
    return next((m for m in db.messages if m.can_id == usage.can_id), None)


def vector_reference_check_database(db: Database) -> List[Difference]:
    """Apply the published Vector legacy DBC rules (Technical Reference v1.12)."""
    diffs: List[Difference] = []

    def add(
        severity: str,
        rule_id: str,
        field_name: str,
        description: str,
        msg: Optional[Message] = None,
        sig: Optional[Signal] = None,
        dbc_value: Any = "",
    ) -> None:
        diffs.append(Difference(
            severity=severity,
            category=f"Vector手册v{VECTOR_MANUAL_VERSION}",
            can_id=msg.can_id if msg else None,
            message_name=msg.name if msg else "",
            signal_name=sig.name if sig else "",
            field_name=field_name,
            matrix_value="",
            dbc_value=value_display(dbc_value),
            description=description,
            rule_id=rule_id,
        ))

    # 1) Attribute definition scope/type/range. Local 0..0 placeholders become one warning.
    for (scope, key), definition in db.attribute_definitions.items():
        spec = vector_spec_for(scope, definition.name)
        if spec is None:
            continue
        expected_scopes = tuple(spec.get("scopes", ()))
        expected_types = tuple(spec.get("types", ()))
        if expected_scopes and scope not in expected_scopes:
            add("错误", "VEC_ATTR_SCOPE_001", "属性对象类型",
                f"属性“{definition.name}”定义在 {scope}，Vector手册要求对象类型为 {', '.join(expected_scopes)}。",
                dbc_value=scope)
        if expected_types and definition.value_type not in expected_types:
            add("错误", "VEC_ATTR_TYPE_001", "属性数据类型",
                f"属性“{definition.name}”类型为 {definition.value_type}，Vector手册要求 {', '.join(expected_types)}。",
                dbc_value=definition.value_type)
        if definition.value_type in {"INT", "HEX", "FLOAT"}:
            spec_min = spec.get("minimum")
            spec_max = spec.get("maximum")
            if spec_min is not None and spec_max is not None and (
                definition.minimum is None or definition.maximum is None
                or not nearly_equal(definition.minimum, float(spec_min))
                or not nearly_equal(definition.maximum, float(spec_max))
            ):
                add("警告", "VEC_ATTR_DEF_RANGE_001", "属性定义范围",
                    f"属性“{definition.name}”的BA_DEF_范围为 {value_display(definition.minimum)}~{value_display(definition.maximum)}，"
                    f"Vector手册v{VECTOR_MANUAL_VERSION}给出的范围为 {value_display(spec_min)}~{value_display(spec_max)}。"
                    "项目模板若使用0~0占位可保留，但实际值仍按手册范围校验。",
                    dbc_value=f"{value_display(definition.minimum)}~{value_display(definition.maximum)}")
        expected_order = VECTOR_ENUM_ORDER.get(key)
        if expected_order and definition.value_type == "ENUM":
            actual = tuple(norm_enum(v) for v in definition.enum_values)
            expected = tuple(norm_enum(v) for v in expected_order)
            if actual != expected:
                add("警告", "VEC_ATTR_ENUM_ORDER_001", "枚举名称/顺序",
                    f"属性“{definition.name}”的枚举名称或顺序与Vector手册不一致。枚举索引具有配置含义，建议按手册顺序定义。",
                    dbc_value=list(definition.enum_values))

    # 2) Every known explicit usage is checked against the published range/value set.
    for usage in db.attribute_usages:
        spec = vector_spec_for(usage.scope, usage.name)
        if spec is None:
            continue
        msg = _msg_for_usage(db, usage)
        sig = None
        if msg and usage.signal_name:
            sig = next((item for item in msg.signals if item.name == usage.signal_name), None)
        allowed = spec.get("allowed")
        if allowed:
            allowed_norm = {norm_enum(item) for item in allowed}
            raw_index = parse_int(usage.raw_value) if re.fullmatch(r"[+-]?\d+", usage.raw_value.strip()) else None
            valid_indices = set(spec.get("valid_indices", ()))
            # Vector手册特别强调枚举排序。项目DBC可使用不同显示名称；只要BA_使用的
            # 数值索引仍落在手册定义槽位，就不对每条报文重复报错，名称/顺序差异在
            # BA_DEF_层面统一给出一次警告。
            index_is_valid = raw_index is not None and raw_index in valid_indices
            if not index_is_valid and norm_enum(usage.decoded_value) not in allowed_norm:
                add("警告", "VEC_ATTR_VALUE_001", "属性值",
                    f"属性“{usage.name}”值“{value_display(usage.decoded_value)}”不在Vector手册标准显示值 {list(allowed)} 中。"
                    "若项目通过兼容枚举槽位扩展名称，请结合转换工具版本确认。",
                    msg, sig, usage.decoded_value)
        if "minimum" in spec or "maximum" in spec:
            number = parse_float(usage.decoded_value)
            if number is None:
                add("错误", "VEC_ATTR_NUMERIC_001", "属性值",
                    f"属性“{usage.name}”应为数值，当前为“{value_display(usage.decoded_value)}”。",
                    msg, sig, usage.decoded_value)
            else:
                minimum = spec.get("minimum")
                maximum = spec.get("maximum")
                if minimum is not None and number < minimum or maximum is not None and number > maximum:
                    add("错误", "VEC_ATTR_RANGE_001", "属性值",
                        f"属性“{usage.name}”值 {value_display(number)} 超出Vector手册范围 {value_display(minimum)}~{value_display(maximum)}。",
                        msg, sig, number)

    # 3) General CAN / CAN-FD consistency.
    fd_messages = [m for m in db.messages if m.bus_format == "can_fd" or (m.dlc or 0) > 8]
    bus_type = _db_global(db, "BusType")
    if fd_messages and norm_enum(bus_type) != "canfd":
        add("错误", "VEC_GENERAL_BUSTYPE_001", "BusType",
            "DBC包含CAN-FD报文，但网络属性BusType未设置为“CAN FD”。",
            fd_messages[0], dbc_value=bus_type)
    manufacturer = _db_global(db, "Manufacturer")
    if manufacturer not in (None, "") and norm_enum(manufacturer) != "vector":
        add("警告", "VEC_GENERAL_MANUFACTURER_001", "Manufacturer",
            "Vector手册要求Manufacturer值为“Vector”。", dbc_value=manufacturer)

    # 4) COM-dependent rules.
    timeout_defs = {key for (scope, key), _d in db.attribute_definitions.items() if scope == "SG" and key.startswith("gensigtimeouttime_")}
    for msg in db.messages:
        repetition = parse_int(effective_attribute(msg, db, "GenMsgNrOfRepetition")[0]) or 0
        fast_time = parse_float(effective_attribute(msg, db, "GenMsgCycleTimeFast")[0])
        if repetition > 0 and (fast_time is None or fast_time <= 0):
            add("错误", "VEC_COM_REPETITION_001", "GenMsgCycleTimeFast",
                "GenMsgNrOfRepetition大于0时，必须用GenMsgCycleTimeFast定义重复发送间隔。",
                msg, dbc_value=fast_time)
        for sig in msg.signals:
            sig_send = effective_attribute(sig, db, "GenSigSendType")[0]
            if norm_enum(sig_send) in {"onchange", "onchangewithrepetition"} and (sig.length or 0) > 32:
                add("错误", "VEC_COM_ONCHANGE_SIZE_001", "GenSigSendType",
                    "Vector手册规定OnChange仅支持长度不大于4 Byte的信号。",
                    msg, sig, f"{sig_send}/{sig.length}bit")
            # If timeout attributes are used in the DBC, each receiver needs a dedicated definition.
            if timeout_defs:
                for receiver in sig.receivers:
                    expected_key = f"gensigtimeouttime_{receiver}".lower()
                    if expected_key not in timeout_defs:
                        add("警告", "VEC_COM_TIMEOUT_DEF_001", "GenSigTimeoutTime_<Ecu>",
                            f"信号接收节点“{receiver}”没有对应的属性定义 GenSigTimeoutTime_{receiver}。",
                            msg, sig, receiver)

    # 5) E2E consistency without inventing profile-specific requirements.
    for msg in db.messages:
        profile = effective_attribute(msg, db, "E2EProfile")[0]
        e2e_length = parse_float(effective_attribute(msg, db, "E2EDataLength")[0])
        e2e_keys = {key for key in msg.attributes if key.startswith("e2e")}
        for sig in msg.signals:
            e2e_keys.update(key for key in sig.attributes if key.startswith("e2e"))
        if profile not in (None, ""):
            if norm_enum(profile) not in {norm_enum(v) for v in ("P01", "P02", "P04", "P05", "P06", "P07", "P11")}:
                add("错误", "VEC_E2E_PROFILE_001", "E2EProfile",
                    "E2EProfile不在Vector手册支持的P01/P02/P04/P05/P06/P07/P11集合中。",
                    msg, dbc_value=profile)
            if e2e_length is None or e2e_length <= 0:
                suggestion = e2e_data_length_suggestion(msg)
                add("错误", "VEC_E2E_LENGTH_001", "E2EDataLength",
                    f"配置E2EProfile时，E2EDataLength应存在且大于0。{suggestion}",
                    msg, dbc_value=e2e_length)
        elif e2e_keys:
            add("警告", "VEC_E2E_PROFILE_MISSING_001", "E2EProfile",
                "报文或信号存在E2E属性，但报文未配置E2EProfile；请确认属性映射层级。",
                msg, dbc_value=sorted(e2e_keys))

    # 6) SecOC relations that can be derived from the DBC alone.
    for msg in db.messages:
        sc_message = effective_attribute(msg, db, "SC_Message")[0]
        sc_norm = norm_enum(sc_message)
        if sc_norm in {"2", "yessplitmessage", "splitmessage"}:
            for attr in ("SC_LinkPos", "SCL_LinkLen", "SCL_CryptographicMessageID"):
                value = effective_attribute(msg, db, attr)[0]
                if value in (None, ""):
                    add("错误", "VEC_SECOC_SPLIT_001", attr,
                        f"SC_Message为split message时必须配置{attr}。", msg)
        auth_tx = parse_int(effective_attribute(msg, db, "SCP_AuthInfoTxLength")[0])
        fresh_tx = parse_int(effective_attribute(msg, db, "SCP_FreshnessValueTxLength")[0])
        fresh_full = parse_int(effective_attribute(msg, db, "SCP_FreshnessValueLength")[0])
        if auth_tx is not None and fresh_tx is not None and (auth_tx + fresh_tx) % 8 != 0:
            add("错误", "VEC_SECOC_BYTE_BOUNDARY_001", "SecOC长度",
                "SCP_AuthInfoTxLength与SCP_FreshnessValueTxLength之和必须落在字节边界。",
                msg, dbc_value=f"{auth_tx}+{fresh_tx}")
        if fresh_tx is not None and fresh_full is not None and fresh_tx > fresh_full:
            add("错误", "VEC_SECOC_FRESHNESS_001", "Freshness长度",
                "截断Freshness Value长度不能大于完整Freshness Value长度。",
                msg, dbc_value=f"{fresh_tx}>{fresh_full}")

    # 7) AUTOSAR NM network rules.
    asr_nm_messages = [m for m in db.messages if parse_yes(effective_attribute(m, db, "NmAsrMessage")[0]) is True]
    nm_type = _db_global(db, "NmType")
    if asr_nm_messages and norm_enum(nm_type) != "nmasr":
        add("错误", "VEC_NM_TYPE_001", "NmType",
            "存在NmAsrMessage=Yes的报文时，网络NmType必须设置为“NmAsr”。",
            asr_nm_messages[0], dbc_value=nm_type)
    nm_count = parse_int(_db_global(db, "NmAsrMessageCount"))
    nm_base = parse_int(_db_global(db, "NmAsrBaseAddress"))
    if nm_count is not None and not _is_power_of_two(nm_count):
        add("错误", "VEC_NM_COUNT_001", "NmAsrMessageCount",
            "NmAsrMessageCount必须是2的自然数次幂。", dbc_value=nm_count)
    if nm_count and nm_base is not None and nm_base % nm_count != 0:
        add("错误", "VEC_NM_BASE_001", "NmAsrBaseAddress",
            "NmAsrBaseAddress必须是NmAsrMessageCount的整数倍。", dbc_value=nm_base)
    if nm_count and nm_base is not None:
        nm_max = nm_base + nm_count - 1
        for msg in asr_nm_messages:
            if msg.can_id is not None and not (nm_base <= msg.can_id <= nm_max):
                add("错误", "VEC_NM_RANGE_001", "NM报文范围",
                    f"NmAsrMessage=Yes，但CAN ID不在网络定义范围 {format_can_id(nm_base)}~{format_can_id(nm_max)}。",
                    msg, dbc_value=format_can_id(msg.can_id))
    nm_cycle = parse_float(_db_global(db, "NmAsrCanMsgCycleTime"))
    for node in db.nodes:
        reduced = parse_float(_node_attr(db, node, "NmAsrCanMsgReducedTime"))
        offset = parse_float(_node_attr(db, node, "NmAsrCanMsgCycleOffset"))
        if reduced is not None and nm_cycle is not None and not (0.5 * nm_cycle <= reduced < nm_cycle):
            add("错误", "VEC_NM_REDUCED_001", "NmAsrCanMsgReducedTime",
                f"节点“{node}”的ReducedTime必须大于等于周期的一半且小于周期。",
                dbc_value=f"{reduced}/{nm_cycle}")
        if offset is not None and nm_cycle is not None and offset >= nm_cycle:
            add("错误", "VEC_NM_OFFSET_001", "NmAsrCanMsgCycleOffset",
                f"节点“{node}”的CycleOffset必须小于NmAsrCanMsgCycleTime。",
                dbc_value=f"{offset}/{nm_cycle}")
    for msg in asr_nm_messages:
        if msg.sender and msg.sender != "Vector__XXX":
            node_flag = _node_attr(db, msg.sender, "NmAsrNode")
            if parse_yes(node_flag) is not True:
                add("警告", "VEC_NM_NODE_001", "NmAsrNode",
                    f"NM报文发送节点“{msg.sender}”未明确配置NmAsrNode=Yes。",
                    msg, dbc_value=node_flag)
        for sig in msg.signals:
            upper_name = sig.name.upper()
            if upper_name.endswith("_CBV") or upper_name.endswith("_SNI"):
                bits = occupied_bits(sig) or []
                byte_positions = {bit // 8 for bit in bits}
                if sig.length != 8 or len(byte_positions) != 1 or next(iter(byte_positions), -1) not in {0, 1}:
                    add("错误", "VEC_NM_SIGNAL_POS_001", "NM信号位置",
                        "_CBV/_SNI信号应为8 bit，且位于字节0或字节1。",
                        msg, sig, f"start={sig.start_bit}, len={sig.length}")

    # 8) OSEK-NM count/base rules.
    osek_count = parse_int(_db_global(db, "NmMessageCount"))
    osek_base = parse_int(_db_global(db, "NmBaseAddress"))
    if osek_count is not None and not _is_power_of_two(osek_count):
        add("错误", "VEC_OSEK_COUNT_001", "NmMessageCount",
            "OSEK-NM的NmMessageCount必须是2的自然数次幂。", dbc_value=osek_count)
    if osek_count and osek_base is not None and osek_base % osek_count != 0:
        add("错误", "VEC_OSEK_BASE_001", "NmBaseAddress",
            "NmBaseAddress必须是NmMessageCount的整数倍。", dbc_value=osek_base)

    # 9) DCM / CanTp relations.
    connections: Dict[int, List[Message]] = {}
    for msg in db.messages:
        diag_flags = {
            "DiagState": parse_yes(effective_attribute(msg, db, "DiagState")[0]) is True,
            "DiagRequest": parse_yes(effective_attribute(msg, db, "DiagRequest")[0]) is True,
            "DiagResponse": parse_yes(effective_attribute(msg, db, "DiagResponse")[0]) is True,
        }
        if sum(diag_flags.values()) > 1:
            add("错误", "VEC_DCM_ROLE_001", "诊断角色",
                "同一报文不应同时标记为Functional Request、Physical Request和Physical Response中的多个角色。",
                msg, dbc_value=[name for name, yes in diag_flags.items() if yes])
        if diag_flags["DiagRequest"] or diag_flags["DiagResponse"]:
            connection = parse_int(effective_attribute(msg, db, "DiagConnection")[0])
            if connection is None:
                add("错误", "VEC_DCM_CONNECTION_001", "DiagConnection",
                    "物理诊断请求/响应必须配置DiagConnection，以把请求和响应归入同一连接。", msg)
            else:
                connections.setdefault(connection, []).append(msg)
        tp_index = parse_int(effective_attribute(msg, db, "TpTxIndex")[0])
        if tp_index is not None and tp_index > 0:
            add("提示", "VEC_CANTP_TPTXINDEX_001", "TpTxIndex",
                "TpTxIndex非0会创建N-PDU；对应CanTpChannel和N-SDU仍需在DaVinci Configurator中手工创建。",
                msg, dbc_value=tp_index)
    for connection, group in connections.items():
        reqs = [m for m in group if parse_yes(effective_attribute(m, db, "DiagRequest")[0]) is True]
        resps = [m for m in group if parse_yes(effective_attribute(m, db, "DiagResponse")[0]) is True]
        if not reqs or not resps:
            exemplar = group[0]
            add("警告", "VEC_DCM_PAIR_001", "诊断连接配对",
                f"DiagConnection={connection}未同时找到物理请求和物理响应报文。",
                exemplar, dbc_value=[m.name for m in group])
        fd_values = {norm_enum(effective_attribute(m, db, "DiagFdOnly")[0]) for m in group if effective_attribute(m, db, "DiagFdOnly")[0] not in (None, "")}
        if len(fd_values) > 1:
            add("错误", "VEC_DCM_FDONLY_001", "DiagFdOnly",
                f"同一DiagConnection={connection}的请求和响应必须使用相同DiagFdOnly值。",
                group[0], dbc_value=sorted(fd_values))

    # 10) J1939 consistency.
    protocol_type = _db_global(db, "ProtocolType")
    if norm_enum(protocol_type) == "j1939":
        for msg in db.messages:
            if msg.frame_format != "extended":
                add("错误", "VEC_J1939_FRAME_001", "J1939帧格式",
                    "J1939报文应使用扩展ID（J1939 PG）。", msg, dbc_value=msg.frame_format)
            if parse_yes(effective_attribute(msg, db, "TpJ1939VarDlc")[0]) is True and msg.signals and msg.dlc is not None:
                last_sig = max(msg.signals, key=lambda item: max(occupied_bits(item) or [-1]))
                last_end = max(occupied_bits(last_sig) or [-1]) + 1
                if last_end != msg.dlc * 8:
                    add("错误", "VEC_J1939_VARDLC_001", "TpJ1939VarDlc",
                        "TpJ1939VarDlc=Yes时，最后一个信号必须延伸到报文布局末端。",
                        msg, last_sig, f"end={last_end}, dlc_bits={msg.dlc * 8}")

    # 11) XCP/CCP and generic CDD/application messages.
    for msg in db.messages:
        if not re.search(r"(?i)(xcp|ccp)", msg.name):
            continue
        requirements = {
            "GenMsgILSupport": False,
            "NmMessage": False,
            "NmAsrMessage": False,
            "DiagState": False,
            "DiagRequest": False,
            "DiagResponse": False,
        }
        for attr, expected in requirements.items():
            value = effective_attribute(msg, db, attr)[0]
            if value not in (None, "") and parse_yes(value) is not expected:
                add("错误", "VEC_XCP_LAYER_001", attr,
                    f"XCP/CCP报文的{attr}应为No（或手册允许缺省的属性不提供）。",
                    msg, dbc_value=value)
        full_bits = set(range((msg.dlc or 0) * 8))
        full_payload_signals = []
        for sig in msg.signals:
            bits = set(occupied_bits(sig) or [])
            if full_bits and bits == full_bits and sig.receivers:
                full_payload_signals.append(sig)
        if not full_payload_signals:
            add("错误", "VEC_XCP_PAYLOAD_001", "XCP载荷信号",
                "XCP/CCP报文应定义一个覆盖完整报文且用于Rx映射的信号。",
                msg, dbc_value=f"DLC={msg.dlc}")

    # 12) Update bits and invalid values.
    for msg in db.messages:
        il_support = parse_yes(effective_attribute(msg, db, "GenMsgILSupport")[0])
        signal_names = {sig.name for sig in msg.signals}
        group_names = set(msg.signal_groups)
        for sig in msg.signals:
            if sig.name.endswith("_UB"):
                base = sig.name[:-3]
                if base not in signal_names and base not in group_names:
                    add("错误", "VEC_UPDATE_BASE_001", "Update Bit关联",
                        f"Update Bit“{sig.name}”在同一报文中找不到对应信号或信号组“{base}”。",
                        msg, sig)
                if sig.length != 1:
                    add("错误", "VEC_UPDATE_LENGTH_001", "Update Bit长度",
                        "Update Bit信号长度必须为1 bit。", msg, sig, sig.length)
                send_type = effective_attribute(sig, db, "GenSigSendType")[0]
                if norm_enum(send_type) != "nosigsendtype":
                    add("错误", "VEC_UPDATE_SENDTYPE_001", "GenSigSendType",
                        "Update Bit的GenSigSendType必须为NoSigSendType。",
                        msg, sig, send_type)
                if il_support is not True:
                    add("错误", "VEC_UPDATE_IL_001", "GenMsgILSupport",
                        "使用Update Bit时，报文GenMsgILSupport必须为Yes。",
                        msg, sig, il_support)
            inactive_attr_name, inactive, inactive_explicit = effective_attribute_any(sig, db, ("GenSigInactiveValue", "GenSigInvalidValue"))
            sna_values = [raw for raw, desc in sig.value_table.items() if desc.strip().upper() == "SNA"]
            # BA_DEF_DEF_默认值会自动显示在所有信号上，但不代表项目启用了无效值机制。
            # 仅对显式BA_赋值或已经存在SNA描述的信号执行一致性检查。
            if inactive_explicit and inactive not in (None, "") and not sna_values:
                add("警告", "VEC_INVALID_SNA_001", "无效值/SNA",
                    "信号配置了无效值属性，但值表中没有名为“SNA”的值描述。Vector手册使用SNA定义AUTOSAR无效值。",
                    msg, sig, inactive)
            if sna_values and inactive_explicit and inactive not in (None, ""):
                inactive_num = parse_int(inactive)
                if inactive_num is not None and inactive_num not in sna_values:
                    add("警告", "VEC_INVALID_SNA_MISMATCH_001", "无效值/SNA",
                        "属性无效值与值表SNA编码不一致。", msg, sig,
                        f"attr={inactive_num}, SNA={sna_values}")

    # 13) System signal fan-out naming consistency.
    fanout_groups: Dict[Tuple[str, str], List[Tuple[Message, Signal]]] = {}
    for msg in db.messages:
        if parse_yes(effective_attribute(msg, db, "GenMsgILSupport")[0]) is not True:
            continue
        for sig in msg.signals:
            fanout_groups.setdefault((msg.sender, sig.name), []).append((msg, sig))
    for (_sender, _sig_name), members in fanout_groups.items():
        if len(members) < 2:
            continue
        sizes = {sig.length for _m, sig in members}
        tables = {tuple(sorted(sig.value_table.items())) for _m, sig in members}
        if len(sizes) > 1 or len(tables) > 1:
            msg, sig = members[0]
            add("警告", "VEC_SYSTEM_SIGNAL_001", "System Signal Fan-out",
                "同一发送节点存在同名COM信号，但信号长度或值表不一致，不能按手册规则合并为同一个AUTOSAR System Signal。",
                msg, sig, [f"{m.name}:{s.length}" for m, s in members])

    return diffs


DEFAULT_RULES_CONFIG: Dict[str, Any] = {
    "version": 5,
    "description": "项目培训覆盖规则。Vector Technical Reference v1.12作为内置基线，本文件只保存项目特定的NM、UDS、XCP及经验规则。",
    "rules": [
        {
            "id": "NM_PROJECT_001",
            "enabled": True,
            "severity": "错误",
            "target": "DBC",
            "type": "project_nm_message",
            "name_regex": r"(?i)(^|[_\-])(can)?nm([_\-]|$)|network\s*management",
            "allowed_send_types": ["Cyclic"],
            "require_positive_cycle": True,
            "nm_attribute_names": ["NmMessage", "NMAsrMessage", "NmhMessage"],
            "field_name": "NM属性",
            "description": "按当前项目准出规范，NM报文应显式配置为Cyclic、周期大于0，并设置NM标识属性。"
        },
        {
            "id": "UDS_REQUEST_PROJECT_001",
            "enabled": True,
            "severity": "错误",
            "target": "DBC",
            "type": "project_diag_message",
            "name_regex": r"(?i)(diag|diagnostic|uds|obd|isotp|iso.?tp)",
            "request_only": True,
            "request_attribute_names": ["DiagRequest"],
            "request_name_regex": r"(?i)(req|request|diagreq|obdreq|udsreq)",
            "allowed_send_types": ["NotUsed"],
            "expected_cycle": 0,
            "require_explicit_send_type": True,
            "require_explicit_cycle": True,
            "require_connection": True,
            "connection_attribute_names": ["diagConnection", "DiagConnection"],
            "expected_connection": 57345,
            "require_explicit_connection": True,
            "field_name": "UDS接收报文属性",
            "description": "按培训项目规范，UDS接收请求报文应显式配置GenMsgCycleTime=0、GenMsgSendType=NotUsed、diagConnection=57345。"
        },
        {
            "id": "XCP_PROJECT_001",
            "enabled": True,
            "severity": "错误",
            "target": "DBC",
            "type": "project_xcp_message",
            "name_regex": r"(?i)(^|[_\-])xcp([_\-]|$)",
            "forbidden_explicit_attributes": ["GenMsgCycleTime"],
            "field_name": "XCP周期属性",
            "description": "XCP报文不应显式配置GenMsgCycleTime，避免被周期性持续发送。"
        },
        {
            "id": "ALIVE_COUNTER_001",
            "enabled": True,
            "severity": "提示",
            "target": "DBC",
            "type": "signal_name_length_range",
            "name_regex": r"(?i)(alive.*counter|rolling.*counter|message.*counter|(^|[_\-])(alv|alive|cnt|counter)([_\-]|$))",
            "min_length": 2,
            "max_length": 8,
            "field_name": "信号长度",
            "description": "疑似Alive Counter信号的位宽不在常见2~8 bit范围内，请按项目E2E规范确认。"
        },
        {
            "id": "CRC_CHECKSUM_001",
            "enabled": True,
            "severity": "提示",
            "target": "DBC",
            "type": "signal_name_length_allowed",
            "name_regex": r"(?i)(crc|checksum|check_sum|check-sum)",
            "allowed_lengths": [8, 16, 32],
            "field_name": "信号长度",
            "description": "疑似CRC/Checksum信号的位宽不是常见8/16/32 bit，请按实际算法确认。"
        },
        {
            "id": "E2E_REQUIRED_001",
            "enabled": True,
            "severity": "警告",
            "target": "DBC",
            "type": "message_name_require_signal_patterns",
            "name_regex": r"(?i)(^|[_\-])e2e([_\-]|$)|end.?to.?end",
            "required_signal_patterns": [r"(?i)(crc|checksum)", r"(?i)(alive.*counter|rolling.*counter|(^|[_\-])(cnt|counter)([_\-]|$))"],
            "field_name": "E2E组成",
            "description": "报文名称或描述疑似E2E报文，但未同时找到CRC和Counter类信号。"
        },
        {
            "id": "BOOL_SIGNAL_001",
            "enabled": True,
            "severity": "提示",
            "target": "DBC",
            "type": "signal_binary_enum_length",
            "field_name": "信号长度",
            "description": "该信号枚举仅定义0和1，但位宽不是1 bit；请确认是否为预留编码。"
        },
        {
            "id": "PROJECT_CYCLE_TIME_001",
            "enabled": False,
            "severity": "警告",
            "target": "DBC",
            "type": "message_cycle_allowed",
            "name_regex": r".*",
            "allowed_cycle_times": [5, 10, 20, 50, 100, 200, 500, 1000],
            "field_name": "周期时间",
            "description": "周期时间不在项目允许列表中。确认项目规范后再启用。"
        }
    ]
}


class RuleConfigError(RuntimeError):
    pass


def default_rules_path() -> Path:
    return Path(__file__).resolve().with_name("can_rules.json")


def ensure_default_rules_file(path: Optional[Path] = None) -> Path:
    target = path or default_rules_path()
    if not target.exists():
        target.write_text(json.dumps(DEFAULT_RULES_CONFIG, ensure_ascii=False, indent=2), encoding="utf-8")
    return target


def load_rules(path: Optional[str] = None) -> List[Dict[str, Any]]:
    if path:
        rule_path = Path(path)
        if not rule_path.is_file():
            raise RuleConfigError(f"规则文件不存在：{rule_path}")
        try:
            config = json.loads(rule_path.read_text(encoding="utf-8-sig"))
        except (OSError, json.JSONDecodeError) as exc:
            raise RuleConfigError(f"规则文件读取失败：{exc}") from exc
    else:
        config = DEFAULT_RULES_CONFIG

    rules = config.get("rules") if isinstance(config, dict) else None
    if not isinstance(rules, list):
        raise RuleConfigError("规则文件必须包含 rules 数组。")

    validated: List[Dict[str, Any]] = []
    seen_ids: set[str] = set()
    for index, raw_rule in enumerate(rules, start=1):
        if not isinstance(raw_rule, dict):
            raise RuleConfigError(f"第 {index} 条规则不是JSON对象。")
        rule = dict(raw_rule)
        rule_id = text(rule.get("id")) or f"RULE_{index:03d}"
        if rule_id in seen_ids:
            raise RuleConfigError(f"规则ID重复：{rule_id}")
        seen_ids.add(rule_id)
        rule["id"] = rule_id
        if text(rule.get("severity")) not in {"错误", "警告", "提示"}:
            rule["severity"] = "警告"
        rule.setdefault("enabled", True)
        rule.setdefault("target", "DBC")
        validated.append(rule)
    return validated


def regex_matches(pattern: Any, value: Any) -> bool:
    try:
        return bool(re.search(text(pattern), text(value)))
    except re.error as exc:
        raise RuleConfigError(f"无效正则表达式 {pattern!r}：{exc}") from exc


def semantic_check_database(db: Database, label: str, rules: Sequence[Dict[str, Any]]) -> List[Difference]:
    diffs: List[Difference] = []
    label_key = "dbc" if label.lower() == "dbc" else "matrix"

    def add(rule: Dict[str, Any], msg: Message, signal: Optional[Signal], actual: Any, detail: str = "", field_name: str = "") -> None:
        description = text(rule.get("description"))
        if detail:
            description = f"{description} {detail}" if description else detail
        matrix_value = value_display(actual) if label_key == "matrix" else ""
        dbc_value = value_display(actual) if label_key == "dbc" else ""
        diffs.append(Difference(
            severity=text(rule.get("severity")) or "警告",
            category=f"{label}规则检查",
            can_id=msg.can_id,
            message_name=msg.name,
            signal_name=signal.name if signal else "",
            field_name=field_name or text(rule.get("field_name")) or "规则检查",
            matrix_value=matrix_value,
            dbc_value=dbc_value,
            description=description,
            rule_id=text(rule.get("id")),
        ))

    for rule in rules:
        if not bool(rule.get("enabled", True)):
            continue
        target = text(rule.get("target") or "DBC").lower()
        if target not in {label_key, "both", "all", "两者"}:
            continue
        rule_type = text(rule.get("type"))
        name_pattern = rule.get("name_regex", r".*")

        if rule_type == "project_nm_message":
            allowed = {norm_enum(v) for v in rule.get("allowed_send_types", ["Cyclic"])}
            attr_names = tuple(rule.get("nm_attribute_names", ["NmMessage", "NMAsrMessage", "NmhMessage"]))
            for msg in db.messages:
                flags = message_kind_flags(msg, db)
                if not flags["nm"] and not regex_matches(name_pattern, f"{msg.name} {msg.comment}"):
                    continue
                _send_name, send_value, send_explicit = effective_attribute_any(msg, db, ("GenMsgSendType", "VFrameSendType", "SendType"))
                _cycle_name, cycle_value, cycle_explicit = effective_attribute_any(msg, db, ("GenMsgCycleTime", "CycleTime"))
                send_norm = norm_enum(send_value)
                cycle = parse_float(cycle_value)
                if allowed and send_norm not in allowed:
                    add(rule, msg, None, send_value, f"当前发送类型={value_display(send_value)}，允许值={rule.get('allowed_send_types')}。", "GenMsgSendType")
                if not send_explicit:
                    add(rule, msg, None, send_value, "GenMsgSendType没有通过BA_显式赋值。", "GenMsgSendType")
                if bool(rule.get("require_positive_cycle", True)) and (cycle is None or cycle <= 0):
                    add(rule, msg, None, cycle_value, "GenMsgCycleTime必须大于0。", "GenMsgCycleTime")
                if bool(rule.get("require_positive_cycle", True)) and not cycle_explicit:
                    add(rule, msg, None, cycle_value, "GenMsgCycleTime没有通过BA_显式赋值。", "GenMsgCycleTime")
                _attr_name, nm_value, nm_explicit = effective_attribute_any(msg, db, attr_names)
                if parse_yes(nm_value) is not True:
                    add(rule, msg, None, nm_value, f"缺少有效NM标识属性，候选属性={list(attr_names)}。", "NM标识")
                elif not nm_explicit:
                    add(rule, msg, None, nm_value, "NM标识仅来自默认值，建议显式赋值。", "NM标识")

        elif rule_type == "project_diag_message":
            allowed = {norm_enum(v) for v in rule.get("allowed_send_types", [])}
            expected_cycle = parse_float(rule.get("expected_cycle"))
            request_only = bool(rule.get("request_only", False))
            request_attr_names = tuple(rule.get("request_attribute_names", ["DiagRequest"]))
            request_name_pattern = rule.get("request_name_regex", r"(?i)(req|request)")
            for msg in db.messages:
                flags = message_kind_flags(msg, db)
                if not flags["diag"] and not regex_matches(name_pattern, f"{msg.name} {msg.comment}"):
                    continue
                if request_only:
                    _request_attr, request_value, _request_explicit = effective_attribute_any(msg, db, request_attr_names)
                    request_by_attr = parse_yes(request_value) is True
                    request_by_name = regex_matches(request_name_pattern, f"{msg.name} {msg.comment}")
                    if not request_by_attr and not request_by_name:
                        continue

                _send_name, send_value, send_explicit = effective_attribute_any(msg, db, ("GenMsgSendType", "VFrameSendType", "SendType"))
                _cycle_name, cycle_value, cycle_explicit = effective_attribute_any(msg, db, ("GenMsgCycleTime", "CycleTime"))
                if allowed and norm_enum(send_value) not in allowed:
                    add(rule, msg, None, send_value, f"当前发送类型={value_display(send_value)}，要求值={rule.get('allowed_send_types')}。", "GenMsgSendType")
                if bool(rule.get("require_explicit_send_type", False)) and not send_explicit:
                    add(rule, msg, None, send_value, "GenMsgSendType必须通过BA_显式配置，不能只使用默认值。", "GenMsgSendType")
                if expected_cycle is not None and not nearly_equal(parse_float(cycle_value), expected_cycle):
                    add(rule, msg, None, cycle_value, f"期望GenMsgCycleTime={value_display(expected_cycle)}。", "GenMsgCycleTime")
                if bool(rule.get("require_explicit_cycle", False)) and not cycle_explicit:
                    add(rule, msg, None, cycle_value, "GenMsgCycleTime必须通过BA_显式配置，不能只使用默认值。", "GenMsgCycleTime")
                if bool(rule.get("require_connection", False)):
                    names = tuple(rule.get("connection_attribute_names", ["diagConnection", "DiagConnection"]))
                    _conn_name, conn_value, conn_explicit = effective_attribute_any(msg, db, names)
                    expected = rule.get("expected_connection")
                    if conn_value in (None, "") or (expected is not None and not nearly_equal(parse_float(conn_value), parse_float(expected))):
                        add(rule, msg, None, conn_value, f"期望{list(names)}={expected}。", "diagConnection")
                    elif bool(rule.get("require_explicit_connection", False)) and not conn_explicit:
                        add(rule, msg, None, conn_value, "diagConnection必须通过BA_显式配置，不能只使用默认值。", "diagConnection")
                    elif not conn_explicit:
                        add(rule, msg, None, conn_value, "diagConnection仅来自默认值，建议显式赋值。", "diagConnection")

        elif rule_type == "project_xcp_message":
            forbidden = [text(v) for v in rule.get("forbidden_explicit_attributes", ["GenMsgCycleTime"])]
            for msg in db.messages:
                flags = message_kind_flags(msg, db)
                if not flags["xcp"] and not regex_matches(name_pattern, f"{msg.name} {msg.comment}"):
                    continue
                for attr_name in forbidden:
                    if attr_name.lower() in msg.explicit_attributes:
                        add(rule, msg, None, msg.attributes.get(attr_name.lower()), f"属性{attr_name}不应显式配置。", attr_name)

        elif rule_type == "message_name_send_type_forbidden":
            forbidden = rule.get("send_type_regex", r"(?i)cycle|cyclic|periodic")
            for msg in db.messages:
                if regex_matches(name_pattern, f"{msg.name} {msg.comment}") and msg.send_type and regex_matches(forbidden, msg.send_type):
                    add(rule, msg, None, msg.send_type, f"当前值：{msg.send_type}")

        elif rule_type == "message_name_dlc_allowed":
            allowed = {int(v) for v in rule.get("allowed_dlc", [])}
            for msg in db.messages:
                if regex_matches(name_pattern, f"{msg.name} {msg.comment}") and msg.dlc is not None and msg.dlc not in allowed:
                    add(rule, msg, None, msg.dlc, f"当前值：{msg.dlc}")

        elif rule_type == "message_dlc_max":
            max_dlc = parse_int(rule.get("max_dlc"))
            if max_dlc is None:
                continue
            for msg in db.messages:
                if msg.dlc is not None and msg.dlc > max_dlc:
                    add(rule, msg, None, msg.dlc, f"当前值：{msg.dlc}")

        elif rule_type == "signal_name_length_allowed":
            allowed = {int(v) for v in rule.get("allowed_lengths", [])}
            for msg in db.messages:
                for sig in msg.signals:
                    if regex_matches(name_pattern, f"{sig.name} {sig.comment}") and sig.length is not None and sig.length not in allowed:
                        add(rule, msg, sig, sig.length, f"当前值：{sig.length}")

        elif rule_type == "signal_name_length_range":
            minimum = parse_int(rule.get("min_length")); maximum = parse_int(rule.get("max_length"))
            for msg in db.messages:
                for sig in msg.signals:
                    if not regex_matches(name_pattern, f"{sig.name} {sig.comment}") or sig.length is None:
                        continue
                    if (minimum is not None and sig.length < minimum) or (maximum is not None and sig.length > maximum):
                        add(rule, msg, sig, sig.length, f"当前值：{sig.length}")

        elif rule_type == "message_name_require_signal_patterns":
            patterns = rule.get("required_signal_patterns", [])
            if not isinstance(patterns, list):
                raise RuleConfigError(f"规则 {rule.get('id')} 的 required_signal_patterns 必须是数组。")
            for msg in db.messages:
                if not regex_matches(name_pattern, f"{msg.name} {msg.comment}"):
                    continue
                missing = [text(p) for p in patterns if not any(regex_matches(p, f"{sig.name} {sig.comment}") for sig in msg.signals)]
                if missing:
                    add(rule, msg, None, "缺少信号", "缺少匹配项：" + ", ".join(missing))

        elif rule_type == "signal_binary_enum_length":
            for msg in db.messages:
                for sig in msg.signals:
                    keys = set(sig.value_table.keys())
                    if keys == {0, 1} and sig.length is not None and sig.length != 1:
                        add(rule, msg, sig, sig.length, f"当前值：{sig.length} bit，枚举={sorted(keys)}")

        elif rule_type == "message_cycle_allowed":
            allowed_values = [parse_float(v) for v in rule.get("allowed_cycle_times", [])]
            allowed_values = [v for v in allowed_values if v is not None]
            for msg in db.messages:
                if not regex_matches(name_pattern, f"{msg.name} {msg.comment}") or msg.cycle_time is None:
                    continue
                if not any(nearly_equal(msg.cycle_time, v) for v in allowed_values):
                    add(rule, msg, None, msg.cycle_time, f"当前值：{value_display(msg.cycle_time)}")

        elif rule_type == "message_name_frame_format_allowed":
            allowed = {norm_enum(v) for v in rule.get("allowed_formats", [])}
            for msg in db.messages:
                if regex_matches(name_pattern, f"{msg.name} {msg.comment}") and msg.frame_format and norm_enum(msg.frame_format) not in allowed:
                    add(rule, msg, None, msg.frame_format, f"当前值：{msg.frame_format}")

        elif rule_type == "message_name_id_range":
            min_id = parse_can_id(rule.get("min_id")); max_id = parse_can_id(rule.get("max_id"))
            for msg in db.messages:
                if not regex_matches(name_pattern, f"{msg.name} {msg.comment}") or msg.can_id is None:
                    continue
                if (min_id is not None and msg.can_id < min_id) or (max_id is not None and msg.can_id > max_id):
                    add(rule, msg, None, format_can_id(msg.can_id), f"当前值：{format_can_id(msg.can_id)}")
        else:
            raise RuleConfigError(f"不支持的规则类型：{rule_type}（规则 {rule.get('id')}）")

    return diffs


def check_dbc_only(
    dbc: Database,
    semantic_rules: Optional[Sequence[Dict[str, Any]]] = None,
    vector_rules_enabled: bool = True,
) -> List[Difference]:
    """无通信矩阵时，执行DBC自身、Vector手册与项目语义规则。"""
    diffs = self_check_database(dbc, "DBC")
    if vector_rules_enabled:
        diffs.extend(vector_reference_check_database(dbc))
    if semantic_rules:
        diffs.extend(semantic_check_database(dbc, "DBC", semantic_rules))
    severity_order = {"错误": 0, "警告": 1, "提示": 2}
    return sorted(
        diffs,
        key=lambda d: (
            severity_order.get(d.severity, 9),
            d.can_id if d.can_id is not None else 1 << 30,
            d.message_name.lower(),
            d.signal_name.lower(),
            d.field_name,
        ),
    )


def compare_databases(
    matrix: Database,
    dbc: Database,
    semantic_rules: Optional[Sequence[Dict[str, Any]]] = None,
    vector_rules_enabled: bool = True,
    rename_mapping: Optional[Dict[str, Any]] = None,
) -> List[Difference]:
    diffs: List[Difference] = []

    mapping = rename_mapping if isinstance(rename_mapping, dict) else {}
    mapped_messages = mapping.get("mapping", {}).get("messages", {}) if isinstance(mapping.get("mapping", {}), dict) else {}
    mapped_signals = mapping.get("mapping", {}).get("signals", {}) if isinstance(mapping.get("mapping", {}), dict) else {}

    def dbc_message_key(msg: Message) -> str:
        raw_id = int(msg.can_id or 0)
        if msg.id_frame_format == "extended":
            raw_id |= 0x80000000
        return object_key(raw_id)

    def mapped_entry(entries: Dict[str, Any], key: str) -> Dict[str, Any]:
        value = entries.get(key, {})
        return value if isinstance(value, dict) else {}

    def mapped_name_matches(entries: Dict[str, Any], key: str, left: Any, right: Any) -> bool:
        entry = mapped_entry(entries, key)
        original = text(entry.get("original_name"))
        generated = text(entry.get("generated_name"))
        return bool(original and generated and text(left) == original and text(right) == generated)

    dbc_by_id: Dict[int, List[Message]] = {}
    dbc_by_name: Dict[str, List[Message]] = {}
    for msg in dbc.messages:
        if msg.can_id is not None:
            dbc_by_id.setdefault(msg.can_id, []).append(msg)
        dbc_by_name.setdefault(norm_name(msg.name), []).append(msg)

    matched_dbc_messages: set[int] = set()

    def add_diff(
        severity: str,
        category: str,
        matrix_msg: Message,
        dbc_msg: Optional[Message],
        signal_name: str,
        field_name: str,
        matrix_value: Any,
        dbc_value: Any,
        description: str,
    ) -> None:
        diffs.append(Difference(
            severity=severity,
            category=category,
            can_id=matrix_msg.can_id if matrix_msg.can_id is not None else (dbc_msg.can_id if dbc_msg else None),
            message_name=matrix_msg.name or (dbc_msg.name if dbc_msg else ""),
            signal_name=signal_name,
            field_name=field_name,
            matrix_value=value_display(matrix_value),
            dbc_value=value_display(dbc_value),
            description=description,
        ))

    for matrix_msg in matrix.messages:
        candidates: List[Message] = []
        match_mode = ""
        if matrix_msg.can_id is not None:
            candidates = dbc_by_id.get(matrix_msg.can_id, [])
            if candidates:
                match_mode = "ID"
        if not candidates and matrix_msg.name:
            candidates = dbc_by_name.get(norm_name(matrix_msg.name), [])
            if candidates:
                match_mode = "名称"

        if not candidates:
            add_diff(
                "错误", "缺失报文", matrix_msg, None, "", "报文", matrix_msg.name, "",
                "通信矩阵中存在该报文，但 DBC 中未找到。",
            )
            continue

        dbc_msg = candidates[0]
        matched_dbc_messages.add(id(dbc_msg))
        if len(candidates) > 1:
            add_diff(
                "警告", "报文匹配", matrix_msg, dbc_msg, "", "匹配唯一性", len(candidates), len(candidates),
                f"按{match_mode}找到多个 DBC 报文，当前使用第一个进行比较。",
            )

        if matrix_msg.can_id is not None and dbc_msg.can_id is not None and matrix_msg.can_id != dbc_msg.can_id:
            add_diff("错误", "报文差异", matrix_msg, dbc_msg, "", "CAN ID", format_can_id(matrix_msg.can_id),
                     format_can_id(dbc_msg.can_id), "报文名称匹配，但 CAN ID 不一致。")

        message_mapping_key = dbc_message_key(dbc_msg)
        message_fields = [
            ("报文名称", matrix_msg.name, dbc_msg.name, "name"),
            ("DLC", matrix_msg.dlc, dbc_msg.dlc, "number"),
            ("发送节点", matrix_msg.sender, dbc_msg.sender, "name"),
            ("周期时间", matrix_msg.cycle_time, dbc_msg.cycle_time, "number"),
            ("发送类型", matrix_msg.send_type, dbc_msg.send_type, "enum"),
            ("帧格式", matrix_msg.frame_format, dbc_msg.frame_format, "enum"),
            ("总线格式", matrix_msg.bus_format, dbc_msg.bus_format, "enum"),
        ]
        for field_name, matrix_value, dbc_value, kind in message_fields:
            # 矩阵空白表示该字段不作为基准，不制造无意义差异。
            if matrix_value in (None, ""):
                continue
            equal = False
            if kind == "number":
                equal = nearly_equal(parse_float(matrix_value), parse_float(dbc_value))
            elif kind == "name":
                equal = norm_name(matrix_value) == norm_name(dbc_value)
                if field_name == "报文名称":
                    equal = equal or mapped_name_matches(mapped_messages, message_mapping_key, matrix_value, dbc_value)
            else:
                equal = norm_enum(matrix_value) == norm_enum(dbc_value)
            if not equal:
                add_diff(
                    "错误" if field_name in {"DLC", "帧格式", "总线格式"} else "警告",
                    "报文差异", matrix_msg, dbc_msg, "", field_name, matrix_value, dbc_value,
                    f"报文{field_name}与通信矩阵不一致。",
                )
            elif kind == "name" and text(matrix_value) != text(dbc_value):
                add_diff(
                    "提示", "命名格式", matrix_msg, dbc_msg, "", field_name, matrix_value, dbc_value,
                    f"归一化后可匹配，但{field_name}的大小写、下划线或分隔符不同。",
                )

        dbc_sig_by_name: Dict[str, List[Signal]] = {}
        for sig in dbc_msg.signals:
            dbc_sig_by_name.setdefault(norm_name(sig.name), []).append(sig)
        matched_dbc_signals: set[int] = set()

        for matrix_sig in matrix_msg.signals:
            sig_candidates = dbc_sig_by_name.get(norm_name(matrix_sig.name), [])
            if not sig_candidates:
                for mapping_key, mapping_value in mapped_signals.items():
                    if not mapping_key.startswith(message_mapping_key + "|") or not isinstance(mapping_value, dict):
                        continue
                    if text(mapping_value.get("original_name")) != text(matrix_sig.name):
                        continue
                    generated_name = text(mapping_value.get("generated_name"))
                    if generated_name:
                        sig_candidates = dbc_sig_by_name.get(norm_name(generated_name), [])
                        if sig_candidates:
                            break
            if not sig_candidates:
                add_diff(
                    "错误", "缺失信号", matrix_msg, dbc_msg, matrix_sig.name, "信号", matrix_sig.name, "",
                    "通信矩阵中存在该信号，但 DBC 对应报文中未找到。",
                )
                continue
            dbc_sig = sig_candidates[0]
            matched_dbc_signals.add(id(dbc_sig))
            if len(sig_candidates) > 1:
                add_diff(
                    "警告", "信号匹配", matrix_msg, dbc_msg, matrix_sig.name, "匹配唯一性",
                    len(sig_candidates), len(sig_candidates), "归一化后找到多个同名 DBC 信号。",
                )

            if matrix_sig.start_bit is not None and not start_bits_equivalent(matrix_sig, dbc_sig):
                add_diff(
                    "错误", "信号差异", matrix_msg, dbc_msg, matrix_sig.name, "起始位",
                    matrix_sig.start_bit, dbc_sig.start_bit,
                    "信号占用bit与通信矩阵不一致。Motorola信号已自动换算矩阵LSB位号和DBC MSB位号。",
                )

            fields = [
                ("信号名称", matrix_sig.name, dbc_sig.name, "name", "提示"),
                ("信号长度", matrix_sig.length, dbc_sig.length, "number", "错误"),
                ("字节序", matrix_sig.byte_order, dbc_sig.byte_order, "enum", "错误"),
                ("符号类型", matrix_sig.signed, dbc_sig.signed, "bool", "错误"),
                ("精度/Factor", matrix_sig.factor, dbc_sig.factor, "number", "错误"),
                ("偏移量/Offset", matrix_sig.offset, dbc_sig.offset, "number", "错误"),
                ("最小值", matrix_sig.minimum, dbc_sig.minimum, "number", "警告"),
                ("最大值", matrix_sig.maximum, dbc_sig.maximum, "number", "警告"),
                ("单位", matrix_sig.unit, dbc_sig.unit, "text", "警告"),
                ("接收节点", matrix_sig.receivers, dbc_sig.receivers, "nodes", "警告"),
                ("初始值", matrix_sig.initial_value, dbc_sig.initial_value, "number", "警告"),
                ("无效值", matrix_sig.invalid_value, dbc_sig.invalid_value, "number", "警告"),
                ("复用定义", matrix_sig.multiplex, dbc_sig.multiplex, "enum", "错误"),
            ]
            for field_name, matrix_value, dbc_value, kind, severity in fields:
                if matrix_value in (None, "", tuple()):
                    continue
                if kind == "number":
                    equal = nearly_equal(parse_float(matrix_value), parse_float(dbc_value))
                elif kind == "name":
                    equal = norm_name(matrix_value) == norm_name(dbc_value) or mapped_name_matches(
                        mapped_signals,
                        signal_key(message_mapping_key, matrix_sig.name),
                        matrix_value,
                        dbc_value,
                    )
                elif kind == "enum":
                    equal = norm_enum(matrix_value) == norm_enum(dbc_value)
                elif kind == "bool":
                    equal = matrix_value is dbc_value
                elif kind == "nodes":
                    equal = {norm_name(x) for x in matrix_value} == {norm_name(x) for x in dbc_value}
                else:
                    equal = text(matrix_value) == text(dbc_value)

                if not equal:
                    add_diff(
                        severity, "信号差异", matrix_msg, dbc_msg, matrix_sig.name, field_name,
                        matrix_value, dbc_value, f"信号{field_name}与通信矩阵不一致。",
                    )
                elif kind == "name" and text(matrix_value) != text(dbc_value):
                    add_diff(
                        "提示", "命名格式", matrix_msg, dbc_msg, matrix_sig.name, field_name,
                        matrix_value, dbc_value, "归一化后可匹配，但信号名称格式不同。",
                    )

            if matrix_sig.value_table:
                matrix_table = {int(k): text(v) for k, v in matrix_sig.value_table.items()}
                dbc_table = {int(k): text(v) for k, v in dbc_sig.value_table.items()}
                if matrix_table != dbc_table:
                    add_diff(
                        "警告", "信号差异", matrix_msg, dbc_msg, matrix_sig.name, "枚举值",
                        "; ".join(f"{k}={v}" for k, v in sorted(matrix_table.items())),
                        "; ".join(f"{k}={v}" for k, v in sorted(dbc_table.items())),
                        "Value Table 与通信矩阵不一致。",
                    )

        for dbc_sig in dbc_msg.signals:
            if id(dbc_sig) not in matched_dbc_signals:
                add_diff(
                    "警告", "DBC多余信号", matrix_msg, dbc_msg, dbc_sig.name, "信号", "", dbc_sig.name,
                    "DBC 中存在该信号，但通信矩阵对应报文中没有。",
                )

    for dbc_msg in dbc.messages:
        if id(dbc_msg) not in matched_dbc_messages:
            diffs.append(Difference(
                "警告", "DBC多余报文", dbc_msg.can_id, dbc_msg.name, "", "报文", "", dbc_msg.name,
                "DBC 中存在该报文，但通信矩阵中没有匹配项。",
            ))

    diffs.extend(self_check_database(matrix, "矩阵"))
    diffs.extend(self_check_database(dbc, "DBC"))
    if vector_rules_enabled:
        diffs.extend(vector_reference_check_database(dbc))
    if semantic_rules:
        diffs.extend(semantic_check_database(matrix, "矩阵", semantic_rules))
        diffs.extend(semantic_check_database(dbc, "DBC", semantic_rules))

    severity_order = {"错误": 0, "警告": 1, "提示": 2}
    return sorted(
        diffs,
        key=lambda d: (
            severity_order.get(d.severity, 9),
            d.can_id if d.can_id is not None else 1 << 30,
            d.message_name.lower(),
            d.signal_name.lower(),
            d.field_name,
        ),
    )


def stats_for(diffs: Sequence[Difference]) -> Dict[str, int]:
    result = {"错误": 0, "警告": 0, "提示": 0, "总计": len(diffs)}
    for diff in diffs:
        result[diff.severity] = result.get(diff.severity, 0) + 1
    return result


def export_csv_report(path: str, diffs: Sequence[Difference], matrix: Optional[Database], dbc: Database) -> None:
    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        writer = csv.writer(f)
        writer.writerow(["严重级别", "类别", "规则ID", "CAN ID（16进制）", "CAN ID（10进制）", "报文名称", "信号名称", "差异项", "矩阵值", "DBC值", "说明"])
        for d in diffs:
            writer.writerow([
                d.severity, d.category, d.rule_id, format_can_id(d.can_id), format_can_id_decimal(d.can_id),
                d.message_name, d.signal_name, d.field_name, d.matrix_value, d.dbc_value, d.description,
            ])


def export_xlsx_report(path: str, diffs: Sequence[Difference], matrix: Optional[Database], dbc: Database) -> None:
    if Workbook is None:
        raise RuntimeError("导出 Excel 报告需要 openpyxl。")

    wb = Workbook()
    ws_summary = wb.active
    ws_summary.title = "检查汇总"
    stats = stats_for(diffs)

    report_title = "DBC无矩阵规则检查报告" if matrix is None else "CAN矩阵-DBC一致性检查报告"
    summary_rows = [
        [report_title, ""],
        ["工具版本", APP_VERSION],
        ["Vector规则基准", f"{VECTOR_MANUAL_TITLE} v{VECTOR_MANUAL_VERSION}"],
        ["检查模式", "无CAN矩阵：DBC结构/Vector手册/项目规则" if matrix is None else "通信矩阵与DBC一致性检查"],
        ["通信矩阵", matrix.source if matrix is not None else "未提供"],
        ["DBC文件", dbc.source],
        ["矩阵报文数", len(matrix.messages) if matrix is not None else 0],
        ["矩阵信号数", sum(len(m.signals) for m in matrix.messages) if matrix is not None else 0],
        ["DBC报文数", len(dbc.messages)],
        ["DBC信号数", sum(len(m.signals) for m in dbc.messages)],
        ["DBC节点数", len(dbc.nodes)],
        ["DBC属性定义数", len(dbc.attribute_definitions)],
        ["DBC显式属性赋值数", len(dbc.attribute_usages)],
        ["错误", stats["错误"]],
        ["警告", stats["警告"]],
        ["提示", stats["提示"]],
        ["差异总计", stats["总计"]],
        ["触发的规则ID数", len({d.rule_id for d in diffs if d.rule_id})],
    ]
    for row in summary_rows:
        ws_summary.append(row)
    ws_summary.merge_cells("A1:B1")
    ws_summary["A1"].font = Font(size=16, bold=True)
    ws_summary["A1"].alignment = Alignment(horizontal="center")
    ws_summary.column_dimensions["A"].width = 22
    ws_summary.column_dimensions["B"].width = 55
    for cell in ws_summary[1]:
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.font = Font(color="FFFFFF", bold=True)

    ws = wb.create_sheet("全部差异")
    headers = ["序号", "严重级别", "类别", "规则ID", "CAN ID（16进制）", "CAN ID（10进制）", "报文名称", "信号名称", "差异项", "矩阵值", "DBC值", "说明"]
    ws.append(headers)
    for idx, d in enumerate(diffs, start=1):
        ws.append([
            idx, d.severity, d.category, d.rule_id, format_can_id(d.can_id), format_can_id_decimal(d.can_id),
            d.message_name, d.signal_name, d.field_name, d.matrix_value, d.dbc_value, d.description,
        ])

    header_fill = PatternFill("solid", fgColor="1F4E78")
    error_fill = PatternFill("solid", fgColor="F4CCCC")
    warning_fill = PatternFill("solid", fgColor="FFF2CC")
    info_fill = PatternFill("solid", fgColor="D9EAF7")
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for row in ws.iter_rows(min_row=2):
        severity = row[1].value
        fill = error_fill if severity == "错误" else warning_fill if severity == "警告" else info_fill
        row[1].fill = fill
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    widths = [8, 10, 18, 20, 15, 15, 24, 28, 18, 28, 28, 55]
    for idx, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    # 按严重等级拆分，便于项目审查。
    for severity in ("错误", "警告", "提示"):
        sheet = wb.create_sheet(severity)
        sheet.append(headers)
        filtered = [d for d in diffs if d.severity == severity]
        for idx, d in enumerate(filtered, start=1):
            sheet.append([
                idx, d.severity, d.category, d.rule_id, format_can_id(d.can_id), format_can_id_decimal(d.can_id),
                d.message_name, d.signal_name, d.field_name, d.matrix_value, d.dbc_value, d.description,
            ])
        for cell in sheet[1]:
            cell.fill = header_fill
            cell.font = Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(horizontal="center")
        for col_idx, width in enumerate(widths, start=1):
            sheet.column_dimensions[get_column_letter(col_idx)].width = width
        sheet.freeze_panes = "A2"
        if sheet.max_row >= 1:
            sheet.auto_filter.ref = sheet.dimensions
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)

    ws_attr = wb.create_sheet("DBC属性概览")
    attr_headers = [
        "CAN ID（16进制）", "CAN ID（10进制）", "报文名称", "发送节点", "DLC", "总线格式", "ID格式", "GenMsgSendType", "显式赋值",
        "GenMsgCycleTime", "显式赋值", "VFrameFormat", "NmMessage", "MsgType", "DiagRequest", "DiagResponse",
        "DiagState", "CANFD_BRS", "识别类型",
    ]
    ws_attr.append(attr_headers)
    for msg in dbc.messages:
        _send_n, send_v, send_e = effective_attribute_any(msg, dbc, ("GenMsgSendType", "VFrameSendType", "SendType"))
        _cycle_n, cycle_v, cycle_e = effective_attribute_any(msg, dbc, ("GenMsgCycleTime", "CycleTime"))
        _vf_n, vf_v, _ = effective_attribute_any(msg, dbc, ("VFrameFormat",))
        _nm_n, nm_v, _ = effective_attribute_any(msg, dbc, ("NmMessage", "NMAsrMessage", "NmhMessage"))
        _type_n, type_v, _ = effective_attribute_any(msg, dbc, ("MsgType",))
        diag_req = effective_attribute(msg, dbc, "DiagRequest")[0]
        diag_resp = effective_attribute(msg, dbc, "DiagResponse")[0]
        diag_state = effective_attribute(msg, dbc, "DiagState")[0]
        brs = effective_attribute(msg, dbc, "CANFD_BRS")[0]
        kinds = message_kind_flags(msg, dbc)
        kind_text = ",".join(name.upper() for name, enabled in kinds.items() if enabled)
        ws_attr.append([
            format_can_id(msg.can_id), format_can_id_decimal(msg.can_id), msg.name, msg.sender, msg.dlc, msg.bus_format, msg.frame_format,
            value_display(send_v), "是" if send_e else "否", value_display(cycle_v), "是" if cycle_e else "否",
            value_display(vf_v), value_display(nm_v), value_display(type_v), value_display(diag_req),
            value_display(diag_resp), value_display(diag_state), value_display(brs), kind_text,
        ])
    for cell in ws_attr[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for idx, width in enumerate([15, 15, 28, 18, 8, 12, 10, 20, 10, 18, 10, 20, 14, 18, 14, 14, 14, 12, 16], start=1):
        ws_attr.column_dimensions[get_column_letter(idx)].width = width
    ws_attr.freeze_panes = "A2"
    ws_attr.auto_filter.ref = ws_attr.dimensions

    ws_vector = wb.create_sheet("Vector手册规则说明")
    vector_rows = [
        ["基准文档", f"{VECTOR_MANUAL_TITLE}, Version {VECTOR_MANUAL_VERSION}, Released"],
        ["规则层级", "内置Vector手册基线 + can_rules.json项目培训规则"],
        ["通用属性", "BusType、VFrameFormat、波特率/采样点、属性对象类型/类型/范围/枚举顺序"],
        ["COM", "ILSupport、报文/信号发送类型、周期、重复、Delay、StartValue、TimeoutTime_<Ecu>"],
        ["E2E/SecOC", "E2E Profile/DataId/DataLength基础一致性；SecOC split与长度边界关系"],
        ["NM", "AUTOSAR NM/OSEK-NM网络、节点、报文、ID范围、周期/offset、CBV/SNI"],
        ["CanTp/DCM", "TpTxIndex、诊断角色、DiagConnection配对、DiagFdOnly一致性"],
        ["J1939", "ProtocolType、扩展ID、动态DLC末端信号"],
        ["XCP/CDD", "非AUTOSAR层属性与完整载荷Rx信号"],
        ["Update Bit/Invalid", "<X>_UB命名/长度/发送类型；SNA无效值"],
        ["边界", "手册自身说明其为属性总览；具体E2E Profile、BSW组件细节仍应以对应MICROSAR技术参考和项目规范为准。"],
    ]
    for row in vector_rows:
        ws_vector.append(row)
    ws_vector.column_dimensions["A"].width = 22
    ws_vector.column_dimensions["B"].width = 100
    for row in ws_vector.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for cell in ws_vector[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)

    wb.save(path)


def wrap_tree_cell(value: Any, width: int, max_lines: int = 3) -> str:
    """为Treeview单元格插入换行；完整内容仍保存在下方详情区。"""
    raw = value_display(value)
    if not raw:
        return ""
    raw = raw.replace("\r\n", "\n").replace("\r", "\n")
    lines: List[str] = []
    for paragraph in raw.split("\n"):
        wrapped = textwrap.wrap(
            paragraph, width=max(6, width), break_long_words=True,
            break_on_hyphens=False, replace_whitespace=False, drop_whitespace=True,
        ) or [""]
        lines.extend(wrapped)
    if len(lines) > max_lines:
        lines = lines[:max_lines]
        if lines[-1]:
            lines[-1] = lines[-1][:-1] + "…" if len(lines[-1]) > 1 else "…"
    return "\n".join(lines)


def difference_detail_text(diff: Difference) -> str:
    """生成选中检查项的完整、可复制详情。"""
    pairs = [
        ("级别", diff.severity),
        ("类别", diff.category),
        ("规则ID", diff.rule_id),
        ("CAN ID（16进制）", format_can_id(diff.can_id)),
        ("CAN ID（10进制）", format_can_id_decimal(diff.can_id)),
        ("报文名称", diff.message_name),
        ("信号名称", diff.signal_name),
        ("差异项", diff.field_name),
        ("矩阵值", diff.matrix_value),
        ("DBC值", diff.dbc_value),
        ("说明", diff.description),
    ]
    return "\n".join(f"{name}：{value_display(value)}" for name, value in pairs if value_display(value))


class CheckerApp:
    def __init__(self, root: tk.Tk) -> None:
        self.root = root
        self.root.title(f"{APP_NAME} v{APP_VERSION}")
        self.root.geometry("1450x900")
        self.root.minsize(1100, 720)

        self.matrix_path = tk.StringVar()
        self.dbc_path = tk.StringVar()
        self.rename_config_path = tk.StringVar()
        default_rule_file = ensure_default_rules_file()
        self.rules_path = tk.StringVar(value=str(default_rule_file))
        self.semantic_rules_enabled = tk.BooleanVar(value=True)
        self.vector_rules_enabled = tk.BooleanVar(value=True)
        self.dbc_only_mode = tk.BooleanVar(value=False)
        self.active_rules_count = 0
        self.sheet_name = tk.StringVar()
        self.e2e_table_path = tk.StringVar()
        self.e2e_sheet_name = tk.StringVar()
        self.status_var = tk.StringVar(value="请选择通信矩阵和 DBC 文件；E2E ID表可选导入。没有矩阵时可勾选“无CAN矩阵”。")
        self.summary_var = tk.StringVar(value="错误：0    警告：0    提示：0    总计：0")
        self.filter_var = tk.StringVar(value="全部")
        self.type_filter_var = tk.StringVar(value="全部")
        self.run_state_var = tk.StringVar(value="就绪")

        self.matrix_db: Optional[Database] = None
        self.dbc_db: Optional[Database] = None
        self.differences: List[Difference] = []
        self.mapping_info: Dict[str, str] = {}
        self.header_row = 0
        self.e2e_table_entries: List[E2EIdEntry] = []
        self.e2e_mapping_info: Dict[str, str] = {}
        self.e2e_header_row = 0
        self.e2e_actual_sheet = ""
        self._ui_queue: "queue.Queue[Tuple[Callable[..., Any], Tuple[Any, ...]]]" = queue.Queue()
        self._tree_item_details: Dict[str, str] = {}
        # 结果表中“单条一键修改E2E”对应的高置信度修复项。
        self._tree_item_e2e_fixes: Dict[str, Tuple[Message, int, str]] = {}
        # 非E2E的一键修复项；具体动作由规则ID和当前检查结果决定。
        self._tree_item_autofix_diffs: Dict[str, Difference] = {}

        self._build_ui()
        self.root.after(50, self._poll_ui_queue)

    def _build_ui(self) -> None:
        style = ttk.Style()
        try:
            style.theme_use("vista" if sys.platform.startswith("win") else "clam")
        except tk.TclError:
            pass
        # Treeview本身不支持每行自动高度；统一提高行高并在插入时手动换行。
        style.configure("Result.Treeview", rowheight=64, font=("Microsoft YaHei UI", 9))
        style.configure("Result.Treeview.Heading", font=("Microsoft YaHei UI", 9, "bold"))

        top = ttk.Frame(self.root, padding=12)
        top.pack(fill="x")
        top.columnconfigure(1, weight=1)

        self.matrix_label = ttk.Label(top, text="CAN通信矩阵：")
        self.matrix_label.grid(row=0, column=0, sticky="w", pady=4)
        self.matrix_entry = ttk.Entry(top, textvariable=self.matrix_path)
        self.matrix_entry.grid(row=0, column=1, sticky="ew", padx=6, pady=4)
        self.matrix_button = ttk.Button(top, text="选择文件", command=self.choose_matrix)
        self.matrix_button.grid(row=0, column=2, padx=4, pady=4)

        self.sheet_label = ttk.Label(top, text="工作表：")
        self.sheet_label.grid(row=0, column=3, sticky="e", padx=(15, 2))
        self.sheet_combo = ttk.Combobox(top, textvariable=self.sheet_name, state="readonly", width=24)
        self.sheet_combo.grid(row=0, column=4, padx=4, pady=4)

        ttk.Label(top, text="DBC文件：").grid(row=1, column=0, sticky="w", pady=4)
        ttk.Entry(top, textvariable=self.dbc_path).grid(row=1, column=1, sticky="ew", padx=6, pady=4)
        ttk.Button(top, text="选择文件", command=self.choose_dbc).grid(row=1, column=2, padx=4, pady=4)

        button_frame = ttk.Frame(top)
        button_frame.grid(row=1, column=3, columnspan=2, sticky="e")
        self.check_button = ttk.Button(button_frame, text="开始检查", command=self.start_check)
        self.check_button.pack(side="left", padx=4)
        self.export_button = ttk.Button(button_frame, text="导出报告", command=self.export_report, state="disabled")
        self.export_button.pack(side="left", padx=4)
        self.fix_e2e_button = ttk.Button(
            button_frame, text="全部一键修改E2E", command=self.auto_fix_all_e2e, state="disabled"
        )
        self.fix_e2e_button.pack(side="left", padx=4)
        ttk.Button(button_frame, text="命名设置", command=self.open_naming_tool).pack(side="left", padx=4)
        ttk.Button(button_frame, text="节点补全", command=self.open_node_completion).pack(side="left", padx=4)
        ttk.Button(button_frame, text="属性修正", command=self.open_attribute_repair).pack(side="left", padx=4)
        ttk.Button(button_frame, text="查看列映射", command=self.show_mapping).pack(side="left", padx=4)

        ttk.Label(top, text="E2E ID表：").grid(row=2, column=0, sticky="w", pady=4)
        self.e2e_table_entry = ttk.Entry(top, textvariable=self.e2e_table_path)
        self.e2e_table_entry.grid(row=2, column=1, sticky="ew", padx=6, pady=4)
        ttk.Button(top, text="选择E2E表", command=self.choose_e2e_table).grid(row=2, column=2, padx=4, pady=4)
        e2e_sheet_frame = ttk.Frame(top)
        e2e_sheet_frame.grid(row=2, column=3, columnspan=2, sticky="e")
        ttk.Label(e2e_sheet_frame, text="E2E工作表：").pack(side="left", padx=(4, 2))
        self.e2e_sheet_combo = ttk.Combobox(e2e_sheet_frame, textvariable=self.e2e_sheet_name, state="readonly", width=24)
        self.e2e_sheet_combo.pack(side="left", padx=4)
        ttk.Button(e2e_sheet_frame, text="查看E2E映射", command=self.show_e2e_mapping).pack(side="left", padx=4)

        ttk.Label(top, text="语义规则：").grid(row=3, column=0, sticky="w", pady=4)
        ttk.Entry(top, textvariable=self.rules_path).grid(row=3, column=1, sticky="ew", padx=6, pady=4)
        ttk.Button(top, text="选择规则", command=self.choose_rules).grid(row=3, column=2, padx=4, pady=4)
        rule_frame = ttk.Frame(top)
        rule_frame.grid(row=3, column=3, columnspan=2, sticky="e")
        ttk.Checkbutton(
            rule_frame, text="无CAN矩阵（仅DBC规则检查）", variable=self.dbc_only_mode,
            command=self._toggle_check_mode,
        ).pack(side="left", padx=4)
        ttk.Checkbutton(
            rule_frame, text=f"启用Vector手册v{VECTOR_MANUAL_VERSION}规则", variable=self.vector_rules_enabled
        ).pack(side="left", padx=4)
        ttk.Checkbutton(rule_frame, text="启用项目培训规则", variable=self.semantic_rules_enabled).pack(side="left", padx=4)
        ttk.Button(rule_frame, text="查看规则", command=self.show_rules).pack(side="left", padx=4)

        info = ttk.Frame(self.root, padding=(12, 0, 12, 8))
        info.pack(fill="x")
        ttk.Label(info, textvariable=self.summary_var, font=("Microsoft YaHei UI", 11, "bold")).pack(side="left")
        ttk.Label(info, text="级别：").pack(side="left", padx=(28, 4))
        filter_combo = ttk.Combobox(
            info, textvariable=self.filter_var, values=["全部", "错误", "警告", "提示"], state="readonly", width=8
        )
        filter_combo.pack(side="left")
        filter_combo.bind("<<ComboboxSelected>>", lambda _e: self.refresh_tree())
        ttk.Label(info, text="类型：").pack(side="left", padx=(16, 4))
        type_filter_combo = ttk.Combobox(
            info, textvariable=self.type_filter_var,
            values=["全部", "DBC自身/规则", "矩阵差异", "E2E全部", "E2E DataID表", "E2E DataLength", "NM", "UDS/诊断", "XCP"],
            state="readonly", width=15,
        )
        type_filter_combo.pack(side="left")
        type_filter_combo.bind("<<ComboboxSelected>>", lambda _e: self.refresh_tree())

        columns = ("severity", "category", "rule_id", "can_id", "can_id_dec", "message", "signal", "field", "matrix", "dbc", "action", "description")

        # 结果表格和完整详情区放进可拖动的垂直分隔窗格。
        # 用户可拖动中间横向分隔条，自由调整上下两个区域的高度。
        self.result_pane = tk.PanedWindow(
            self.root,
            orient=tk.VERTICAL,
            sashwidth=8,
            sashrelief=tk.RAISED,
            showhandle=True,
            handlesize=12,
            handlepad=6,
            borderwidth=0,
            relief=tk.FLAT,
            background="#D0D0D0",
        )
        self.result_pane.pack(fill="both", expand=True, padx=12, pady=(0, 6))

        tree_frame = ttk.Frame(self.result_pane)
        self.tree = ttk.Treeview(tree_frame, columns=columns, show="headings", style="Result.Treeview")
        headings = {
            "severity": "级别", "category": "类别", "rule_id": "规则ID",
            "can_id": "CAN ID（16进制）", "can_id_dec": "CAN ID（10进制）", "message": "报文名称",
            "signal": "信号名称", "field": "差异项", "matrix": "矩阵值", "dbc": "DBC值",
            "action": "操作", "description": "说明",
        }
        widths = {
            "severity": 65, "category": 120, "rule_id": 155, "can_id": 115, "can_id_dec": 115,
            "message": 150, "signal": 170, "field": 105, "matrix": 170, "dbc": 170,
            "action": 100, "description": 340,
        }
        for col in columns:
            self.tree.heading(col, text=headings[col])
            self.tree.column(col, width=widths[col], minwidth=55, anchor="w")

        yscroll = ttk.Scrollbar(tree_frame, orient="vertical", command=self.tree.yview)
        xscroll = ttk.Scrollbar(tree_frame, orient="horizontal", command=self.tree.xview)
        self.tree.configure(yscrollcommand=yscroll.set, xscrollcommand=xscroll.set)
        self.tree.grid(row=0, column=0, sticky="nsew")
        yscroll.grid(row=0, column=1, sticky="ns")
        xscroll.grid(row=1, column=0, sticky="ew")
        tree_frame.rowconfigure(0, weight=1)
        tree_frame.columnconfigure(0, weight=1)
        self.tree.bind("<<TreeviewSelect>>", self._show_selected_detail)
        self.tree.bind("<ButtonRelease-1>", self._on_tree_action_click, add="+")

        self.tree.tag_configure("错误", background="#FCE8E6")
        self.tree.tag_configure("警告", background="#FFF8E1")
        self.tree.tag_configure("提示", background="#EAF3FA")
        self.tree.tag_configure("通过", background="#E8F5E9")

        detail_frame = ttk.LabelFrame(
            self.result_pane,
            text="选中项完整详情（拖动上方分隔条可调整此框大小）",
            padding=(8, 5),
        )
        self.detail_text = tk.Text(
            detail_frame, height=5, wrap="word", relief="flat", borderwidth=0,
            font=("Microsoft YaHei UI", 9), background="#FAFAFA",
        )
        detail_scroll = ttk.Scrollbar(detail_frame, orient="vertical", command=self.detail_text.yview)
        self.detail_text.configure(yscrollcommand=detail_scroll.set, state="disabled")
        self.detail_text.pack(side="left", fill="both", expand=True)
        detail_scroll.pack(side="right", fill="y")

        # minsize 防止任一区域被拖到完全不可见；stretch 允许随窗口共同缩放。
        self.result_pane.add(tree_frame, minsize=220, stretch="always")
        self.result_pane.add(detail_frame, minsize=90, stretch="always")
        self.root.after(120, self._set_initial_result_split)

        status_bar = ttk.Frame(self.root, padding=(12, 4, 12, 8))
        status_bar.pack(fill="x")
        self.progress = ttk.Progressbar(status_bar, mode="determinate", maximum=100, value=0, length=160)
        self.progress.pack(side="left", padx=(0, 10))
        ttk.Label(status_bar, textvariable=self.run_state_var, width=8, anchor="center").pack(side="left", padx=(0, 10))
        ttk.Label(status_bar, textvariable=self.status_var).pack(side="left", fill="x", expand=True)

    def _set_initial_result_split(self) -> None:
        """设置结果表与详情区的初始比例；之后用户可直接拖动分隔条。"""
        try:
            total_height = max(self.result_pane.winfo_height(), 500)
            # 默认约 78% 给结果表，22% 给详情区。
            self.result_pane.sash_place(0, 0, int(total_height * 0.78))
        except (tk.TclError, IndexError):
            pass

    def _post_ui(self, func: Callable[..., Any], *args: Any) -> None:
        self._ui_queue.put((func, args))

    def _poll_ui_queue(self) -> None:
        try:
            while True:
                func, args = self._ui_queue.get_nowait()
                func(*args)
        except queue.Empty:
            pass
        try:
            self.root.after(50, self._poll_ui_queue)
        except tk.TclError:
            pass

    def set_status(self, message: str) -> None:
        self._post_ui(self.status_var.set, message)

    def _toggle_check_mode(self) -> None:
        dbc_only = bool(self.dbc_only_mode.get())
        state = "disabled" if dbc_only else "normal"
        self.matrix_entry.configure(state=state)
        self.matrix_button.configure(state=state)
        self.sheet_combo.configure(state="disabled" if dbc_only else "readonly")
        if dbc_only:
            self.status_var.set("无CAN矩阵模式：将只检查DBC语法、位布局、属性、节点和保守经验规则。")
        else:
            self.status_var.set("矩阵对比模式：请选择通信矩阵、工作表和DBC文件。")

    def choose_matrix(self) -> None:
        path = filedialog.askopenfilename(
            title="选择 CAN 通信矩阵",
            filetypes=[("支持的矩阵文件", "*.xlsx *.xlsm *.csv"), ("Excel", "*.xlsx *.xlsm"), ("CSV", "*.csv")],
        )
        if not path:
            return
        self.matrix_path.set(path)
        try:
            sheets, best_sheet, details = inspect_matrix_sheets(path)
            self.sheet_combo["values"] = sheets
            self.sheet_name.set(best_sheet or (sheets[0] if sheets else ""))
            if best_sheet:
                self.status_var.set(
                    f"已选择通信矩阵：{Path(path).name}；自动选择工作表：{best_sheet}"
                )
            else:
                self.status_var.set(f"所选文件未识别出通信矩阵工作表：{Path(path).name}")
                short_details = "\n".join(
                    f"{name}：{reason.splitlines()[0]}" for name, reason in details.items()
                )
                messagebox.showwarning(
                    "文件不像通信矩阵",
                    "没有工作表包含可识别的报文和信号表头。\n"
                    "你可能选择了导出的检查报告，而不是原始 CAN 通信矩阵。\n\n"
                    f"工作表检测结果：\n{short_details}",
                )
        except Exception as exc:
            messagebox.showerror("读取失败", str(exc))

    def choose_e2e_table(self) -> None:
        path = filedialog.askopenfilename(
            title="选择客户E2E ID表",
            filetypes=[("E2E ID表", "*.xlsx *.xlsm *.csv"), ("Excel", "*.xlsx *.xlsm"), ("CSV", "*.csv")],
        )
        if not path:
            return
        self.e2e_table_path.set(path)
        try:
            sheets, best_sheet, details = inspect_e2e_id_sheets(path)
            recognized = [name for name in sheets if details.get(name, "").startswith("识别到")]
            combo_values = ([E2E_ALL_SHEETS_LABEL] + sheets) if len(recognized) > 1 else sheets
            self.e2e_sheet_combo["values"] = combo_values
            default_sheet = E2E_ALL_SHEETS_LABEL if len(recognized) > 1 else (best_sheet or (sheets[0] if sheets else ""))
            self.e2e_sheet_name.set(default_sheet)
            if recognized:
                if len(recognized) > 1:
                    self.status_var.set(
                        f"已选择E2E ID表：{Path(path).name}；检测到{len(recognized)}个可识别工作表，默认汇总导入。"
                    )
                else:
                    self.status_var.set(f"已选择E2E ID表：{Path(path).name}；自动选择工作表：{default_sheet}")
            else:
                short = "\n".join(f"{k}：{v}" for k, v in details.items())
                messagebox.showwarning(
                    "未识别E2E ID表",
                    "未找到类似 TX | DATA ID | Rx | DATAID 的布局。\n\n" + short,
                )
        except Exception as exc:
            messagebox.showerror("读取E2E ID表失败", str(exc))

    def show_e2e_mapping(self) -> None:
        if not self.e2e_mapping_info or not self.e2e_table_entries or self.dbc_db is None:
            messagebox.showinfo(
                "E2E DataID核对",
                "尚未完成E2E ID表核对。选择E2E ID表并执行“开始检查”后，可在这里直接查看客户Data ID与DBC Data ID。",
            )
            return

        rows = build_e2e_mapping_rows(self.e2e_table_entries, self.dbc_db)
        win = tk.Toplevel(self.root)
        win.title("E2E DataID核对结果")
        win.geometry("1500x680")
        win.minsize(1000, 480)
        win.transient(self.root)

        top = ttk.Frame(win, padding=(10, 10, 10, 6))
        top.pack(fill="x")
        mapping_text = "；".join(f"{k}→{v}" for k, v in self.e2e_mapping_info.items())
        ttk.Label(
            top,
            text=(
                f"工作表：{self.e2e_actual_sheet}    "
                f"表头行：{self.e2e_header_row if self.e2e_header_row > 0 else '各表自动识别'}    "
                f"读取：{len(self.e2e_table_entries)}条    {mapping_text}"
            ),
        ).pack(anchor="w")

        ok_count = sum(1 for r in rows if r.status.startswith("一致"))
        mismatch_count = sum(1 for r in rows if r.status == "不一致")
        missing_count = sum(1 for r in rows if r.status in {"DBC未找到", "DBC未配置", "DBC多处匹配", "DBC多值冲突", "客户表冲突"})
        skipped_count = sum(1 for r in rows if r.status == "客户未给ID")
        ttk.Label(
            top,
            text=f"一致：{ok_count}    不一致：{mismatch_count}    缺失/冲突：{missing_count}    客户未给ID：{skipped_count}    总计：{len(rows)}",
            font=("Microsoft YaHei UI", 10, "bold"),
        ).pack(anchor="w", pady=(6, 0))

        body = ttk.Frame(win, padding=(10, 0, 10, 10))
        body.pack(fill="both", expand=True)
        cols = (
            "dir", "cust_group", "dbc_group", "can_hex", "can_dec", "msg",
            "cust_hex", "cust_dec", "dbc_hex", "dbc_dec", "status", "note",
        )
        tree = ttk.Treeview(body, columns=cols, show="headings")
        headings = {
            "dir": "方向", "cust_group": "客户SignalGroup", "dbc_group": "DBC SignalGroup",
            "can_hex": "CAN ID(16进制)", "can_dec": "CAN ID(10进制)", "msg": "报文名称",
            "cust_hex": "客户DataID(16进制)", "cust_dec": "客户DataID(10进制)",
            "dbc_hex": "DBC DataID(16进制)", "dbc_dec": "DBC DataID(10进制)",
            "status": "结果", "note": "说明",
        }
        widths = {
            "dir": 60, "cust_group": 190, "dbc_group": 190, "can_hex": 115, "can_dec": 115,
            "msg": 210, "cust_hex": 135, "cust_dec": 135, "dbc_hex": 145, "dbc_dec": 145,
            "status": 115, "note": 330,
        }
        for col in cols:
            tree.heading(col, text=headings[col])
            tree.column(col, width=widths[col], minwidth=55, anchor="w")

        ybar = ttk.Scrollbar(body, orient="vertical", command=tree.yview)
        xbar = ttk.Scrollbar(body, orient="horizontal", command=tree.xview)
        tree.configure(yscrollcommand=ybar.set, xscrollcommand=xbar.set)
        tree.grid(row=0, column=0, sticky="nsew")
        ybar.grid(row=0, column=1, sticky="ns")
        xbar.grid(row=1, column=0, sticky="ew")
        body.rowconfigure(0, weight=1)
        body.columnconfigure(0, weight=1)

        tree.tag_configure("一致", background="#EAF6EA")
        tree.tag_configure("异常", background="#FCEAEA")
        tree.tag_configure("提示", background="#FFF8DC")

        for row in rows:
            dbc_hex = " / ".join(format_can_id(v) for v in row.dbc_data_ids) if row.dbc_data_ids else ""
            dbc_dec = " / ".join(str(v) for v in row.dbc_data_ids) if row.dbc_data_ids else ""
            if row.status.startswith("一致"):
                tag = "一致"
            elif row.status == "客户未给ID":
                tag = "提示"
            else:
                tag = "异常"
            tree.insert(
                "", "end",
                values=(
                    row.direction, row.customer_group, row.dbc_group, format_can_id(row.can_id),
                    format_can_id_decimal(row.can_id), row.message_name,
                    format_can_id(row.customer_data_id),
                    "" if row.customer_data_id is None else str(row.customer_data_id),
                    dbc_hex, dbc_dec, row.status, row.note,
                ),
                tags=(tag,),
            )

        bottom = ttk.Frame(win, padding=(10, 0, 10, 10))
        bottom.pack(fill="x")
        ttk.Label(
            bottom,
            text=(
                "说明：DataID是E2E保护对象标识；这里专门核对客户E2E表，不等同于E2EDataLength检查。"
                " 客户表未列出的DBC E2E组默认不判错误；仅对客户明确给出的项做硬核对。"
            ),
        ).pack(side="left")
        ttk.Button(bottom, text="关闭", command=win.destroy).pack(side="right")

    def open_naming_tool(self) -> None:
        """打开真实 CAN ID 命名预览/配置窗口。"""
        dbc_path = self.dbc_path.get().strip()
        if not dbc_path or not os.path.isfile(dbc_path):
            dbc_path = filedialog.askopenfilename(title="选择用于命名的 DBC 文件", filetypes=[("DBC文件", "*.dbc"), ("所有文件", "*.*")])
            if not dbc_path:
                return
            self.dbc_path.set(dbc_path)

        win = tk.Toplevel(self.root)
        win.title("DBC节点/报文/信号命名设置")
        win.geometry("1320x720")
        win.minsize(1000, 560)
        config = default_rename_config()
        plan = None
        row_items: Dict[str, Any] = {}

        top = ttk.Frame(win, padding=10)
        top.pack(fill="x")
        top.columnconfigure(1, weight=1)
        ttk.Label(top, text="DBC文件：").grid(row=0, column=0, sticky="w", pady=3)
        ttk.Label(top, text=dbc_path).grid(row=0, column=1, columnspan=3, sticky="w", pady=3)
        ttk.Label(top, text="默认命名标识：").grid(row=1, column=0, sticky="w", pady=3)
        default_var = tk.StringVar(value="can1")
        ttk.Entry(top, textvariable=default_var, width=18).grid(row=1, column=1, sticky="w", pady=3)
        messages_var = tk.BooleanVar(value=True)
        signals_var = tk.BooleanVar(value=True)
        ttk.Checkbutton(top, text="处理报文", variable=messages_var).grid(row=1, column=2, sticky="w", padx=8)
        ttk.Checkbutton(top, text="处理信号", variable=signals_var).grid(row=1, column=3, sticky="w", padx=8)
        ttk.Label(top, text="配置JSON（可留空）：").grid(row=2, column=0, sticky="w", pady=3)
        config_var = tk.StringVar(value=self.rename_config_path.get().strip())
        ttk.Entry(top, textvariable=config_var).grid(row=2, column=1, sticky="ew", pady=3)

        controls = ttk.Frame(top)
        controls.grid(row=2, column=2, columnspan=2, sticky="e")
        status_var = tk.StringVar(value="只需填命名标识，然后点“保存为新DBC”；配置JSON可不填。")

        body = ttk.Frame(win, padding=(10, 0, 10, 6))
        body.pack(fill="both", expand=True)
        columns = ("type", "frame", "can", "old", "new", "identifier", "status", "reason")
        tree = ttk.Treeview(body, columns=columns, show="headings", selectmode="extended")
        headings = {"type": "对象", "frame": "帧身份", "can": "真实CAN ID", "old": "原名称", "new": "新名称", "identifier": "标识", "status": "状态", "reason": "依据/阻塞原因"}
        widths = {"type": 70, "frame": 75, "can": 100, "old": 220, "new": 285, "identifier": 90, "status": 80, "reason": 360}
        for col in columns:
            tree.heading(col, text=headings[col])
            tree.column(col, width=widths[col], minwidth=55, anchor="w")
        ybar = ttk.Scrollbar(body, orient="vertical", command=tree.yview)
        xbar = ttk.Scrollbar(body, orient="horizontal", command=tree.xview)
        tree.configure(yscrollcommand=ybar.set, xscrollcommand=xbar.set)
        tree.grid(row=0, column=0, sticky="nsew")
        ybar.grid(row=0, column=1, sticky="ns")
        xbar.grid(row=1, column=0, sticky="ew")
        body.rowconfigure(0, weight=1)
        body.columnconfigure(0, weight=1)

        override = ttk.Frame(win, padding=(10, 0, 10, 6))
        override.pack(fill="x")
        ttk.Label(override, text="选中对象覆盖标识：").pack(side="left")
        override_var = tk.StringVar()
        ttk.Entry(override, textvariable=override_var, width=18).pack(side="left", padx=5)

        def render() -> bool:
            nonlocal plan, config
            try:
                config["default_identifier"] = default_var.get().strip()
                config["targets"] = {"messages": bool(messages_var.get()), "signals": bool(signals_var.get())}
                plan = build_rename_plan(dbc_path, config)
            except Exception as exc:
                status_var.set(str(exc))
                messagebox.showerror("命名配置错误", str(exc), parent=win)
                return False
            for item in tree.get_children():
                tree.delete(item)
            row_items.clear()
            for index, item in enumerate(plan.items):
                iid = str(index)
                row_items[iid] = item
                tree.insert("", "end", iid=iid, values=(item.object_type, item.frame_format, format_can_id(item.can_id), item.old_name, item.new_name, item.identifier, item.status, item.reason))
            status_var.set(f"预览 {len(plan.items)} 项；可执行 {len(plan.executable_items)} 项；阻塞 {len(plan.blocked_items)} 项。真实 ID 来自 DBC BO_。")
            return True

        def set_override() -> None:
            value = override_var.get().strip()
            selected = [row_items[iid] for iid in tree.selection() if iid in row_items]
            if not selected:
                messagebox.showinfo("未选择对象", "请在预览表中选择一个或多个报文/信号。", parent=win)
                return
            for item in selected:
                overrides = config["message_overrides"] if item.object_type == "报文" else config["signal_overrides"]
                if value:
                    overrides[item.object_key] = value
                else:
                    overrides.pop(item.object_key, None)
            render()

        def load_config() -> None:
            nonlocal config
            path = filedialog.askopenfilename(parent=win, title="加载命名JSON配置", filetypes=[("JSON配置", "*.json"), ("所有文件", "*.*")])
            if not path:
                return
            try:
                config = load_rename_config(path)
            except Exception as exc:
                messagebox.showerror("加载配置失败", str(exc), parent=win)
                return
            config_var.set(path)
            self.rename_config_path.set(path)
            default_var.set(config.get("default_identifier", "can1"))
            messages_var.set(bool(config.get("targets", {}).get("messages", True)))
            signals_var.set(bool(config.get("targets", {}).get("signals", True)))
            render()

        def save_config() -> bool:
            config["default_identifier"] = default_var.get().strip()
            config["targets"] = {"messages": bool(messages_var.get()), "signals": bool(signals_var.get())}
            path = config_var.get().strip()
            if not path:
                path = filedialog.asksaveasfilename(parent=win, title="保存命名JSON配置", defaultextension=".json", filetypes=[("JSON配置", "*.json")])
            if not path:
                return False
            try:
                save_rename_config(path, config)
            except Exception as exc:
                messagebox.showerror("保存配置失败", str(exc), parent=win)
                return False
            config_var.set(path)
            self.rename_config_path.set(path)
            status_var.set(f"配置已保存：{path}")
            return True

        def apply() -> None:
            nonlocal config
            if not render() or plan is None:
                return
            if plan.blocked_items:
                messagebox.showwarning("存在阻塞项", "预览中存在来源不明、冲突或不支持项，已阻止写回。请先处理阻塞项。", parent=win)
                return
            output = filedialog.asksaveasfilename(parent=win, title="另存重命名后的DBC", initialdir=str(Path(dbc_path).parent), initialfile=f"{Path(dbc_path).stem}_renamed{Path(dbc_path).suffix}", defaultextension=".dbc", filetypes=[("DBC文件", "*.dbc")])
            if not output:
                return
            try:
                saved_path, updated_config = apply_rename_plan(plan, config, output)
                config = updated_config
                config_path = config_var.get().strip() or str(Path(saved_path).with_suffix(".rename.json"))
                save_rename_config(config_path, config)
                parse_dbc(saved_path)
            except Exception as exc:
                messagebox.showerror("命名写回失败", str(exc), parent=win)
                return
            self.rename_config_path.set(config_path)
            status_var.set(f"已另存并重新解析：{saved_path}；映射配置：{config_path}")
            messagebox.showinfo("命名完成", f"已保存：\n{saved_path}\n\n映射配置：\n{config_path}\n\n原 DBC 未被覆盖。", parent=win)

        tree.bind("<<TreeviewSelect>>", lambda _event: override_var.set(row_items[tree.selection()[0]].identifier if tree.selection() and tree.selection()[0] in row_items else ""))
        ttk.Button(override, text="设置选中对象", command=set_override).pack(side="left", padx=4)
        ttk.Button(override, text="清除选中覆盖", command=lambda: (override_var.set(""), set_override())).pack(side="left", padx=4)
        ttk.Label(override, textvariable=status_var).pack(side="left", padx=12)

        bottom = ttk.Frame(win, padding=(10, 0, 10, 10))
        bottom.pack(fill="x")
        ttk.Button(bottom, text="刷新预览", command=render).pack(side="left", padx=3)
        ttk.Button(bottom, text="加载JSON", command=load_config).pack(side="left", padx=3)
        ttk.Button(bottom, text="保存JSON", command=save_config).pack(side="left", padx=3)
        ttk.Button(bottom, text="保存为新DBC", command=apply).pack(side="right", padx=3)
        ttk.Button(bottom, text="关闭", command=win.destroy).pack(side="right", padx=3)
        render()

    def open_node_completion(self) -> None:
        """只补入 DBC 中已明确引用但未声明的真实节点。"""
        dbc_path = self.dbc_path.get().strip()
        if not dbc_path or not os.path.isfile(dbc_path):
            dbc_path = filedialog.askopenfilename(title="选择用于节点补全的 DBC 文件", filetypes=[("DBC文件", "*.dbc"), ("所有文件", "*.*")])
            if not dbc_path:
                return
            self.dbc_path.set(dbc_path)
        try:
            plan = build_node_completion_plan(dbc_path)
        except Exception as exc:
            messagebox.showerror("节点扫描失败", str(exc))
            return
        if not plan.missing_nodes:
            messagebox.showinfo("节点补全", "未发现已明确引用但未声明在 BU_ 中的真实节点。Vector__XXX 不会被当作真实节点补入。")
            return
        ok = messagebox.askyesno(
            "节点补全预览",
            f"发现以下真实节点未在 BU_ 中声明：\n\n{', '.join(plan.missing_nodes)}\n\n"
            "将只补入这些明确引用的节点，不会虚构接收者，也不会替换已有收发关系。是否另存为新 DBC？",
        )
        if not ok:
            return
        output = filedialog.asksaveasfilename(
            title="另存节点补全后的 DBC",
            initialdir=str(Path(dbc_path).parent),
            initialfile=f"{Path(dbc_path).stem}_nodes_completed{Path(dbc_path).suffix}",
            defaultextension=".dbc",
            filetypes=[("DBC文件", "*.dbc")],
        )
        if not output:
            return
        try:
            saved_path, nodes = apply_node_completion(plan, output)
            parse_dbc(saved_path)
        except Exception as exc:
            messagebox.showerror("节点补全失败", str(exc))
            return
        self.status_var.set(f"节点补全完成：已另存 {Path(saved_path).name}，补入 {len(nodes)} 个节点。")
        messagebox.showinfo("节点补全完成", f"已保存：\n{saved_path}\n\n补入节点：{', '.join(nodes)}")

    def open_attribute_repair(self) -> None:
        """预览并在用户确认后移除没有 BA_DEF_ 的属性赋值。"""
        dbc_path = self.dbc_path.get().strip()
        if not dbc_path or not os.path.isfile(dbc_path):
            dbc_path = filedialog.askopenfilename(title="选择用于属性修正的 DBC 文件", filetypes=[("DBC文件", "*.dbc"), ("所有文件", "*.*")])
            if not dbc_path:
                return
            self.dbc_path.set(dbc_path)
        try:
            plan = build_dbc_repair_plan(dbc_path)
        except Exception as exc:
            messagebox.showerror("属性扫描失败", str(exc))
            return
        if not plan.items:
            messagebox.showinfo("属性修正", "未发现没有 BA_DEF_ 定义的 BA_ 属性赋值。")
            return

        preview = "\n".join(
            f"第 {item.line_number} 行：{item.attribute_name}（{item.scope}）"
            for item in plan.items[:20]
        )
        suffix = "\n……" if len(plan.items) > 20 else ""
        approved = messagebox.askyesno(
            "属性修正预览",
            f"发现 {len(plan.items)} 条没有 BA_DEF_ 定义的属性赋值：\n\n{preview}{suffix}\n\n"
            "这些属性没有可靠的类型、作用域或枚举定义，工具不会伪造 BA_DEF_。\n"
            "若确认，工具将只在另存副本中删除上述 BA_ 赋值，使该未定义属性不再阻止 DBC 导入。\n"
            "原 DBC 不会修改。是否继续另存？",
        )
        if not approved:
            return
        output = filedialog.asksaveasfilename(
            title="另存属性修正后的 DBC",
            initialdir=str(Path(dbc_path).parent),
            initialfile=f"{Path(dbc_path).stem}_attributes_repaired{Path(dbc_path).suffix}",
            defaultextension=".dbc",
            filetypes=[("DBC文件", "*.dbc")],
        )
        if not output:
            return
        try:
            saved_path, applied = apply_dbc_repair_plan(
                plan, output, remove_undefined_attributes=True,
            )
            repaired_db = parse_dbc(saved_path)
            remaining = [
                item for item in self_check_database(repaired_db, "DBC")
                if item.rule_id == "DBC_ATTR_UNDEFINED_001"
            ]
            if remaining:
                raise ValueError("另存副本仍包含未定义属性，未报告修复成功。")
        except Exception as exc:
            messagebox.showerror("属性修正失败", str(exc))
            return
        self.status_var.set(f"属性修正完成：已另存 {Path(saved_path).name}，移除 {len(applied)} 条无定义 BA_ 赋值。")
        messagebox.showinfo(
            "属性修正完成",
            f"已保存：\n{saved_path}\n\n已移除 {len(applied)} 条没有 BA_DEF_ 定义的属性赋值。\n"
            "请重新执行“开始检查”查看该副本的其余问题。",
        )

    def choose_dbc(self) -> None:
        path = filedialog.askopenfilename(title="选择 DBC 文件", filetypes=[("DBC文件", "*.dbc"), ("所有文件", "*.*")])
        if path:
            self.dbc_path.set(path)
            self.status_var.set(f"已选择 DBC：{Path(path).name}")

    def choose_rules(self) -> None:
        path = filedialog.askopenfilename(
            title="选择语义规则文件",
            filetypes=[("JSON规则文件", "*.json"), ("所有文件", "*.*")],
        )
        if path:
            self.rules_path.set(path)
            try:
                rules = load_rules(path)
                enabled = sum(1 for rule in rules if rule.get("enabled", True))
                self.status_var.set(f"已选择规则文件：{Path(path).name}，启用 {enabled} 条规则。")
            except Exception as exc:
                messagebox.showerror("规则文件错误", str(exc))

    def show_rules(self) -> None:
        try:
            rules = load_rules(self.rules_path.get().strip() or None)
        except Exception as exc:
            messagebox.showerror("规则文件错误", str(exc))
            return
        lines = [
            f"【内置基线】{VECTOR_MANUAL_TITLE} v{VECTOR_MANUAL_VERSION}",
            "覆盖：General、COM、E2E、SecOC、AUTOSAR/OSEK NM、CanTp/DCM、J1939、XCP、Update Bit、SNA、System Signal。",
            "属性BA_DEF_使用0~0占位时，只对定义给出一次警告；实际BA_值按Vector手册范围判断。",
            "",
            "【项目培训规则 can_rules.json】",
        ]
        for rule in rules:
            state = "启用" if rule.get("enabled", True) else "关闭"
            lines.append(
                f"[{state}] {rule.get('id')} / {rule.get('severity', '警告')}\n"
                f"{rule.get('description', '')}"
            )
        messagebox.showinfo("AUTOSAR/CAN语义规则", "\n\n".join(lines) or "规则文件中没有规则。")

    def start_check(self) -> None:
        matrix_path = self.matrix_path.get().strip()
        dbc_path = self.dbc_path.get().strip()
        sheet_name = self.sheet_name.get().strip() or None
        e2e_table_path = self.e2e_table_path.get().strip()
        e2e_sheet_name = self.e2e_sheet_name.get().strip() or None
        rules_path = self.rules_path.get().strip() or None
        semantic_enabled = bool(self.semantic_rules_enabled.get())
        vector_enabled = bool(self.vector_rules_enabled.get())
        dbc_only = bool(self.dbc_only_mode.get())

        if not dbc_only and (not matrix_path or not os.path.isfile(matrix_path)):
            messagebox.showwarning("缺少文件", "请选择有效的 CAN 通信矩阵，或勾选“无CAN矩阵（仅DBC规则检查）”。")
            return
        if not dbc_path or not os.path.isfile(dbc_path):
            messagebox.showwarning("缺少文件", "请选择有效的 DBC 文件。")
            return
        if e2e_table_path and not os.path.isfile(e2e_table_path):
            messagebox.showwarning("E2E ID表无效", "已填写E2E ID表路径，但文件不存在。请重新选择或清空该路径。")
            return
        if e2e_table_path:
            try:
                if e2e_sheet_name == E2E_ALL_SHEETS_LABEL:
                    _sheets, _best, details = inspect_e2e_id_sheets(e2e_table_path)
                    if not any(v.startswith("识别到") for v in details.values()):
                        raise E2EIdTableError("没有任何工作表能识别为E2E ID表。")
                else:
                    preview, _sheet = _read_tabular_preview(e2e_table_path, e2e_sheet_name)
                    detect_e2e_id_layout(preview)
            except Exception as exc:
                messagebox.showerror("E2E ID表无效", str(exc))
                return

        # 矩阵模式下扫描前50行；无矩阵模式跳过此步骤。
        if not dbc_only:
            try:
                validate_matrix_selection(matrix_path, sheet_name)
            except Exception as exc:
                messagebox.showerror("通信矩阵无效", str(exc))
                self.status_var.set("未开始检查：请选择原始 CAN 通信矩阵及正确工作表。")
                return

        self.check_button.configure(state="disabled")
        self.export_button.configure(state="disabled")
        self.fix_e2e_button.configure(state="disabled")
        self.progress.stop()
        self.progress.configure(mode="indeterminate", value=0)
        self.progress.start(10)
        self.run_state_var.set("检查中")
        self.status_var.set("开始检查……")
        self.matrix_db = None
        self.dbc_db = None
        self.differences = []
        self.e2e_table_entries = []
        self.e2e_mapping_info = {}
        self.e2e_header_row = 0
        self.e2e_actual_sheet = ""
        self.refresh_tree()

        def worker() -> None:
            try:
                matrix_db: Optional[Database] = None
                mapping_info: Dict[str, str] = {}
                header_row = 0
                if not dbc_only:
                    self.set_status("正在读取通信矩阵……")
                    matrix_db, mapping_info, header_row = read_matrix(
                        matrix_path,
                        sheet_name,
                        progress=self.set_status,
                    )
                self.set_status("正在解析 DBC……")
                dbc_db = parse_dbc(dbc_path, progress=self.set_status)
                rename_mapping: Optional[Dict[str, Any]] = None
                mapping_path = self.rename_config_path.get().strip()
                if not mapping_path:
                    sibling_mapping = Path(dbc_path).with_suffix(".rename.json")
                    mapping_path = str(sibling_mapping) if sibling_mapping.is_file() else ""
                if mapping_path and os.path.isfile(mapping_path):
                    loaded_mapping = load_rename_config(mapping_path)
                    target_fingerprint = loaded_mapping.get("mapping", {}).get("target_fingerprint", "")
                    if not target_fingerprint or target_fingerprint == hashlib.sha256(Path(dbc_path).read_bytes()).hexdigest():
                        rename_mapping = loaded_mapping
                        self.set_status(f"已加载命名映射：{Path(mapping_path).name}；用于避免重命名后的矩阵假缺失。")
                    else:
                        self.set_status("命名映射与当前DBC指纹不一致，已忽略旧映射并继续检查。")
                semantic_rules: List[Dict[str, Any]] = []
                if semantic_enabled:
                    self.set_status("正在读取AUTOSAR/CAN语义规则……")
                    semantic_rules = load_rules(rules_path)
                if dbc_only:
                    self.set_status("正在执行DBC结构、属性和经验规则检查……")
                    diffs = check_dbc_only(dbc_db, semantic_rules, vector_rules_enabled=vector_enabled)
                else:
                    self.set_status("正在逐项比较报文、信号并执行规则……")
                    assert matrix_db is not None
                    diffs = compare_databases(
                        matrix_db,
                        dbc_db,
                        semantic_rules,
                        vector_rules_enabled=vector_enabled,
                        rename_mapping=rename_mapping,
                    )

                e2e_entries: List[E2EIdEntry] = []
                e2e_mapping: Dict[str, str] = {}
                e2e_header = 0
                e2e_actual_sheet = ""
                if e2e_table_path:
                    self.set_status("正在读取客户E2E ID表并核对SignalGroup/Data ID……")
                    e2e_entries, e2e_mapping, e2e_header, e2e_actual_sheet = read_e2e_id_table(
                        e2e_table_path, e2e_sheet_name
                    )
                    diffs.extend(compare_e2e_id_table(e2e_entries, dbc_db))

                enabled_count = sum(1 for rule in semantic_rules if rule.get("enabled", True))
                self._post_ui(
                    self._finish_check,
                    matrix_db, dbc_db, diffs, mapping_info, header_row, enabled_count, dbc_only, vector_enabled,
                    e2e_entries, e2e_mapping, e2e_header, e2e_actual_sheet,
                )
            except Exception as exc:
                detail = traceback.format_exc()
                # 不使用闭包捕获异常变量，避免异常回调丢失。
                self._post_ui(self._check_failed, exc, detail)

        threading.Thread(target=worker, daemon=True).start()

    def _finish_check(
        self,
        matrix_db: Optional[Database],
        dbc_db: Database,
        diffs: List[Difference],
        mapping_info: Dict[str, str],
        header_row: int,
        enabled_rules_count: int,
        dbc_only: bool,
        vector_enabled: bool,
        e2e_entries: List[E2EIdEntry],
        e2e_mapping: Dict[str, str],
        e2e_header: int,
        e2e_actual_sheet: str,
    ) -> None:
        self.matrix_db = matrix_db
        self.dbc_db = dbc_db
        self.differences = diffs
        self.mapping_info = mapping_info
        self.header_row = header_row
        self.active_rules_count = enabled_rules_count
        self.e2e_table_entries = e2e_entries
        self.e2e_mapping_info = e2e_mapping
        self.e2e_header_row = e2e_header
        self.e2e_actual_sheet = e2e_actual_sheet
        self.progress.stop()
        self.progress.configure(mode="determinate", maximum=100, value=100)
        self.run_state_var.set("已完成")
        self.check_button.configure(state="normal")
        self.export_button.configure(state="normal")
        auto_fix_count = len(collect_e2e_data_length_autofixes(dbc_db))
        self.fix_e2e_button.configure(state="normal" if auto_fix_count > 0 else "disabled")
        stats = stats_for(diffs)
        self.summary_var.set(
            f"错误：{stats['错误']}    警告：{stats['警告']}    提示：{stats['提示']}    总计：{stats['总计']}"
        )
        if dbc_only:
            self.status_var.set(
                f"DBC经验检查完成：{len(dbc_db.messages)} 条报文 / {sum(len(m.signals) for m in dbc_db.messages)} 个信号；"
                f"Vector手册规则：{'启用' if vector_enabled else '关闭'}；项目规则 {enabled_rules_count} 条；"
                f"E2E ID表：{len(e2e_entries)} 条。"
                "无矩阵模式只能发现结构和规则问题，不能证明与客户需求一致。"
            )
        else:
            assert matrix_db is not None
            self.status_var.set(
                f"检查完成：矩阵 {len(matrix_db.messages)} 条报文 / {sum(len(m.signals) for m in matrix_db.messages)} 个信号；"
                f"DBC {len(dbc_db.messages)} 条报文 / {sum(len(m.signals) for m in dbc_db.messages)} 个信号；"
                f"Vector手册规则：{'启用' if vector_enabled else '关闭'}；项目规则 {enabled_rules_count} 条；"
                f"E2E ID表：{len(e2e_entries)} 条。"
            )
        self.refresh_tree()

    def auto_fix_all_e2e(self) -> None:
        """把高置信度E2EDataLength建议值一键写回当前DBC。"""
        dbc_path = self.dbc_path.get().strip()
        if not dbc_path or not os.path.isfile(dbc_path):
            messagebox.showwarning("缺少DBC", "请选择有效的DBC文件。")
            return
        if self.dbc_db is None:
            messagebox.showinfo("尚未检查", "请先执行一次检查，再使用全部一键修改E2E。")
            return

        fixes = collect_e2e_data_length_autofixes(self.dbc_db)
        if not fixes:
            messagebox.showinfo(
                "没有可自动修改项",
                "当前没有能够从DBC内部唯一确定的E2EDataLength。\n\n"
                "多个SignalGroup或只能用DLC×8估算的低置信度建议不会自动写回。\n"
                "结果表中只有显示‘一键修改’的E2E项才支持单条直接修改。",
            )
            self.fix_e2e_button.configure(state="disabled")
            return

        preview_lines = []
        for msg, bits, _reason in fixes[:12]:
            can_hex = f"0x{msg.can_id:X}" if msg.can_id is not None else ""
            can_dec = str(msg.can_id) if msg.can_id is not None else ""
            preview_lines.append(f"{can_hex} ({can_dec}) {msg.name}: -> {bits} bit")
        if len(fixes) > 12:
            preview_lines.append(f"……另有 {len(fixes) - 12} 项")

        ok = messagebox.askyesno(
            "全部一键修改E2E",
            "将批量修改当前DBC中的E2EDataLength，并在同目录自动创建备份。\n"
            "只修改能够由唯一SignalGroup确定的高置信度项目；模糊项会跳过。\n\n"
            + "\n".join(preview_lines)
            + "\n\n是否继续？",
        )
        if not ok:
            return

        try:
            backup_path, changes = apply_e2e_data_length_autofixes(dbc_path, self.dbc_db, fixes)
        except Exception as exc:
            messagebox.showerror("一键修改失败", str(exc))
            return

        self.status_var.set(f"E2E批量一键修改完成：修改 {len(changes)} 项；已备份 {Path(backup_path).name}。正在重新检查……")
        messagebox.showinfo(
            "修改完成",
            f"已修改 {len(changes)} 项 E2EDataLength。\n"
            f"原DBC备份：\n{backup_path}\n\n"
            "程序将自动重新检查修改后的DBC。",
        )
        self.start_check()

    def _on_tree_action_click(self, event: tk.Event) -> None:
        """点击结果表“操作”列时执行当前行的安全修复或配置弹窗。"""
        try:
            if self.tree.identify_region(event.x, event.y) != "cell":
                return
            row_id = self.tree.identify_row(event.y)
            column_id = self.tree.identify_column(event.x)
            columns = list(self.tree["columns"])
            action_column = f"#{columns.index('action') + 1}"
            if not row_id or column_id != action_column:
                return
            if row_id in self._tree_item_e2e_fixes:
                self.auto_fix_one_e2e(row_id)
            elif row_id in self._tree_item_autofix_diffs:
                self.auto_fix_one_difference(row_id)
        except (tk.TclError, ValueError):
            return

    def auto_fix_one_e2e(self, tree_item_id: str) -> None:
        """只修改结果表中当前这一条高置信度E2EDataLength建议。"""
        dbc_path = self.dbc_path.get().strip()
        if not dbc_path or not os.path.isfile(dbc_path):
            messagebox.showwarning("缺少DBC", "请选择有效的DBC文件。")
            return
        if self.dbc_db is None:
            messagebox.showinfo("尚未检查", "请先执行一次检查。")
            return
        fix = self._tree_item_e2e_fixes.get(tree_item_id)
        if fix is None:
            return

        msg, bits, reason = fix
        try:
            backup_path, changes = apply_e2e_data_length_autofixes(
                dbc_path, self.dbc_db, [fix]
            )
        except Exception as exc:
            messagebox.showerror("单条E2E修改失败", str(exc))
            return

        can_text = f"0x{msg.can_id:X} ({msg.can_id})" if msg.can_id is not None else msg.name
        self.status_var.set(
            f"已修改 {can_text} {msg.name}: E2EDataLength={bits} bit；"
            f"备份 {Path(backup_path).name}。正在重新检查……"
        )
        # 单条修改不再弹二次确认/完成框，真正做到点一次即修改；原文件始终有备份。
        self.start_check()

    def auto_fix_one_difference(self, tree_item_id: str) -> None:
        """处理节点、周期及可由现有定义确定的显式 BO_ 属性。"""
        diff = self._tree_item_autofix_diffs.get(tree_item_id)
        if diff is None:
            return
        if diff.rule_id in {"DBC_NODE_REF_001", "DBC_NODE_REF_002"}:
            # 节点名称已经在 BO_/SG_ 中明确出现，节点补全窗口会显示全部缺失节点并另存。
            self.open_node_completion()
            return
        if diff.rule_id == "DBC_ATTR_UNDEFINED_001":
            # 没有 BA_DEF_ 时无法猜类型；复用属性处理窗口向用户说明并要求确认。
            self.open_attribute_repair()
            return

        dbc_path = self.dbc_path.get().strip()
        if not dbc_path or not os.path.isfile(dbc_path) or self.dbc_db is None:
            messagebox.showwarning("缺少DBC", "请先选择 DBC 并完成一次检查。")
            return
        message = next(
            (item for item in self.dbc_db.messages
             if item.can_id == diff.can_id and item.name == diff.message_name),
            None,
        )
        if message is None:
            messagebox.showerror("修复定位失败", "检查结果对应的报文已变化；请重新检查后再修复。")
            return

        attribute_name = ""
        value: Any = None
        if diff.rule_id == "DBC_TX_CONFLICT_001":
            attribute_name, value = "GenMsgCycleTime", 0
        elif diff.rule_id in {"DBC_TX_EXPLICIT_001", "DBC_TX_EXPLICIT_002"}:
            attribute_name, value = diff.field_name, diff.dbc_value
            if value in (None, ""):
                messagebox.showinfo("需要填写", f"{message.name} 的 {attribute_name} 没有可继承的有效值，请填写后再修复。")
                return
        elif diff.rule_id == "DBC_TX_CYCLE_001":
            value = simpledialog.askinteger(
                "填写周期", f"{message.name}（{format_can_id(message.can_id)}）需要大于 0 的 GenMsgCycleTime（ms）：",
                parent=self.root, minvalue=1,
            )
            if value is None:
                return
            attribute_name = "GenMsgCycleTime"
        else:
            return

        try:
            backup_path, description = apply_message_attribute_autofix(
                dbc_path, self.dbc_db, message, attribute_name, value,
            )
        except Exception as exc:
            messagebox.showerror("一键修复失败", str(exc))
            return
        self.status_var.set(f"一键修复完成：{description} 已备份 {Path(backup_path).name}。正在重新检查……")
        self.start_check()

    def _check_failed(self, exc: Exception, detail: str) -> None:
        self.progress.stop()
        self.progress.configure(mode="determinate", maximum=100, value=0)
        self.run_state_var.set("失败")
        self.check_button.configure(state="normal")
        self.fix_e2e_button.configure(state="disabled")
        self.status_var.set("检查失败。")
        print(detail, file=sys.stderr)
        messagebox.showerror("检查失败", f"{exc}\n\n详细错误已输出到控制台。")

    def _set_detail_text(self, content: str) -> None:
        self.detail_text.configure(state="normal")
        self.detail_text.delete("1.0", "end")
        self.detail_text.insert("1.0", content)
        self.detail_text.configure(state="disabled")

    def _show_selected_detail(self, _event: Optional[tk.Event] = None) -> None:
        selection = self.tree.selection()
        if not selection:
            self._set_detail_text("单击上方任意检查项，可在这里查看完整内容。")
            return
        self._set_detail_text(self._tree_item_details.get(selection[0], ""))

    def _matches_type_filter(self, diff: Difference) -> bool:
        selected = self.type_filter_var.get()
        if selected == "全部":
            return True
        category = (diff.category or "").lower()
        rule_id = (diff.rule_id or "").upper()
        field = (diff.field_name or "").lower().replace("_", "")
        message = (diff.message_name or "").lower()
        signal = (diff.signal_name or "").lower()
        description = (diff.description or "").lower()

        is_e2e_id_table = diff.category == "E2E ID表核对"
        is_e2e_length = "e2edatalength" in field or "E2E_LENGTH" in rule_id or rule_id == "E2E_ATTR_001"
        is_e2e = is_e2e_id_table or is_e2e_length or "E2E" in rule_id or "e2e" in field
        is_matrix = diff.category in {"报文差异", "信号差异", "DBC多余报文", "矩阵自身检查", "矩阵规则检查"}
        is_nm = (
            "NM_" in rule_id or rule_id.startswith("VEC_NM") or "nmasr" in field or
            re.search(r"(^|[_-])(can)?nm([_-]|$)", message, flags=re.IGNORECASE) is not None
        )
        is_diag = (
            "UDS" in rule_id or "DIAG" in rule_id or "DCM" in rule_id or
            field.startswith("diag") or any(x in message for x in ("diag", "uds", "obd"))
        )
        is_xcp = "XCP" in rule_id or any(x in message for x in ("xcp", "ccp"))

        if selected == "DBC自身/规则":
            return not is_matrix and not is_e2e_id_table
        if selected == "矩阵差异":
            return is_matrix
        if selected == "E2E全部":
            return is_e2e
        if selected == "E2E DataID表":
            return is_e2e_id_table
        if selected == "E2E DataLength":
            return is_e2e_length
        if selected == "NM":
            return is_nm
        if selected == "UDS/诊断":
            return is_diag
        if selected == "XCP":
            return is_xcp
        return True

    def refresh_tree(self) -> None:
        for item in self.tree.get_children():
            self.tree.delete(item)
        self._tree_item_details.clear()
        self._tree_item_e2e_fixes.clear()
        self._tree_item_autofix_diffs.clear()
        self._set_detail_text("单击上方任意检查项，可在这里查看完整内容；“操作”列出现按钮文字时，可直接点击修复或填写缺失值。")
        e2e_fix_by_can_id: Dict[int, Tuple[Message, int, str]] = {}
        if self.dbc_db is not None:
            for fix in collect_e2e_data_length_autofixes(self.dbc_db):
                fix_msg = fix[0]
                if fix_msg.can_id is not None:
                    e2e_fix_by_can_id[int(fix_msg.can_id)] = fix
        e2e_action_seen: set[int] = set()
        selected = self.filter_var.get()
        if not self.differences and self.dbc_db is not None:
            if selected in ("全部", "通过"):
                description = (
                    "检查已完成，未发现已启用规则能够识别的问题。无矩阵模式不代表DBC与客户需求完全一致。"
                    if self.matrix_db is None
                    else "检查已完成，未发现矩阵与 DBC 差异，也未命中已启用的语义规则。"
                )
                iid = self.tree.insert(
                    "", "end",
                    values=(
                        "通过", "检查结果", "CHECK_OK", "", "", "", "",
                        "一致性检查", "", "", "", wrap_tree_cell(description, 42, 3),
                    ),
                    tags=("通过",),
                )
                self._tree_item_details[iid] = description
            return
        inserted_count = 0
        for diff in self.differences:
            if selected != "全部" and diff.severity != selected:
                continue
            if not self._matches_type_filter(diff):
                continue
            action_text = ""
            action_fix: Optional[Tuple[Message, int, str]] = None
            generic_action = False
            can_id_int = int(diff.can_id) if diff.can_id is not None else None
            # 同一报文可能同时命中DBC自检和Vector规则。每个E2E只显示一个“一键修改”，避免重复按钮。
            if (
                can_id_int is not None
                and can_id_int in e2e_fix_by_can_id
                and can_id_int not in e2e_action_seen
                and "e2edatalength" in str(diff.field_name or "").replace("_", "").lower()
            ):
                action_text = "一键修改"
                action_fix = e2e_fix_by_can_id[can_id_int]
                e2e_action_seen.add(can_id_int)
            elif diff.rule_id in {"DBC_NODE_REF_001", "DBC_NODE_REF_002"}:
                action_text = "补全节点"
                generic_action = True
            elif diff.rule_id == "DBC_ATTR_UNDEFINED_001":
                action_text = "处理属性"
                generic_action = True
            elif diff.rule_id in {"DBC_TX_CONFLICT_001", "DBC_TX_EXPLICIT_001", "DBC_TX_EXPLICIT_002"}:
                action_text = "一键修复"
                generic_action = True
            elif diff.rule_id == "DBC_TX_CYCLE_001":
                action_text = "填写周期"
                generic_action = True
            values = (
                diff.severity,
                wrap_tree_cell(diff.category, 14, 2),
                wrap_tree_cell(diff.rule_id, 18, 2),
                format_can_id(diff.can_id),
                format_can_id_decimal(diff.can_id),
                wrap_tree_cell(diff.message_name, 22, 3),
                wrap_tree_cell(diff.signal_name, 24, 3),
                wrap_tree_cell(diff.field_name, 14, 2),
                wrap_tree_cell(diff.matrix_value, 24, 3),
                wrap_tree_cell(diff.dbc_value, 24, 3),
                action_text,
                wrap_tree_cell(diff.description, 42, 3),
            )
            iid = self.tree.insert("", "end", values=values, tags=(diff.severity,))
            inserted_count += 1
            detail = difference_detail_text(diff)
            if action_fix is not None:
                _msg, bits, reason = action_fix
                detail += f"\n操作：一键修改 E2EDataLength -> {bits} bit\n建议依据：{reason}"
                self._tree_item_e2e_fixes[iid] = action_fix
            elif generic_action:
                self._tree_item_autofix_diffs[iid] = diff
                if diff.rule_id in {"DBC_NODE_REF_001", "DBC_NODE_REF_002"}:
                    detail += "\n操作：补全当前 DBC 已明确引用、但 BU_ 缺失的节点。"
                elif diff.rule_id == "DBC_ATTR_UNDEFINED_001":
                    detail += "\n操作：属性类型不明，打开属性处理窗口，由你确认删除或补充定义。"
                elif diff.rule_id == "DBC_TX_CYCLE_001":
                    detail += "\n操作：周期值无法从 DBC 推断，点击后填写大于 0 的毫秒值。"
                else:
                    detail += "\n操作：根据当前已定义的属性值直接写回；原 DBC 会自动备份。"
            self._tree_item_details[iid] = detail

        if inserted_count == 0 and self.differences:
            desc = f"当前筛选条件无匹配项：级别={self.filter_var.get()}，类型={self.type_filter_var.get()}。"
            iid = self.tree.insert(
                "", "end",
                values=("提示", "筛选结果", "FILTER_EMPTY", "", "", "", "", "", "", "", "", desc),
                tags=("提示",),
            )
            self._tree_item_details[iid] = desc

    def show_mapping(self) -> None:
        if not self.mapping_info:
            if self.dbc_db is not None and self.matrix_db is None:
                messagebox.showinfo("列映射", "当前为无CAN矩阵模式，没有Excel列映射。")
            else:
                messagebox.showinfo("列映射", "尚未执行矩阵检查。执行检查后可查看自动识别的列映射。")
            return
        aliases_cn = {
            "message_id": "报文ID", "message_name": "报文名称", "dlc": "DLC", "sender": "发送节点",
            "receiver": "接收节点", "cycle_time": "周期时间", "send_type": "发送类型", "frame_format": "帧格式",
            "message_comment": "报文描述", "signal_name": "信号名称", "start_bit": "起始位",
            "start_byte": "起始字节", "bit_in_byte": "字节内起始位", "signal_length": "信号长度",
            "byte_order": "字节序", "signed": "符号类型", "factor": "精度/Factor", "offset": "偏移量",
            "minimum": "最小值", "maximum": "最大值", "unit": "单位", "initial_value": "初始值",
            "invalid_value": "无效值", "signal_comment": "信号描述", "value_table": "枚举值", "multiplex": "复用",
        }
        lines = [f"识别到的表头行：第 {self.header_row} 行", ""]
        for canonical, original in sorted(self.mapping_info.items()):
            lines.append(f"{aliases_cn.get(canonical, canonical)}  ←  {original}")
        messagebox.showinfo("自动列映射", "\n".join(lines))

    def export_report(self) -> None:
        if self.dbc_db is None:
            messagebox.showwarning("无结果", "请先执行检查。")
            return

        base_name = "DBC无矩阵规则检查报告" if self.matrix_db is None else "CAN矩阵_DBC差异报告"
        if Workbook is not None:
            path = filedialog.asksaveasfilename(
                title="导出检查报告",
                defaultextension=".xlsx",
                initialfile=base_name + ".xlsx",
                filetypes=[("Excel报告", "*.xlsx"), ("CSV报告", "*.csv")],
            )
        else:
            path = filedialog.asksaveasfilename(
                title="导出检查报告",
                defaultextension=".csv",
                initialfile=base_name + ".csv",
                filetypes=[("CSV报告", "*.csv")],
            )
        if not path:
            return
        try:
            if Path(path).suffix.lower() == ".csv":
                export_csv_report(path, self.differences, self.matrix_db, self.dbc_db)
            else:
                export_xlsx_report(path, self.differences, self.matrix_db, self.dbc_db)
            self.status_var.set(f"报告已导出：{path}")
            messagebox.showinfo("导出完成", f"报告已保存到：\n{path}")
        except Exception as exc:
            messagebox.showerror("导出失败", str(exc))


def main() -> None:
    if tk is None:
        raise SystemExit("当前 Python 未包含 tkinter，无法启动图形界面；核心解析和离线自测仍可运行。")
    root = tk.Tk()
    CheckerApp(root)
    root.mainloop()


if __name__ == "__main__":
    main()
